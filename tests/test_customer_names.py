import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from checker import ColumnMapping, Options, Source, check_files, compare_value
from test_sbf_salad_viber import BLAST_HEADERS, salad_database, salad_row
from test_sbf_pl_viber import BLAST_HEADERS as PL_HEADERS, pl_database, pl_row
from test_web_checker import csv_upload, xlsx_upload
from web_checker import UploadedSource, prepare_uploaded_check, run_prepared_check, run_uploaded_check


class CustomerNameTests(unittest.TestCase):
    def test_replacement_glyph_matches_only_enye_on_the_same_debtor_record(self):
        database = pl_database([
            ["Ana Peña", "00123", "000001", "a@example.test", "09171234567", 1001],
            ["Bob Jones", "456", "000002", "b@example.test", "09181234567", 1002],
        ])
        rows = [
            pl_row(name="Ana Pe\ufffda"),
            pl_row(name="Ana Pe\ufffda", phone="09181234567"),
            pl_row(name="Ana Pe\ufffde"),
            pl_row(name="Ana Pena"),
            pl_row("1002", "Ana Pe\ufffda", "09181234567"),
            pl_row("unknown", "Ana Pe\ufffda"),
        ]
        run = run_prepared_check(prepare_uploaded_check(database, [
            csv_upload("viber.csv", [PL_HEADERS, *rows]),
        ]))
        result = run.results[0]
        self.assertEqual(result.error, "")
        self.assertEqual([row[1] for row in result.records],
                         ["MATCHED", "DIFFERENT", "DIFFERENT", "DIFFERENT", "DIFFERENT", "NOT FOUND"])
        self.assertEqual([difference[2].split(" ")[0] for difference in result.differences],
                         ["Phone", "FullName", "FullName", "FullName"])
        self.assertEqual(result.records[0][3][0], "Ana Pe\ufffda")
        book = load_workbook(io.BytesIO(run.report))
        try:
            self.assertEqual(book["File 001"]["D2"].value, "Ana Pe\ufffda")
            rules = [row[4] for row in book["Mappings 001"].iter_rows(min_row=2, values_only=True)]
            self.assertTrue(any("U+FFFD" in rule and "Ñ/ñ" in rule for rule in rules))
        finally:
            book.close()

    def test_replacement_glyph_is_not_a_wildcard_or_removed_from_other_fields(self):
        headers = ["CHCODE", "Name", "Email", "Note"]
        database = csv_upload("database.csv", [
            headers, ["00123", "Ana Pena", "peña@example.test", "Peña"],
            ["001Ñ", "Ana Peña", "peña@example.test", "Peña"],
        ])
        upload = csv_upload("glyphs.csv", [
            headers,
            ["00123", "Ana Pe\ufffda", "peña@example.test", "Peña"],
            ["001Ñ", "Ana Pe\ufffda", "pe\ufffda@example.test", "Pe\ufffda"],
            ["001\ufffd", "Ana Peña", "peña@example.test", "Peña"],
        ])
        result = run_prepared_check(prepare_uploaded_check(database, [upload])).results[0]
        self.assertEqual([row[1] for row in result.records], ["DIFFERENT", "DIFFERENT", "NOT FOUND"])
        self.assertEqual([difference[2] for difference in result.differences], ["Name", "Email", "Note"])

    def test_replacement_glyph_supports_explicit_name_mappings_and_case_options(self):
        database = UploadedSource(csv_upload("database.csv", [
            ["CHCODE", "Name"], ["00123", "Ana Peña"],
        ]), "CSV", "CHCODE")
        source = UploadedSource(csv_upload("mapped.csv", [
            ["CHCODE", "Recipient"], ["00123", "Ana Pe\ufffda"],
            ["00123", "ana Pe\ufffda"], ["00123", "Ana PEÑA"],
        ]), "CSV", "CHCODE", (ColumnMapping("CHCODE", "CHCODE"), ColumnMapping("Recipient", "Name")))
        result = run_uploaded_check(database, [source], Options(ignore_case=False, mapped_fields=True)).results[0]
        self.assertEqual([row[1] for row in result.records], ["MATCHED", "DIFFERENT", "DIFFERENT"])
        # A damaged database name is compared by the same rule.
        database = UploadedSource(csv_upload("database.csv", [
            ["CHCODE", "Name"], ["00123", "Ana Pe\ufffda"],
        ]), "CSV", "CHCODE")
        source = UploadedSource(csv_upload("mapped.csv", [
            ["CHCODE", "Name"], ["00123", "Ana Peña"],
        ]), "CSV", "CHCODE")
        self.assertEqual(run_uploaded_check(database, [source], Options(mapped_fields=True)).results[0].count("MATCHED"), 1)

    def test_full_copy_reserves_exact_and_damaged_name_rows_without_reusing_duplicates(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            headers = ["CHCODE", "Name", "Note"]
            database = root / "database.csv"
            source = root / "source.csv"
            database.write_bytes(csv_upload(database.name, [
                headers, ["00123", "Ana Peña", "Ready"], ["00123", "Ana Peña", "Hold"],
            ]).data)
            rows = [headers, ["00123", "Ana Pe\ufffda", "Ready"], ["00123", "Ana Peña", "Hold"]]
            source.write_bytes(csv_upload(source.name, rows).data)
            db, target = Source(database, "CSV", "CHCODE"), Source(source, "CSV", "CHCODE")
            result = check_files(db, [target], Options())[0]
            self.assertEqual(result.verdict, "SAME ROWS")
            self.assertEqual([row[:3] for row in result.records], [(2, "MATCHED", "2"), (3, "MATCHED", "3")])
            source.write_bytes(csv_upload(source.name, [*rows, ["00123", "Ana Pe\ufffda", "Ready"]]).data)
            result = check_files(db, [target], Options())[0]
            self.assertEqual([row[1] for row in result.records], ["MATCHED", "MATCHED", "EXTRA ROW"])
            self.assertEqual(result.missing_rows, [])

    def test_question_marks_are_ignored_on_both_sides_with_existing_case_and_space_options(self):
        self.assertEqual(compare_value(" ?Al?ice Smith?? ", "Text", Options(), customer_name=True), "alice smith")
        self.assertEqual(compare_value("?Alice?", "Text", Options(False, False), customer_name=True), "Alice")
        self.assertEqual(compare_value(" Alice? ", "Text", Options(False, False), customer_name=True), " Alice ")
        self.assertNotEqual(compare_value("Alice?", "Text", Options(ignore_case=False), customer_name=True),
                            compare_value("alice", "Text", Options(ignore_case=False), customer_name=True))
        self.assertNotEqual(compare_value("Anne-Marie?", "Text", Options(), customer_name=True),
                            compare_value("AnneMarie", "Text", Options(), customer_name=True))
        self.assertEqual(compare_value(None, "Text", Options(), customer_name=True), "")
        self.assertEqual(compare_value("test?@example.test", "Text", Options()), "test?@example.test")

    def test_all_supported_name_fields_ignore_questions_and_preserve_report_values(self):
        headers = ["CHCODE", "Name", "FullName", "ChName", "First Name", "Last Name", "Middle Name",
                   "chname_format", "Customer Name", "Name"]
        database = xlsx_upload("database.xlsx", [headers, ["00123", *(["Al?ice Smith"] * 9)]])
        upload = csv_upload("names.csv", [headers, ["00123", *(["Alice? Smith??"] * 9)]])
        run = run_prepared_check(prepare_uploaded_check(database, [upload]))
        result = run.results[0]
        self.assertEqual(result.error, "")
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertEqual(len(result.mappings), 10)
        book = load_workbook(io.BytesIO(run.report))
        try:
            self.assertEqual(book["File 001"]["E2"].value, "Alice? Smith??")
            coverage = list(book["Mappings 001"].values)[1:]
            self.assertTrue(all("question marks are ignored" in row[4] for row in coverage[1:]))
            self.assertNotIn("question marks", coverage[0][4])
        finally:
            book.close()

    def test_viber_name_matches_stay_on_the_same_phone_record(self):
        upload = csv_upload("viber.csv", [
            BLAST_HEADERS, salad_row("Al?ice Smith??"),
            salad_row("?Bob Jones", "09171234567"),
        ])
        result = run_prepared_check(prepare_uploaded_check(salad_database(), [upload])).results[0]
        self.assertEqual([row[:3] for row in result.records], [(2, "MATCHED", "2"), (3, "DIFFERENT", "2")])
        self.assertEqual(result.differences[0][3:], ("?Bob Jones", "Alice Smith"))

    def test_question_marks_in_keys_emails_and_notes_still_produce_issues(self):
        headers = ["CHCODE", "Name", "Email", "Note"]
        database = csv_upload("database.csv", [headers, ["00123", "Alice", "alice@example.test", "Ready"]])
        upload = csv_upload("other_fields.csv", [
            headers,
            ["00123", "Alice?", "alice?@example.test", "Ready?"],
            ["00123?", "Alice", "alice@example.test", "Ready"],
        ])
        result = run_prepared_check(prepare_uploaded_check(database, [upload])).results[0]
        self.assertEqual([row[1] for row in result.records], ["DIFFERENT", "NOT FOUND"])
        self.assertEqual([difference[2] for difference in result.differences], ["Email", "Note"])

    def test_explicit_customer_name_mapping_also_ignores_questions(self):
        database = UploadedSource(csv_upload("database.csv", [
            ["CHCODE", "Name"], ["00123", "Alice?"],
        ]), "CSV", "CHCODE")
        source = UploadedSource(csv_upload("mapped.csv", [
            ["CHCODE", "Recipient"], ["00123", "?Alice"],
        ]), "CSV", "CHCODE", (ColumnMapping("CHCODE", "CHCODE"), ColumnMapping("Recipient", "Name")))
        result = run_uploaded_check(database, [source], Options(mapped_fields=True)).results[0]
        self.assertEqual(result.error, "")
        self.assertEqual(result.count("MATCHED"), 1)

    def test_full_copy_ignores_name_questions_and_keeps_duplicate_counts(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            database = root / "database.csv"
            source = root / "source.csv"
            headers = ["CHCODE", "Name", "Note"]
            database.write_bytes(csv_upload(database.name, [headers, ["00123", "Alice?", "Ready"],
                                                          ["00123", "?Alice", "Ready"]]).data)
            source.write_bytes(csv_upload(source.name, [headers, ["00123", "Alice", "Ready"],
                                                       ["00123", "Alice??", "Ready"]]).data)
            db = Source(database, "CSV", "CHCODE")
            target = Source(source, "CSV", "CHCODE")
            result = check_files(db, [target], Options())[0]
            self.assertEqual(result.verdict, "SAME ROWS")
            source.write_bytes(csv_upload(source.name, [headers, ["00123", "Alice", "Ready"],
                                                       ["00123", "Alice??", "Ready"],
                                                       ["00123", "Alice?", "Ready?"]]).data)
            result = check_files(db, [target], Options())[0]
            self.assertEqual(result.count("MATCHED"), 2)
            self.assertEqual(result.count("EXTRA ROW"), 1)

    def test_full_copy_retains_real_name_and_other_field_differences(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            headers = ["CHCODE", "Name", "Note"]
            database = root / "database.csv"
            target = root / "source.csv"
            database.write_bytes(csv_upload(database.name, [headers, ["00123", "Alice?", "Ready"]]).data)
            target.write_bytes(csv_upload(target.name, [headers, ["00123", "Bob?", "Ready?"]]).data)
            result = check_files(Source(database, "CSV", "CHCODE"), [Source(target, "CSV", "CHCODE")], Options())[0]
            self.assertEqual(result.count("DIFFERENT"), 1)
            self.assertEqual([row[2] for row in result.differences], ["Name", "Note"])


if __name__ == "__main__":
    unittest.main()
