import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook, load_workbook

from checker import (
    ColumnMapping, Options, Source, check_files, compare_value, export_report,
    is_filename_identified_pair, is_hsbc_mobile_pair, match_column, suggest_mappings,
)


class HSBCMobileTests(unittest.TestCase):
    database_headers = ["CHCODE", "Mobile phone", "Name", "Last Name", "PROD TYPE"]
    campaign_headers = ["formatted mobile", "ChCode", "ChName", "LastName", "ProductType"]
    alice = ["0917 123 4567", "00123", "Alice Smith", "Smith", "Credit Card"]

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.options = Options(mapped_fields=True, require_flow_stage=True)

    def workbook(self, name, headers, rows):
        path = self.folder / name
        book = Workbook()
        book.active.append(headers)
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        return Source(path, "Sheet", match_column(headers))

    def database(self, headers=None, rows=None):
        return self.workbook("HSBC DB.xlsx", headers or self.database_headers, rows if rows is not None else [
            ["00123", 9171234567, "Alice Smith", "Smith", "CARDS"],
            ["00456", "+639189876543", "Bob Jones", "Jones", "LOAN"],
        ])

    def campaign(self, rows=None, headers=None, name="BLAST Field.xlsx"):
        return self.workbook(name, headers or self.campaign_headers,
                             rows if rows is not None else [self.alice])

    def test_mobile_blast_without_stage_checks_all_five_fields_and_uses_filename(self):
        campaign = self.campaign()
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertEqual(result.error, "")
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertFalse(result.flow_stage_required)
        self.assertEqual(result.file_identifier, "BLAST Field.xlsx")
        self.assertEqual(result.file_identifier_column, "Filename")
        self.assertEqual(result.mappings, [
            ColumnMapping("formatted mobile", "Mobile phone", "Phone"),
            ColumnMapping("ChCode", "CHCODE"),
            ColumnMapping("ChName", "Name"),
            ColumnMapping("LastName", "Last Name"),
            ColumnMapping("ProductType", "PROD TYPE", "Product"),
        ])
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertEqual(result.records[0][:3], (2, "MATCHED", "2"))
        self.assertEqual(result.ignored_columns, [])
        self.assertEqual(result.missing_rows, [])

    def test_mobile_blast_reports_changed_fields_missing_keys_and_blank_keys(self):
        campaign = self.campaign(rows=[
            self.alice,
            ["0917 123 4567", "00456", "Changed Name", "Wrong surname", "Credit Card"],
            ["0917 123 4567", "00999", "Unknown Person", "Person", "Credit Card"],
            ["0917 123 4567", None, "No Key", "Key", "Credit Card"],
        ])
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertEqual(result.error, "")
        self.assertEqual(result.verdict, "ROW ISSUES")
        self.assertEqual([record[1] for record in result.records],
                         ["MATCHED", "DIFFERENT", "NOT FOUND", "BLANK KEY"])
        self.assertEqual({difference[2] for difference in result.differences}, {
            "formatted mobile → Mobile phone", "ChName → Name",
            "LastName → Last Name", "ProductType → PROD TYPE",
        })
        self.assertTrue(all(difference[:2] == (3, 3) for difference in result.differences))

    def test_both_blasts_can_match_the_same_database_rows(self):
        results = check_files(self.database(), [self.campaign(), self.campaign(name="BLAST NEW ENDO.xlsx")],
                              self.options)
        self.assertEqual([result.verdict for result in results], ["MAPPED ROWS MATCH"] * 2)
        self.assertEqual([result.records[0][2] for result in results], ["2", "2"])
        self.assertEqual(results[1].file_identifier, "BLAST NEW ENDO.xlsx")

    def test_leading_titles_are_ignored_on_either_side_but_surnames_and_suffixes_are_checked(self):
        for actual, expected in [
            ("Smith", "Mr./Ms. Smith"), ("Smith", "mR. / mS. Smith"),
            ("Smith", "Mr. Smith"), ("Smith", "Ms. Smith"),
            ("Mr./Ms. Smith", "Smith"), (" Smith ", " Mr./Ms. Smith "),
            ("Lim Jr", "Mr./Ms. Lim Jr"),
        ]:
            with self.subTest(actual=actual, expected=expected):
                database = self.database(rows=[["00123", "09171234567", "Alice Smith", expected, "CARDS"]])
                row = [*self.alice[:3], actual, self.alice[4]]
                result = check_files(database, [self.campaign([row])], self.options)[0]
                self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
                self.assertEqual(result.differences, [])
        for actual, expected in [
            ("Jones", "Mr./Ms. Smith"), ("Lim", "Mr./Ms. Lim Jr"),
            ("Smith", "Mr./Ms.Smith"), ("Smith", "The Mr./Ms. Smith"),
            ("", "Mr./Ms."), ("Smith", "Dr. Smith"),
        ]:
            with self.subTest(actual=actual, expected=expected):
                database = self.database(rows=[["00123", "09171234567", "Alice Smith", expected, "CARDS"]])
                result = check_files(database, [self.campaign([[*self.alice[:3], actual, self.alice[4]]])], self.options)[0]
                self.assertEqual(result.count("DIFFERENT"), 1)
                self.assertEqual(len(result.differences), 1)
                self.assertEqual(result.differences[0][2], "LastName → Last Name")
        database = self.database(rows=[["00123", "09171234567", "Alice Smith", "Mr./Ms. Smith", "CARDS"]])
        result = check_files(database, [self.campaign([[*self.alice[:3], "smith", self.alice[4]]])],
                             replace(self.options, ignore_case=False))[0]
        self.assertEqual(result.count("DIFFERENT"), 1)

    def test_ignoring_titles_never_skips_other_mapped_fields_or_changes_keys(self):
        database = self.database(rows=[["00123", "09171234567", "Alice Smith", "Mr./Ms. Smith", "CARDS"]])
        rows = [self.alice,
                ["09171234567", "00123", "Mr./Ms. Alice Smith", "Smith", "Credit Card"],
                ["09181234567", "00123", "Alice Smith", "Smith", "Credit Card"],
                ["09171234567", "00123", "Alice Smith", "Smith", "Personal Loan"],
                ["09171234567", "Mr./Ms. 00123", "Alice Smith", "Smith", "Credit Card"]]
        result = check_files(database, [self.campaign(rows)], self.options)[0]
        self.assertEqual([row[1] for row in result.records], ["MATCHED", "DIFFERENT", "DIFFERENT", "DIFFERENT", "NOT FOUND"])
        self.assertEqual({row[2] for row in result.differences},
                         {"ChName → Name", "formatted mobile → Mobile phone", "ProductType → PROD TYPE"})

    def test_titles_in_other_layouts_and_full_copy_are_still_compared(self):
        database = self.database(["CHCODE", "Last Name"], [["00123", "Mr./Ms. Smith"]])
        source = self.campaign([["00123", "Smith"]], ["CHCODE", "LastName"])
        result = check_files(database, [source], Options(mapped_fields=True))[0]
        self.assertEqual(result.count("DIFFERENT"), 1)
        source = self.campaign([["00123", "Smith"]], ["CHCODE", "Last Name"])
        result = check_files(database, [source], Options())[0]
        self.assertEqual(result.count("DIFFERENT"), 1)

    def test_duplicate_chcodes_match_every_field_on_one_row_after_ignoring_titles(self):
        database = self.database(rows=[
            ["00123", "09171234567", "Alice Smith", "Mr./Ms. Smith", "CARDS"],
            ["00123", "09189876543", "Bob Jones", "Mr./Ms. Jones", "LOAN"],
        ])
        source = self.campaign([
            ["09189876543", "00123", "Bob Jones", "Jones", "Personal Loan"],
            ["09189876543", "00123", "Bob Jones", "Smith", "Personal Loan"],
        ])
        result = check_files(database, [source], self.options)[0]
        self.assertEqual([row[:3] for row in result.records], [(2, "MATCHED", "3"), (3, "DIFFERENT", "3")])
        self.assertEqual(result.differences[0][4], "Mr./Ms. Jones")

    def test_title_rule_is_disclosed_and_original_values_remain_in_report(self):
        database = self.database(rows=[["00123", "09171234567", "Alice Smith", "Mr./Ms. Smith", "CARDS"]])
        source = self.campaign([self.alice, [*self.alice[:3], "Jones", self.alice[4]]])
        originals = {item.path: item.path.read_bytes() for item in (database, source)}
        result = check_files(database, [source], self.options)[0]
        path = self.folder / "titles-report.xlsx"
        export_report(path, database, self.options, [result])
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            self.assertEqual(book["File 001"]["G2"].value, "Smith")
            self.assertEqual(book["Differences 001"]["E2"].value, "Mr./Ms. Smith")
            coverage = {row[0]: row[1:] for row in book["Mappings 001"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(coverage["LastName"][2], "CHECKED")
            self.assertIn("Mr./Ms.", coverage["LastName"][3])
            self.assertIn("suffixes remain checked", coverage["LastName"][3])
            self.assertNotIn("titles are ignored", coverage["ChName"][3])
        finally:
            book.close()
        self.assertTrue(all(path.read_bytes() == content for path, content in originals.items()))

    def test_aliases_preserve_header_case_and_product_comparison(self):
        headers = ["FORMATTED MOBILE", "chcode", "chNAME", "LASTNAME", "PRODUCTTYPE"]
        database_headers = ["chcode", "MOBILE PHONE", "NAME", "LAST NAME", "prod type"]
        self.assertEqual(suggest_mappings(headers, database_headers, "chcode", "chcode"), [
            ColumnMapping("FORMATTED MOBILE", "MOBILE PHONE", "Phone"),
            ColumnMapping("chcode", "chcode"),
            ColumnMapping("chNAME", "NAME"),
            ColumnMapping("LASTNAME", "LAST NAME"),
            ColumnMapping("PRODUCTTYPE", "prod type", "Product"),
        ])
        campaign = self.campaign([["09189876543", "00456", "Bob Jones", "Jones", "Personal Loan"]])
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")

    def test_recognition_requires_full_schemas_and_chcode_keys(self):
        self.assertTrue(is_hsbc_mobile_pair(self.campaign_headers, self.database_headers))
        self.assertTrue(is_hsbc_mobile_pair(
            [f" {name.upper()} " for name in self.campaign_headers],
            [f" {name.lower()} " for name in self.database_headers], " chCODE ", " CHcode "))
        for index in range(len(self.campaign_headers)):
            with self.subTest(missing_file_header=self.campaign_headers[index]):
                self.assertFalse(is_hsbc_mobile_pair(
                    self.campaign_headers[:index] + self.campaign_headers[index + 1:], self.database_headers))
        for index in range(len(self.database_headers)):
            with self.subTest(missing_database_header=self.database_headers[index]):
                self.assertFalse(is_hsbc_mobile_pair(
                    self.campaign_headers, self.database_headers[:index] + self.database_headers[index + 1:]))
        self.assertFalse(is_hsbc_mobile_pair(self.campaign_headers, self.database_headers, "ChName", "CHCODE"))
        self.assertFalse(is_hsbc_mobile_pair(self.campaign_headers, self.database_headers, "ChCode", "Name"))

    def test_shared_recognition_preserves_datagrid_and_supports_iterators(self):
        self.assertTrue(is_filename_identified_pair(iter(self.campaign_headers), iter(self.database_headers)))
        self.assertTrue(is_filename_identified_pair(
            ["email", "chname", "accountnumber", "indexone_uid"],
            ["Name", "Old I.C.", "Account No.", "Email"]))
        self.assertFalse(is_filename_identified_pair(["CHCODE", "Name"], self.database_headers))

    def test_incomplete_schemas_cannot_bypass_stage_requirement(self):
        campaign = self.campaign(headers=self.campaign_headers[:-1], rows=[self.alice[:-1]])
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertIn("FLOW_STAGE header", result.error)
        database = self.database(headers=self.database_headers[:-1], rows=[
            ["00123", 9171234567, "Alice Smith", "Smith"],
        ])
        result = check_files(database, [self.campaign()], self.options)[0]
        self.assertIn("FLOW_STAGE header", result.error)

    def test_changing_either_lookup_key_keeps_stage_requirement(self):
        for file_key, database_key in [("ChName", "CHCODE"), ("ChCode", "Name")]:
            with self.subTest(file_key=file_key, database_key=database_key):
                result = check_files(replace(self.database(), column=database_key),
                                     [replace(self.campaign(), column=file_key)], self.options)[0]
                self.assertIn("FLOW_STAGE header", result.error)

    def test_present_stage_is_validated_and_remains_file_metadata(self):
        for rows, expected_error in [
            ([[*self.alice, None]], "FLOW_STAGE is blank"),
            ([[*self.alice, "FIELD"], [*self.alice, "NEW_ENDO"]], "multiple FLOW_STAGE"),
            ([[*self.alice, "FIELD"]], ""),
        ]:
            with self.subTest(expected_error=expected_error):
                campaign = self.campaign(rows, [*self.campaign_headers, "FLOW_STAGE"])
                result = check_files(self.database(), [campaign], self.options)[0]
                self.assertTrue(result.flow_stage_required)
                if expected_error:
                    self.assertIn(expected_error, result.error)
                else:
                    self.assertEqual(result.verdict, "EXACT MATCH")
                    self.assertEqual(result.file_identifier, "FIELD")
                    self.assertEqual(len(result.mappings), 5)
                    self.assertNotIn("FLOW_STAGE", result.ignored_columns)

    def test_full_copy_mode_keeps_stage_requirement(self):
        result = check_files(self.database(), [self.campaign()], Options(require_flow_stage=True))[0]
        self.assertEqual(result.verdict, "ERROR")
        self.assertIn("FLOW_STAGE header", result.error)

    def test_empty_mobile_blast_is_an_error(self):
        result = check_files(self.database(), [self.campaign(rows=[])], self.options)[0]
        self.assertEqual(result.verdict, "ERROR")
        self.assertIn("no data rows", result.error)

    def test_report_preserves_mobile_values_and_discloses_all_field_coverage(self):
        campaign = self.campaign(rows=[[*self.alice, "FIELD campaign"]],
                                 headers=[*self.campaign_headers, "Campaign"])
        database = self.database()
        results = check_files(database, [campaign], self.options)
        report = self.folder / "report.xlsx"
        export_report(report, database, self.options, results)
        book = load_workbook(report, read_only=True, data_only=True)
        try:
            self.assertEqual(book["Summary"]["A9"].value, "BLAST Field.xlsx")
            self.assertEqual(book["Summary"]["K9"].value, "MAPPED ROWS MATCH")
            self.assertEqual(book["Summary"]["U9"].value, "Campaign")
            self.assertIsNone(book["Summary"]["W9"].value)
            rows = list(book["Mappings 001"].values)[1:]
            self.assertEqual(rows[0][:4], ("formatted mobile", "Mobile phone", "Phone", "CHECKED"))
            self.assertIn("Philippine mobile", rows[0][4])
            self.assertEqual([row[3] for row in rows], ["CHECKED"] * 5 + ["NOT CHECKED"])
            self.assertEqual(book["File 001"]["D2"].value, self.alice[0])
            self.assertEqual(book["File 001"]["E2"].value, "00123")
        finally:
            book.close()


class PhilippineMobileComparisonTests(unittest.TestCase):
    def test_recognized_mobile_formats_are_equivalent_including_numeric_excel_cells(self):
        values = [
            "+639171234567", "639171234567", "09171234567", "9171234567",
            "+63 (917) 123-4567", "(0917) 123 4567", " 09 17-123-4567 ",
            639171234567, 639171234567.0, 9171234567, 9171234567.0,
        ]
        for value in values:
            with self.subTest(value=value):
                self.assertEqual(compare_value(value, "Phone", Options()), "+639171234567")

    def test_invalid_and_other_phones_retain_text_and_never_merge_multiple_numbers(self):
        values = [
            "0917ABC1234567", "09171234567 / 09181234567", "09171234567,09181234567",
            "09171234567 09181234567", "0917.123.4567", "09171234567 ext 2",
            "0917123456", "091712345678", "+649171234567", "0281234567",
        ]
        for value in values:
            with self.subTest(value=value):
                self.assertEqual(compare_value(value, "Phone", Options(ignore_case=False)), value)
                self.assertNotEqual(compare_value(value, "Phone", Options()), "+639171234567")
        self.assertNotEqual(compare_value("09171234567", "Phone", Options()),
                            compare_value("09171234568", "Phone", Options()))

    def test_invalid_phones_follow_case_and_space_options(self):
        self.assertEqual(compare_value(" Unknown Phone ", "Phone", Options()), "unknown phone")
        self.assertEqual(compare_value(" Unknown Phone ", "Phone", Options(False, False)), " Unknown Phone ")
        self.assertEqual(compare_value(None, "Phone", Options()), "")
        self.assertEqual(compare_value("", "Phone", Options()), "")


if __name__ == "__main__":
    unittest.main()
