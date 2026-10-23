import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from checker import ColumnMapping, Options
from web_checker import Upload, UploadedSource, clean_filename, configuration_signature, inspect_upload, run_uploaded_check


def csv_upload(name, rows):
    stream = io.StringIO(newline="")
    csv.writer(stream).writerows(rows)
    return Upload(name, stream.getvalue().encode("utf-8-sig"))


def xlsx_upload(name, rows):
    stream = io.BytesIO()
    book = Workbook()
    try:
        for row in rows:
            book.active.append(row)
        book.save(stream)
    finally:
        book.close()
    return Upload(name, stream.getvalue())


def database_upload():
    stream = io.BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.title = "Accounts"
    sheet.append(["CHCODE", "CHCode", "Name", "Last Name", "Email address", "OB", "PROD TYPE"])
    sheet.append(["00123", "00123", "Alice Smith", "Smith", "alice@example.test", 1234.5, "CARDS"])
    sheet.append(["456", "456", "Bob Jones", "Jones", "bob@example.test", 500, "LOAN"])
    book.create_sheet("Other").append(["Other"])
    book.save(stream)
    book.close()
    return Upload("database.xlsx", stream.getvalue())


def datagrid_upload():
    return xlsx_upload("DataGrid.xlsx", [
        ["DataGrid export"], [],
        ["Name", "Old I.C.", "Account No.", "Email"],
        ["Alice Smith", "00123", "001234567890", "alice@example.test"],
        ["Bob Jones", "456", "009876543210", "bob@example.test"],
    ])


def email_blast_upload(name="email_blast.csv", *, changed=False):
    return csv_upload(name, [
        ["email", "cc", "bcc", "indexone_uid", "primary_key", "chname", "accountnumber"],
        ["alice@example.test", "", "", "00123", "00123", "Alice Smith", "001234567890"],
        ["bob@example.test", "", "", "456", "456", "Changed Name" if changed else "Bob Jones", "009876543210"],
    ])


def cbs_apd_uploads(*, total_rows=2, changed_emails=0, changed_secondary_uid=False):
    database_rows = [
        ["DataGrid export"], [],
        ["Name", "Old I.C.", "Account No.", "Email", "Account Type"],
    ]
    blast_rows = [["email", "cc", "bcc", "indexone_uid", "primary_key", "chname", "accountnumber", "campaignbucket"]]
    for index in range(total_rows):
        key = f"{index + 1:05d}"
        name = f"Account {index + 1}"
        account = f"{index + 1:012d}"
        email = f"account{index + 1}@example.test"
        database_rows.append([name, key, account, email, "91-120"])
        blast_rows.append([
            f"changed{index + 1}@example.test" if index < changed_emails else email,
            "", "", "secondary-only" if changed_secondary_uid and index == 0 else key,
            key, name, account, "91-120",
        ])
    return xlsx_upload("DataGrid (1).xlsx", database_rows), csv_upload("CBS APD.csv", blast_rows)


def cbs_pl_uploads():
    return datagrid_upload(), csv_upload("CBS PL.csv", [
        ["email", "cc", "bcc", "indexone_uid", "primary_key", "chname", "accountnumber", "campaignbucket"],
        ["alice@example.test", "", "", "00123", "different-primary", "Alice Smith", "001234567890", "PL"],
        ["bob@example.test", "", "", "456", "456", "Bob Jones", "009876543210", "PL"],
    ])


class WebCheckerTests(unittest.TestCase):
    def database(self):
        return UploadedSource(database_upload(), "Accounts", "CHCODE")

    def test_workbook_upload_inspection_supports_sheets_and_duplicate_headers(self):
        info = inspect_upload(database_upload())
        self.assertEqual(info.sheets, ("Accounts", "Other"))
        self.assertEqual(info.column, "CHCODE")
        self.assertEqual(info.headers[:2], ("CHCODE", "CHCode [2]"))
        self.assertEqual(info.header_row, 1)
        self.assertEqual(inspect_upload(database_upload(), "Other").column, "")

    def test_datagrid_email_blast_inspection_and_download_preserve_key_and_filename(self):
        upload = datagrid_upload()
        database_info = inspect_upload(upload)
        blast = email_blast_upload()
        blast_info = inspect_upload(blast)
        self.assertEqual(database_info.column, "Old I.C.")
        self.assertEqual(database_info.header_row, 3)
        self.assertEqual(blast_info.column, "indexone_uid")
        self.assertEqual(blast_info.header_row, 1)
        database = UploadedSource(upload, database_info.sheet, database_info.column)
        source = UploadedSource(blast, blast_info.sheet, blast_info.column)
        run = run_uploaded_check(database, [source], Options(mapped_fields=True, require_flow_stage=True))
        result = run.results[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertFalse(result.flow_stage_required)
        self.assertEqual(result.file_identifier, "email_blast.csv")
        self.assertEqual(result.ignored_columns, ["cc", "bcc"])
        self.assertEqual(len(result.mappings), 5)
        self.assertEqual([record[:3] for record in result.records], [(2, "MATCHED", "4"), (3, "MATCHED", "5")])
        book = load_workbook(io.BytesIO(run.report))
        try:
            self.assertEqual(book["Summary"]["B9"].value, "email_blast.csv")
            self.assertEqual(book["File 001"]["G2"].value, "00123")
            self.assertEqual(book["File 001"]["G2"].data_type, "s")
        finally:
            book.close()

    def test_cbs_apd_uses_primary_key_and_checks_bucket_in_downloadable_report(self):
        database_upload, blast = cbs_apd_uploads()
        database_info = inspect_upload(database_upload)
        blast_info = inspect_upload(blast)
        self.assertEqual(database_info.column, "Old I.C.")
        self.assertEqual(blast_info.column, "primary_key")
        database = UploadedSource(database_upload, database_info.sheet, database_info.column)
        source = UploadedSource(blast, blast_info.sheet, blast_info.column)
        run = run_uploaded_check(database, [source], Options(mapped_fields=True, require_flow_stage=True))
        result = run.results[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertFalse(result.flow_stage_required)
        self.assertEqual(result.file_identifier, "CBS APD.csv")
        self.assertEqual(result.ignored_columns, ["cc", "bcc"])
        self.assertEqual(len(result.mappings), 6)
        self.assertIn(ColumnMapping("campaignbucket", "Account Type"), result.mappings)
        self.assertEqual([record[:3] for record in result.records], [(2, "MATCHED", "4"), (3, "MATCHED", "5")])
        book = load_workbook(io.BytesIO(run.report))
        try:
            self.assertEqual(book["Summary"]["B1"].value, "DataGrid (1).xlsx")
            self.assertEqual(book["Summary"]["B9"].value, "CBS APD.csv")
            self.assertEqual(book["File 001"]["H2"].value, "00001")
            self.assertEqual(book["File 001"]["H2"].data_type, "s")
            coverage = {row[0]: row[1:4] for row in book["Mappings 001"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(coverage["primary_key"], ("Old I.C.", "Text", "CHECKED"))
            self.assertEqual(coverage["campaignbucket"], ("Account Type", "Text", "CHECKED"))
            self.assertEqual(coverage["cc"][2], "NOT CHECKED")
            self.assertEqual(coverage["bcc"][2], "NOT CHECKED")
        finally:
            book.close()

    def test_cbs_apd_joins_primary_key_when_secondary_uid_differs(self):
        database_upload, blast = cbs_apd_uploads(changed_secondary_uid=True)
        database_info = inspect_upload(database_upload)
        blast_info = inspect_upload(blast)
        database = UploadedSource(database_upload, database_info.sheet, database_info.column)
        source = UploadedSource(blast, blast_info.sheet, blast_info.column)
        result = run_uploaded_check(database, [source], Options(mapped_fields=True, require_flow_stage=True)).results[0]
        self.assertEqual(result.source.column, "primary_key")
        self.assertFalse(result.error)
        self.assertEqual([record[:3] for record in result.records], [(2, "DIFFERENT", "4"), (3, "MATCHED", "5")])
        self.assertEqual(len(result.differences), 1)
        difference = result.differences[0]
        self.assertEqual(difference[:2], (2, 4))
        self.assertTrue(difference[2].startswith("indexone_uid"))
        self.assertEqual(difference[3:], ("secondary-only", "00001"))

    def test_cbs_pl_explicit_uid_key_joins_old_ic_and_checks_secondary_identifier(self):
        upload, blast = cbs_pl_uploads()
        database_info = inspect_upload(upload)
        blast_info = inspect_upload(blast)
        self.assertEqual(blast_info.column, "primary_key")
        database = UploadedSource(upload, database_info.sheet, database_info.column)
        source = UploadedSource(blast, blast_info.sheet, "indexone_uid")
        result = run_uploaded_check(database, [source], Options(mapped_fields=True, require_flow_stage=True)).results[0]
        self.assertFalse(result.error)
        self.assertEqual(result.source.column, "indexone_uid")
        self.assertEqual([record[:3] for record in result.records], [(2, "DIFFERENT", "4"), (3, "MATCHED", "5")])
        self.assertEqual(result.count("NOT FOUND"), 0)
        self.assertEqual(len(result.differences), 1)
        self.assertTrue(result.differences[0][2].startswith("primary_key"))
        self.assertEqual(result.differences[0][3:], ("different-primary", "00123"))

    def test_database_header_after_preamble_preserves_database_row_numbers(self):
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            with self.subTest(format=suffix):
                upload = make_upload(f"database.{suffix}", [
                    ["Database export"], [],
                    ["Name", " cHcOdE ", "Email address"],
                    ["Alice Smith", "00123", "alice@example.test"],
                    ["Bob Jones", "456", "bob@example.test"],
                ])
                info = inspect_upload(upload)
                self.assertEqual(info.header_row, 3)
                self.assertEqual(info.headers, ("Name", "cHcOdE", "Email address"))
                self.assertEqual(info.column, "cHcOdE")
                database = UploadedSource(upload, info.sheet, info.column)
                target = UploadedSource(csv_upload("campaign.csv", [
                    ["CHCODE", "FLOW_STAGE"], ["00123", "MAIL"], ["456", "MAIL"],
                ]), "CSV", "CHCODE")
                run = run_uploaded_check(database, [target], Options(mapped_fields=True, require_flow_stage=True))
                result = run.results[0]
                self.assertEqual(result.verdict, "EXACT MATCH")
                self.assertTrue(run.database.has_headers)
                self.assertEqual([record[:3] for record in result.records], [
                    (2, "MATCHED", "4"), (3, "MATCHED", "5"),
                ])
                book = load_workbook(io.BytesIO(run.report))
                try:
                    detail = book["File 001"]
                    self.assertEqual(detail["C2"].value, "4")
                    self.assertEqual(detail["C3"].value, "5")
                    self.assertEqual(detail["D2"].value, "00123")
                    self.assertEqual(detail["D2"].data_type, "s")
                finally:
                    book.close()

    def test_headered_upload_after_preamble_keeps_stage_checks_and_source_rows(self):
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            with self.subTest(format=suffix):
                upload = make_upload(f"campaign.{suffix}", [
                    ["Campaign export"], [],
                    ["Name", " chCODE ", "FLOW_STAGE"],
                    ["Alice Smith", "00123", "NO_REPLY"],
                    ["Changed Name", "456", "NO_REPLY"],
                ])
                info = inspect_upload(upload)
                self.assertEqual(info.header_row, 3)
                self.assertEqual(info.headers, ("Name", "chCODE", "FLOW_STAGE"))
                source = UploadedSource(upload, info.sheet, info.column)
                run = run_uploaded_check(self.database(), [source], Options(mapped_fields=True, require_flow_stage=True))
                result = run.results[0]
                self.assertFalse(result.error)
                self.assertTrue(result.flow_stage_required)
                self.assertEqual(result.flow_stage, "NO_REPLY")
                self.assertEqual([record[:3] for record in result.records], [
                    (4, "MATCHED", "2"), (5, "DIFFERENT", "3"),
                ])
                book = load_workbook(io.BytesIO(run.report))
                try:
                    detail = book["File 001"]
                    self.assertEqual(detail["A2"].value, 4)
                    self.assertEqual(detail["A3"].value, 5)
                    self.assertEqual(detail["E2"].value, "00123")
                    self.assertEqual(book["Differences 001"]["A2"].value, 5)
                    self.assertEqual(book["Differences 001"]["B2"].value, 3)
                finally:
                    book.close()

                missing_stage = make_upload(f"missing_stage.{suffix}", [
                    ["Campaign export"], [], ["Name", "CHCODE"], ["Alice Smith", "00123"],
                ])
                invalid = UploadedSource(missing_stage, info.sheet, "CHCODE")
                invalid_run = run_uploaded_check(self.database(), [invalid], Options(mapped_fields=True, require_flow_stage=True))
                self.assertEqual(invalid_run.results[0].verdict, "ERROR")
                self.assertIn("FLOW_STAGE", invalid_run.results[0].error)

    def test_headerless_column_b_upload_matches_a_database_with_preamble(self):
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            with self.subTest(format=suffix):
                database_upload = make_upload(f"database.{suffix}", [
                    ["Database export"], [], ["Name", "CHCODE"],
                    ["Alice Smith", "00123"], ["Bob Jones", "456"],
                ])
                database_info = inspect_upload(database_upload)
                database = UploadedSource(database_upload, database_info.sheet, database_info.column)
                upload = make_upload(f"headerless.{suffix}", [
                    ["First account", "00123"], ["Second account", "456"],
                ])
                info = inspect_upload(upload, has_headers=False)
                self.assertIsNone(info.header_row)
                source = UploadedSource(upload, info.sheet, info.column, has_headers=False)
                run = run_uploaded_check(database, [source], Options(mapped_fields=True, require_flow_stage=True))
                result = run.results[0]
                self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
                self.assertFalse(result.flow_stage_required)
                self.assertEqual([record[:3] for record in result.records], [
                    (1, "MATCHED", "4"), (2, "MATCHED", "5"),
                ])
                book = load_workbook(io.BytesIO(run.report))
                try:
                    detail = book["File 001"]
                    self.assertEqual(detail["A2"].value, 1)
                    self.assertEqual(detail["C2"].value, "4")
                    self.assertEqual(detail["E2"].value, "00123")
                    self.assertEqual(detail["E2"].data_type, "s")
                finally:
                    book.close()

    def test_browser_uploads_generate_a_downloadable_report(self):
        first = csv_upload("campaign.csv", [
            ["ChCode", "primary_key", "email", "chname", "chname_format", "OB_AT_ENDORSEMENT", "product_type", "endo"],
            ["00123", "00123", "alice@example.test", "Alice Smith", "Smith", "1,234.50", "Credit Card", "12OCT ENDO"],
        ])
        second = csv_upload("changed.csv", [["CHCODE", "Name"], ["456", "Changed"]])
        progress = []
        run = run_uploaded_check(self.database(), [UploadedSource(first, "CSV", "ChCode"), UploadedSource(second, "CSV", "CHCODE")],
                                 Options(mapped_fields=True), lambda index, total, result: progress.append((index, total)))
        self.assertEqual(run.results[0].verdict, "MAPPED ROWS MATCH")
        self.assertEqual(run.results[1].count("DIFFERENT"), 1)
        self.assertEqual(progress, [(1, 2), (2, 2)])
        self.assertTrue(run.report.startswith(b"PK"))
        self.assertEqual(run.database.path, Path("database.xlsx"))
        book = load_workbook(io.BytesIO(run.report))
        try:
            self.assertEqual(book["Summary"]["B1"].value, "database.xlsx")
            self.assertEqual(book["Summary"]["B9"].value, "campaign.csv")
            self.assertEqual(book["File 001"]["D2"].value, "00123")
            self.assertEqual(book["Mappings 001"]["D9"].value, "NOT CHECKED")
        finally:
            book.close()

    def test_uploads_with_same_filenames_are_kept_separate(self):
        database = UploadedSource(csv_upload("same.csv", [["CHCODE"], ["A"]]), "CSV", "CHCODE")
        first = UploadedSource(csv_upload("same.csv", [["CHCODE"], ["A"]]), "CSV", "CHCODE")
        second = UploadedSource(csv_upload("same.csv", [["CHCODE"], ["X"]]), "CSV", "CHCODE")
        run = run_uploaded_check(database, [first, second], Options(mapped_fields=True))
        self.assertEqual(run.results[0].count("MATCHED"), 1)
        self.assertEqual(run.results[1].count("NOT FOUND"), 1)

    def test_corrupt_target_does_not_stop_valid_uploads(self):
        invalid = UploadedSource(Upload("broken.xlsx", b"not an Excel workbook"), "Sheet1", "CHCODE")
        valid = UploadedSource(csv_upload("valid.csv", [["CHCODE"], ["00123"]]), "CSV", "CHCODE")
        run = run_uploaded_check(self.database(), [invalid, valid], Options(mapped_fields=True))
        self.assertEqual(run.results[0].verdict, "ERROR")
        self.assertEqual(run.results[1].count("MATCHED"), 1)

    def test_custom_field_mappings_survive_browser_adapter(self):
        upload = csv_upload("custom.csv", [["CHCODE", "Recipient"], ["00123", "alice@example.test"]])
        mappings = (ColumnMapping("CHCODE", "CHCODE"), ColumnMapping("Recipient", "Email address"))
        run = run_uploaded_check(self.database(), [UploadedSource(upload, "CSV", "CHCODE", mappings)], Options(mapped_fields=True))
        self.assertEqual(run.results[0].count("MATCHED"), 1)
        self.assertEqual(run.results[0].mappings, list(mappings))

    def test_headerless_uploads_preserve_first_row_keys_and_report_values(self):
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            with self.subTest(format=suffix):
                upload = make_upload(f"headerless.{suffix}", [
                    ["", "00123", "Original first row"],
                    ["Unknown account", "00999", "Keep this value"],
                    ["Blank account", "", "Flag this row"],
                ])
                info = inspect_upload(upload, has_headers=False)
                self.assertEqual(info.headers, ("Column A", "CHCODE", "Column C"))
                self.assertEqual(info.column, "CHCODE")
                source = UploadedSource(upload, info.sheet, info.column, has_headers=False)
                run = run_uploaded_check(self.database(), [source], Options(mapped_fields=True, require_flow_stage=True))
                result = run.results[0]
                self.assertFalse(result.error)
                self.assertFalse(result.source.has_headers)
                self.assertEqual([record[:3] for record in result.records], [
                    (1, "MATCHED", "2"), (2, "NOT FOUND", ""), (3, "BLANK KEY", ""),
                ])
                self.assertEqual(result.records[0][3][1:], ("00123", "Original first row"))
                self.assertEqual(result.records[1][3], ("Unknown account", "00999", "Keep this value"))
                self.assertEqual(result.mappings, [ColumnMapping("CHCODE", "CHCODE")])
                book = load_workbook(io.BytesIO(run.report))
                try:
                    detail = book["File 001"]
                    self.assertEqual(tuple(cell.value for cell in detail[1]),
                                     ("Source row", "Result", "Database row(s)", "Column A", "CHCODE", "Column C"))
                    self.assertEqual(detail["A2"].value, 1)
                    self.assertIsNone(detail["D2"].value)
                    self.assertEqual(detail["E2"].value, "00123")
                    self.assertEqual(detail["E2"].data_type, "s")
                    self.assertEqual(detail["F2"].value, "Original first row")
                    self.assertEqual(detail["A3"].value, 2)
                    self.assertEqual(detail["D3"].value, "Unknown account")
                    self.assertEqual(detail["E3"].value, "00999")
                    self.assertEqual(detail["A4"].value, 3)
                    self.assertEqual(detail["B4"].value, "BLANK KEY")
                    self.assertEqual(book["Summary"]["B9"].value, upload.name)
                finally:
                    book.close()

    def test_mixed_uploads_keep_flow_stage_validation_for_headered_files(self):
        headered = UploadedSource(csv_upload("campaign.csv", [
            ["CHCODE", "FLOW_STAGE"], ["00123", "MAIL"],
        ]), "CSV", "CHCODE")
        headerless = UploadedSource(csv_upload("without_headers.csv", [
            ["Other data", "456"],
        ]), "CSV", "CHCODE", has_headers=False)
        missing_stage = UploadedSource(csv_upload("missing_stage.csv", [
            ["CHCODE"], ["456"],
        ]), "CSV", "CHCODE")
        run = run_uploaded_check(self.database(), [headered, headerless, missing_stage],
                                 Options(mapped_fields=True, require_flow_stage=True))
        self.assertEqual(run.results[0].verdict, "EXACT MATCH")
        self.assertEqual(run.results[0].flow_stage, "MAIL")
        self.assertEqual(run.results[1].verdict, "MAPPED ROWS MATCH")
        self.assertFalse(run.results[1].flow_stage_required)
        self.assertEqual(run.results[1].source.path.name, "without_headers.csv")
        self.assertEqual(run.results[1].records[0][0], 1)
        self.assertEqual(run.results[2].verdict, "ERROR")
        self.assertIn("FLOW_STAGE", run.results[2].error)

    def test_headerless_uploads_require_column_b(self):
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            with self.subTest(format=suffix):
                upload = make_upload(f"one_column.{suffix}", [["00123"], ["456"]])
                with self.assertRaisesRegex(ValueError, "[Cc]olumn B"):
                    inspect_upload(upload, has_headers=False)
                sheet = "CSV" if suffix == "csv" else "Sheet"
                run = run_uploaded_check(self.database(), [UploadedSource(upload, sheet, "CHCODE", has_headers=False)],
                                         Options(mapped_fields=True))
                self.assertEqual(run.results[0].verdict, "ERROR")
                self.assertRegex(run.results[0].error, "[Cc]olumn B")

    def test_database_still_requires_headers(self):
        database = UploadedSource(csv_upload("database.csv", [["meta", "00123"]]), "CSV", "CHCODE", has_headers=False)
        target = UploadedSource(csv_upload("target.csv", [["CHCODE"], ["00123"]]), "CSV", "CHCODE")
        with self.assertRaisesRegex(ValueError, "[Dd]atabase.*[Hh]eader"):
            run_uploaded_check(database, [target], Options(mapped_fields=True))

    def test_input_or_option_changes_invalidate_signature(self):
        database = self.database()
        first = UploadedSource(csv_upload("file.csv", [["CHCODE"], ["00123"]]), "CSV", "CHCODE")
        second = UploadedSource(csv_upload("file.csv", [["CHCODE"], ["456"]]), "CSV", "CHCODE")
        signature = configuration_signature(database, [first], Options(mapped_fields=True))
        self.assertNotEqual(signature, configuration_signature(database, [second], Options(mapped_fields=True)))
        self.assertNotEqual(signature, configuration_signature(database, [first], Options(False, True, True)))
        changed_mapping = UploadedSource(first.upload, first.sheet, first.column, (ColumnMapping("CHCODE", "CHCODE"),))
        self.assertNotEqual(signature, configuration_signature(database, [changed_mapping], Options(mapped_fields=True)))
        changed_headers = UploadedSource(first.upload, first.sheet, first.column, has_headers=False)
        self.assertNotEqual(signature, configuration_signature(database, [changed_headers], Options(mapped_fields=True)))
        changed_file_key = UploadedSource(first.upload, first.sheet, "primary_key")
        self.assertNotEqual(signature, configuration_signature(database, [changed_file_key], Options(mapped_fields=True)))
        changed_database_key = UploadedSource(database.upload, database.sheet, "Old I.C.")
        self.assertNotEqual(signature, configuration_signature(changed_database_key, [first], Options(mapped_fields=True)))

    def test_temporary_uploads_are_removed_after_success_and_failure(self):
        original = tempfile.TemporaryDirectory
        directories = []

        def tracked(*args, **kwargs):
            directory = original(*args, **kwargs)
            directories.append(Path(directory.name))
            return directory

        valid = UploadedSource(csv_upload("valid.csv", [["CHCODE"], ["00123"]]), "CSV", "CHCODE")
        with patch("web_checker.tempfile.TemporaryDirectory", side_effect=tracked):
            run_uploaded_check(self.database(), [valid], Options(mapped_fields=True))
            invalid_database = UploadedSource(csv_upload("invalid.csv", [["Other"], ["A"]]), "CSV", "CHCODE")
            with self.assertRaises(ValueError):
                run_uploaded_check(invalid_database, [valid], Options(mapped_fields=True))
        self.assertEqual(len(directories), 2)
        self.assertTrue(all(not directory.exists() for directory in directories))

    def test_client_paths_cannot_escape_upload_directory(self):
        self.assertEqual(clean_filename("../../database.xlsx"), "database.xlsx")
        self.assertEqual(clean_filename(r"C:\Users\someone\file.csv"), "file.csv")
        with self.assertRaises(ValueError):
            clean_filename("..")

    def test_no_target_files_is_an_error(self):
        with self.assertRaisesRegex(ValueError, "at least one"):
            run_uploaded_check(self.database(), [], Options(mapped_fields=True))
