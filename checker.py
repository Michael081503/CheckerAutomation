"""File readers and matching engine for the local database checker."""

from __future__ import annotations

import csv
import re
from collections import defaultdict, deque
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from decimal import Decimal
from itertools import chain
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook, load_workbook

SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}
COMPARISON_VERSION = 4
CHCODE_ALIASES = ("chcode", "primary_key", "indexone_uid")
PHONE_ALIASES = ("mobilephone", "formattedmobile", "phonenumber", "mobilenumber",
                 "phone", "mobile", "contactnumber", "lastcalledno")
CUSTOMER_NAME_ALIASES = {"name", "fullname", "chname", "chnameformat", "firstname",
                         "middlename", "lastname", "surname", "givenname", "customername",
                         "customerfullname", "debtorname", "borrowername", "accountname", "casecustomername"}
HSBC_LAST_NAME_TITLE = re.compile(r"^(?:mr\.[ \t]*/[ \t]*ms\.|mr\.|ms\.)[ \t]+(?=\S)", re.IGNORECASE)


@dataclass
class Table:
    path: Path
    sheet: str
    headers: list[str]
    rows: list[tuple[int, tuple]]
    header_row: int | None = 1


@dataclass(frozen=True)
class ColumnMapping:
    file_column: str
    database_column: str
    comparison: str = "Text"


@dataclass(frozen=True)
class Source:
    path: Path
    sheet: str
    column: str
    mappings: tuple[ColumnMapping, ...] | None = None
    has_headers: bool = True
    detection_error: str = ""


@dataclass(frozen=True)
class Options:
    ignore_case: bool = True
    trim_spaces: bool = True
    mapped_fields: bool = False
    require_flow_stage: bool = False
    allow_missing_flow_stage: bool = False


@dataclass
class FileResult:
    source: Source
    headers: list[str]
    # Source row, status, matching database rows, original source values.
    records: list[tuple[int, str, str, tuple]]
    error: str = ""
    database_headers: list[str] = field(default_factory=list)
    missing_rows: list[tuple[int, tuple]] = field(default_factory=list)
    differences: list[tuple[int | None, int | None, str, object, object]] = field(default_factory=list)
    missing_columns: list[str] = field(default_factory=list)
    extra_columns: list[str] = field(default_factory=list)
    comparison_mode: str = "Full database copy"
    mappings: list[ColumnMapping] = field(default_factory=list)
    ignored_columns: list[str] = field(default_factory=list)
    flow_stage_required: bool = False
    flow_stage_column: str = ""
    flow_stage: str = ""
    database_match_column: str = ""

    @property
    def file_identifier(self) -> str:
        return self.flow_stage or self.source.path.name

    @property
    def file_identifier_column(self) -> str:
        return self.flow_stage_column or "Filename"

    def count(self, status: str) -> int:
        if status == "MISSING FROM FILE":
            return len(self.missing_rows)
        return sum(record[1] == status for record in self.records)

    @property
    def verdict(self) -> str:
        if self.error:
            return "ERROR"
        if self.comparison_mode == "Uploaded rows":
            if self.records and all(row[1] == "MATCHED" for row in self.records):
                return "EXACT MATCH" if self.flow_stage_required else "MAPPED ROWS MATCH"
            return "ROW ISSUES"
        if (self.missing_rows or self.missing_columns or self.extra_columns
                or any(record[1] != "MATCHED" for record in self.records)):
            return "DIFFERENT ROWS"
        return "SAME ROWS"


class Cancelled(Exception):
    pass


def _validate(path: Path) -> None:
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("Choose an .xlsx, .xlsm, or .csv file. Convert older .xls files to .xlsx first.")
    if not path.is_file():
        raise ValueError(f"File no longer exists: {path.name}")


def sheet_names(path: Path) -> list[str]:
    _validate(path)
    if path.suffix.lower() == ".csv":
        return ["CSV"]
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def _has_chcode(values: Iterable) -> bool:
    return any(str(value).strip().casefold() == "chcode" for value in values if value is not None)


def _header_names(values: Iterable) -> set[str]:
    return {str(value).strip().casefold() for value in values if value is not None}


def chcode_columns(headers: Iterable[str]) -> list[str]:
    """Find CHCODE aliases in priority order while preserving source names."""
    headers = list(headers)
    return [header for alias in CHCODE_ALIASES for header in headers
            if header.strip().casefold() == alias]


def phone_column(headers: Iterable[str]) -> str:
    """Find a named phone column, accepting spaces, underscores and letter case."""
    names = {re.sub(r"[^a-z0-9]", "", header.casefold()): header for header in headers}
    return next((names[alias] for alias in PHONE_ALIASES if alias in names), "")


def is_datagrid_database(headers: Iterable) -> bool:
    """Recognize the supported DataGrid schema, rather than an arbitrary key name."""
    names = _header_names(headers)
    return {"name", "old i.c.", "account no.", "email"}.issubset(names)


def is_salad_endo_database(headers: Iterable) -> bool:
    """Recognize SBF SALAD and regular PL endorsement source columns."""
    return {"caseno", "caseaccountno", "casecustomername", "casecustomeremail", "collector"}.issubset(
        _header_names(headers))


def is_datagrid_email_blast(headers: Iterable) -> bool:
    """Recognize email-blast fields that have DataGrid counterparts."""
    names = _header_names(headers)
    return ({"email", "chname", "accountnumber"}.issubset(names)
            and bool(set(CHCODE_ALIASES[1:]) & names))


def is_datagrid_viber_blast(headers: Iterable) -> bool:
    """Recognize SALAD Viber exports whose CHCode contains campaign labels."""
    return {"fullname", "firstname", "lastname", "phone", "chcode", "ob", "monthly",
            "terms", "promo", "level endo"}.issubset(_header_names(headers))


def is_datagrid_pl_viber_blast(headers: Iterable) -> bool:
    """Recognize PL Viber exports whose Debtorid identifies the customer."""
    return {"fullname", "firstname", "lastname", "phone", "chcode", "ob", "monthly",
            "terms", "promo", "level", "debtorid"}.issubset(_header_names(headers))


def is_cbs_apd_email_blast(headers: Iterable) -> bool:
    """Recognize the CBS APD layout whose primary_key contains the CHCODE."""
    names = _header_names(headers)
    return is_datagrid_email_blast(names) and {"primary_key", "campaignbucket"}.issubset(names)


def match_column(headers: Iterable[str], database_headers: Iterable[str] | None = None) -> str:
    """Select the lookup column, using the database schema for account blasts."""
    headers = list(headers)
    database_headers = list(database_headers) if database_headers is not None else None
    if is_datagrid_pl_viber_blast(headers):
        return next(header for header in headers if header.strip().casefold() == "debtorid")
    if is_datagrid_viber_blast(headers):
        return next(header for header in headers if header.strip().casefold() == "phone")
    for header in headers:
        if header.strip().casefold() == "chcode":
            return header
    if is_datagrid_database(headers):
        return next(header for header in headers if header.strip().casefold() == "old i.c.")
    if is_salad_endo_database(headers):
        return next(header for header in headers if header.strip().casefold() == "caseaccountno")
    if (database_headers is not None and not _has_chcode(database_headers)
            and is_salad_endo_database(database_headers)
            and {"email", "chname", "accountnumber"}.issubset(_header_names(headers))):
        return next(header for header in headers if header.strip().casefold() == "accountnumber")
    if is_datagrid_email_blast(headers):
        keys = ("primary_key", "indexone_uid") if is_cbs_apd_email_blast(headers) else ("indexone_uid", "primary_key")
        for expected in keys:
            for header in headers:
                if header.strip().casefold() == expected:
                    return header
    aliases = chcode_columns(headers)
    return aliases[0] if aliases else ""


def is_salad_endo_pair(file_headers: Iterable[str], database_headers: Iterable[str],
                       file_key: str | None = None, database_key: str | None = None) -> bool:
    file_headers, database_headers = list(file_headers), list(database_headers)
    file_key = match_column(file_headers, database_headers) if file_key is None else file_key
    database_key = match_column(database_headers) if database_key is None else database_key
    return (not _has_chcode(file_headers) and not _has_chcode(database_headers)
            and is_salad_endo_database(database_headers)
            and {"email", "chname", "accountnumber"}.issubset(_header_names(file_headers))
            and file_key.strip().casefold() == "accountnumber"
            and database_key.strip().casefold() == "caseaccountno")


def is_datagrid_pair(file_headers: Iterable[str], database_headers: Iterable[str],
                     file_key: str | None = None, database_key: str | None = None) -> bool:
    """Allow the stage-less layout only with its corresponding DataGrid database."""
    file_headers, database_headers = list(file_headers), list(database_headers)
    file_key = match_column(file_headers) if file_key is None else file_key
    database_key = match_column(database_headers) if database_key is None else database_key
    return (not _has_chcode(file_headers) and not _has_chcode(database_headers)
            and is_datagrid_email_blast(file_headers) and is_datagrid_database(database_headers)
            and file_key.strip().casefold() in CHCODE_ALIASES[1:]
            and database_key.strip().casefold() == "old i.c.")


def is_datagrid_viber_pair(file_headers: Iterable[str], database_headers: Iterable[str],
                           file_key: str | None = None, database_key: str | None = None) -> bool:
    file_headers, database_headers = list(file_headers), list(database_headers)
    file_key = match_column(file_headers) if file_key is None else file_key
    database_key = match_column(database_headers) if database_key is None else database_key
    return (is_datagrid_viber_blast(file_headers) and is_datagrid_database(database_headers)
            and "last called no." in _header_names(database_headers)
            and file_key.strip().casefold() == "phone"
            and database_key.strip().casefold() in {"old i.c.", "last called no."})


def is_datagrid_pl_viber_pair(file_headers: Iterable[str], database_headers: Iterable[str],
                              file_key: str | None = None, database_key: str | None = None) -> bool:
    file_headers, database_headers = list(file_headers), list(database_headers)
    file_key = match_column(file_headers) if file_key is None else file_key
    database_key = match_column(database_headers) if database_key is None else database_key
    return (is_datagrid_pl_viber_blast(file_headers) and is_datagrid_database(database_headers)
            and {"last called no.", "debtor id"}.issubset(_header_names(database_headers))
            and file_key.strip().casefold() == "debtorid"
            and database_key.strip().casefold() in {"old i.c.", "debtor id"})


def is_hsbc_mobile_pair(file_headers: Iterable[str], database_headers: Iterable[str],
                        file_key: str | None = None, database_key: str | None = None) -> bool:
    """Recognize the HSBC mobile blast and the database fields it validates."""
    file_headers, database_headers = list(file_headers), list(database_headers)
    file_key = match_column(file_headers) if file_key is None else file_key
    database_key = match_column(database_headers) if database_key is None else database_key
    return ({"formatted mobile", "chcode", "chname", "lastname", "producttype"}.issubset(
                _header_names(file_headers))
            and {"chcode", "mobile phone", "name", "last name", "prod type"}.issubset(
                _header_names(database_headers))
            and file_key.strip().casefold() == "chcode"
            and database_key.strip().casefold() == "chcode")


def is_filename_identified_pair(file_headers: Iterable[str], database_headers: Iterable[str],
                                file_key: str | None = None, database_key: str | None = None) -> bool:
    """Supported mapped layouts may use filenames when FLOW_STAGE is absent."""
    file_headers, database_headers = list(file_headers), list(database_headers)
    return (is_datagrid_pair(file_headers, database_headers, file_key, database_key)
            or is_datagrid_viber_pair(file_headers, database_headers, file_key, database_key)
            or is_datagrid_pl_viber_pair(file_headers, database_headers, file_key, database_key)
            or is_salad_endo_pair(file_headers, database_headers, file_key, database_key)
            or is_hsbc_mobile_pair(file_headers, database_headers, file_key, database_key))


def _has_supported_header(values: Iterable) -> bool:
    values = list(values)
    names = _header_names(values)
    identifier_layout = bool(names & set(CHCODE_ALIASES[1:])) and bool(
        names & {"flow_stage", "name", "email", "chname", "accountnumber"})
    return (_has_chcode(values) or bool(phone_column(names)) or is_datagrid_database(values)
            or is_datagrid_email_blast(values) or is_salad_endo_database(values)
            or {"email", "chname", "accountnumber"}.issubset(names) or identifier_layout)


def _table(path: Path, sheet: str, rows: Iterable, *, has_headers: bool = True,
           headers_only: bool = False) -> Table:
    iterator = iter(rows)
    first = next(iterator, None)
    if first is None:
        raise ValueError(f"{path.name}: the selected sheet is empty.")
    if not has_headers:
        records = []
        width = 0
        for number, values in enumerate(chain((first,), iterator), 1):
            values = tuple(values)
            if any(value is not None and str(value).strip() for value in values):
                width = max(width, len(values))
                records.append((number, values))
        if not records:
            raise ValueError(f"{path.name}: the selected sheet is empty.")
        if width < 2:
            raise ValueError(f"{path.name}: headerless files need at least two columns; CHCODE must be in column B.")
        from openpyxl.utils import get_column_letter
        headers = ["CHCODE" if index == 2 else f"Column {get_column_letter(index)}"
                   for index in range(1, width + 1)]
        return Table(path, sheet, headers,
                     [] if headers_only else
                     [(number, values + (None,) * (width - len(values))) for number, values in records],
                     header_row=None)
    header_row = 1
    if not _has_supported_header(first):
        preamble = []
        for number, values in enumerate(iterator, 2):
            if _has_supported_header(values):
                first = values
                header_row = number
                break
            preamble.append(values)
        else:
            # Keep the generic reader's missing-key behavior for unknown layouts.
            iterator = iter(preamble)
    # Keep internal empty headers so column positions never shift.
    headers = [str(value).strip() if value is not None else "" for value in first]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not header for header in headers):
        raise ValueError(f"{path.name}: header row {header_row} must contain a name for every data column.")
    # Real database exports repeat names. Preserve every position and give later
    # occurrences visible suffixes; the first CHCODE remains the primary key.
    used = set()
    for index, name in enumerate(headers):
        candidate = name
        occurrence = 2
        while candidate.casefold() in used:
            candidate = f"{name} [{occurrence}]"
            occurrence += 1
        headers[index] = candidate
        used.add(candidate.casefold())
    records = []
    if headers_only:
        return Table(path, sheet, headers, records, header_row=header_row)
    for number, values in enumerate(iterator, header_row + 1):
        values = tuple(values)
        if any(value is not None and str(value).strip() for value in values):
            if any(value is not None and str(value).strip() for value in values[len(headers):]):
                raise ValueError(f"{path.name}: row {number} contains data under an unnamed column.")
            records.append((number, values[:len(headers)] + (None,) * max(0, len(headers) - len(values))))
    return Table(path, sheet, headers, records, header_row=header_row)


def read_table(path: Path, sheet: str | None = None, *, headers_only: bool = False,
               has_headers: bool = True) -> Table:
    path = Path(path)
    _validate(path)
    if path.suffix.lower() == ".csv":
        try:
            contents = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            contents = path.read_text(encoding="cp1252")
        try:
            dialect = csv.Sniffer().sniff(contents[:65536], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        import io
        delimiter = dialect.delimiter
        if has_headers:
            # Preambles with different row widths can defeat Sniffer. Recover
            # the delimiter using a supported header, including later rows.
            candidates = [delimiter, *(value for value in ",;\t|" if value != delimiter)]
            for candidate in candidates:
                probe = csv.reader(io.StringIO(contents, newline=""), dialect, delimiter=candidate)
                if any(_has_supported_header(values) for values in probe):
                    delimiter = candidate
                    break
        reader = csv.reader(io.StringIO(contents, newline=""), dialect, delimiter=delimiter)
        return _table(path, "CSV", reader, has_headers=has_headers, headers_only=headers_only)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        selected = sheet or book.sheetnames[0]
        if selected not in book.sheetnames:
            raise ValueError(f"Sheet '{selected}' no longer exists in {path.name}.")
        reader = book[selected].iter_rows(values_only=True)
        return _table(path, selected, reader, has_headers=has_headers, headers_only=headers_only)
    finally:
        book.close()


def normalize(value, options: Options) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        result = str(value)
    elif isinstance(value, (int, float, Decimal)):
        result = format(Decimal(str(value)).normalize(), "f")
    elif isinstance(value, (datetime, date)):
        result = value.isoformat()
    else:
        result = str(value)
    if options.trim_spaces:
        result = result.strip()
    return result.casefold() if options.ignore_case else result


def suggest_mappings(file_headers: list[str], database_headers: list[str],
                     file_key: str, database_key: str) -> list[ColumnMapping]:
    """Direct header matches and explicit HSBC export aliases; never infer by position."""
    database_names = {name.casefold(): name for name in database_headers}
    if is_salad_endo_pair(file_headers, database_headers, file_key, database_key):
        # Campaign agent assignments and blank legacy IDs are not account keys.
        fields = {"accountnumber": "caseAccountNo", "chname": "caseCustomerName", "email": "caseCustomerEmail"}
        return [ColumnMapping(name, database_names[fields[name.casefold()].casefold()])
                for name in file_headers if name.casefold() in fields]
    if is_datagrid_pl_viber_pair(file_headers, database_headers, file_key, database_key):
        # PL campaign columns contain staff names/codes and template values.
        # Only these three fields have established database counterparts.
        fields = {"debtorid": ("Debtor ID", "Text"), "fullname": ("Name", "Text"),
                  "phone": ("Last Called No.", "Phone")}
        return [ColumnMapping(name, database_names[fields[name.casefold()][0].casefold()],
                              fields[name.casefold()][1])
                for name in file_headers if name.casefold() in fields]
    aliases = {
        "email": ("Email address", "Text"),
        "primary_key": (database_key, "Text"),
        "indexone_uid": (database_key, "Text"),
        "chname": ("Name", "Text"),
        "chname_format": ("Last Name", "Text"),
        "lastname": ("Last Name", "Text"),
        "formatted mobile": ("Mobile phone", "Phone"),
        "product_type": ("PROD TYPE", "Product"),
        "producttype": ("PROD TYPE", "Product"),
        "ob_at_endorsement": ("OB", "Number"),
    }
    datagrid_aliases = {}
    if is_datagrid_viber_pair(file_headers, database_headers, file_key, database_key):
        datagrid_aliases["fullname"] = ("Name", "Text")
    if is_datagrid_pair(file_headers, database_headers, file_key, database_key):
        datagrid_aliases = {
            "indexone_uid": (database_key, "Text"),
            "primary_key": (database_key, "Text"),
            "chname": ("Name", "Text"),
            "email": ("Email", "Text"),
            "accountnumber": ("Account No.", "Text"),
        }
        if is_cbs_apd_email_blast(file_headers) and "account type" in database_names:
            datagrid_aliases["campaignbucket"] = ("Account Type", "Text")
        if "financing/card no." in database_names:
            datagrid_aliases["financingnumber"] = ("Financing/Card No.", "Text")
    mappings = []
    for name in file_headers:
        if name == file_key:
            comparison = "Phone" if phone_column([name]) and phone_column([database_key]) else "Text"
            mappings.append(ColumnMapping(name, database_key, comparison))
        elif phone_column([name]):
            expected = database_names.get(name.casefold()) or phone_column(database_headers)
            if expected:
                mappings.append(ColumnMapping(name, expected, "Phone"))
        elif name.casefold() in datagrid_aliases:
            expected, comparison = datagrid_aliases[name.casefold()]
            mappings.append(ColumnMapping(name, database_names[expected.casefold()], comparison))
        elif name.casefold() in database_names:
            mappings.append(ColumnMapping(name, database_names[name.casefold()]))
        elif name.casefold() in aliases:
            expected, comparison = aliases[name.casefold()]
            if expected.casefold() in database_names:
                mappings.append(ColumnMapping(name, database_names[expected.casefold()], comparison))
    return mappings


def is_customer_name_column(header: str) -> bool:
    name = re.sub(r"\s+\[\d+\]$", "", header.strip())
    return re.sub(r"[^a-z0-9]", "", name.casefold()) in CUSTOMER_NAME_ALIASES


def compare_value(value, comparison: str, options: Options, *, customer_name: bool = False,
                  ignore_name_title: bool = False) -> str:
    if comparison == "Text" and customer_name and isinstance(value, str):
        value = value.replace("?", "")
        if ignore_name_title:
            value = HSBC_LAST_NAME_TITLE.sub("", value.strip() if options.trim_spaces else value, count=1)
    if comparison == "Phone":
        text = normalize(value, options)
        if re.fullmatch(r"[+0-9()\s-]+", text):
            compact = re.sub(r"[()\s-]", "", text)
            mobile = re.fullmatch(r"(?:\+?63|0)?(9[0-9]{9})", compact)
            if mobile:
                return "+63" + mobile.group(1)
        return text
    if comparison == "Number" and value is not None and str(value).strip():
        try:
            from decimal import InvalidOperation
            number = Decimal(str(value).strip().replace(",", ""))
            if number.is_finite():
                return format(number, "f").rstrip("0").rstrip(".") if "." in format(number, "f") else format(number, "f")
        except InvalidOperation:
            pass
    if comparison == "Product":
        product = str(value or "").strip().casefold()
        family = {"credit card": "cards", "cards": "cards", "personal loan": "loan", "loan": "loan"}
        if product in family:
            return family[product]
    return normalize(value, options)


def _name_text_matches(actual: str, expected: str) -> bool:
    """Compare normalized names, allowing a lost enye only at the same position.

    U+FFFD is already present in some exports, so decoding cannot recover its
    original letter. Never treat it as a wildcard for any other character.
    """
    if actual == expected:
        return True
    return len(actual) == len(expected) and all(
        left == right
        or (left == "\ufffd" and right in "\u00d1\u00f1")
        or (right == "\ufffd" and left in "\u00d1\u00f1")
        for left, right in zip(actual, expected)
    )


def _signature_difference_count(actual: tuple, expected: tuple, name_positions: set[int]) -> int:
    return sum(
        not (_name_text_matches(left, right) if index in name_positions else left == right)
        for index, (left, right) in enumerate(zip(actual, expected))
    )


def find_header(headers: Iterable[str], name: str) -> str:
    return next((header for header in headers if header.casefold() == name.strip().casefold()), "")


def _identify_flow_stage(target: Table, result: FileResult) -> None:
    result.flow_stage_column = find_header(target.headers, "FLOW_STAGE")
    if not result.flow_stage_column:
        raise ValueError("The uploaded file must have a FLOW_STAGE header to identify the file type.")
    position = target.headers.index(result.flow_stage_column)
    stages = {}
    blank_rows = []
    for number, values in target.rows:
        value = "" if values[position] is None else str(values[position]).strip()
        if not value:
            blank_rows.append(number)
        else:
            stages.setdefault(value.casefold(), value.upper())
    if blank_rows:
        raise ValueError(f"FLOW_STAGE is blank at source row(s) {', '.join(map(str, blank_rows[:20]))}. Every row needs the file identifier.")
    if len(stages) > 1:
        raise ValueError(f"The file contains multiple FLOW_STAGE identifiers ({', '.join(stages.values())}). Upload a separate file for each stage.")
    if not stages:
        raise ValueError("The file has no data rows with a FLOW_STAGE identifier.")
    result.flow_stage = next(iter(stages.values()))


def _compare_mapped(table: Table, target: Table, database: Source, source: Source,
                    result: FileResult, options: Options, lookup: dict,
                    cancelled: Callable[[], bool]) -> None:
    result.comparison_mode = "Uploaded rows"
    if source.mappings is not None:
        mappings = list(source.mappings)
    elif source.has_headers:
        mappings = suggest_mappings(target.headers, table.headers, source.column, database.column)
    else:
        mappings = [ColumnMapping(source.column, database.column)]
    if result.flow_stage_required:
        if source.mappings is not None and any(item.file_column == result.flow_stage_column for item in mappings):
            raise ValueError("FLOW_STAGE is the file identifier and must not be mapped to a database column.")
        mappings = [item for item in mappings if item.file_column != result.flow_stage_column]
    if not any(item.file_column == source.column and item.database_column == database.column for item in mappings):
        raise ValueError(f"The {source.column}-to-{database.column} mapping is required.")
    if len({item.file_column for item in mappings}) != len(mappings):
        raise ValueError("Each file column can have only one database mapping.")
    for item in mappings:
        if item.file_column not in target.headers or item.database_column not in table.headers:
            raise ValueError(f"Mapping no longer exists: {item.file_column} → {item.database_column}. Review field mappings.")
        if item.comparison not in {"Text", "Number", "Product", "Phone"}:
            raise ValueError(f"Unknown comparison type: {item.comparison}")
    result.mappings = mappings
    result.ignored_columns = [name for name in target.headers if name not in {item.file_column for item in mappings}
                              and name != result.flow_stage_column]
    if not target.rows:
        raise ValueError("The selected file has column headers but no data rows to validate.")
    positions = [(target.headers.index(item.file_column), table.headers.index(item.database_column), item) for item in mappings]
    customer_name_fields = {item.file_column for item in mappings
                            if is_customer_name_column(item.file_column) or is_customer_name_column(item.database_column)}
    last_name_fields = ({item.file_column for item in mappings if item.comparison == "Text"
                         and re.sub(r"[^a-z0-9]", "", item.database_column.casefold()) == "lastname"
                         and item.file_column != source.column}
                        if is_hsbc_mobile_pair(target.headers, table.headers, source.column, database.column) else set())
    name_positions = {index for index, (_, _, item) in enumerate(positions)
                      if item.comparison == "Text" and item.file_column in customer_name_fields
                      and item.file_column != source.column}
    database_rows = dict(table.rows)
    # An uploaded row is validated independently; database rows may be reused by
    # repeated uploaded CHCODEs and by separate campaign files.
    key_position = target.headers.index(source.column)
    key_mapping = next(item for item in mappings if item.file_column == source.column
                       and item.database_column == database.column)
    required_phone = (phone_column(target.headers) if (
        options.allow_missing_flow_stage and source.has_headers and not find_header(target.headers, "FLOW_STAGE")
        or is_datagrid_pl_viber_pair(target.headers, table.headers, source.column, database.column)) else "")
    phone_position = target.headers.index(required_phone) if required_phone else None
    database_signatures = {}
    exact = {}
    for key, numbers in lookup.items():
        if cancelled():
            raise Cancelled()
        for number in numbers:
            values = database_rows[number]
            signature = tuple(compare_value(values[db_pos], item.comparison, options,
                                             customer_name=item.file_column in customer_name_fields,
                                             ignore_name_title=item.file_column in last_name_fields)
                              for _, db_pos, item in positions)
            database_signatures[number] = signature
            exact.setdefault((key, signature), number)
    for row_number, values in target.rows:
        if cancelled():
            raise Cancelled()
        key = compare_value(values[key_position], key_mapping.comparison, options,
                            customer_name=key_mapping.file_column in customer_name_fields)
        signature = tuple(compare_value(values[file_pos], item.comparison, options,
                                         customer_name=item.file_column in customer_name_fields,
                                         ignore_name_title=item.file_column in last_name_fields)
                          for file_pos, _, item in positions)
        matched = exact.get((key, signature))
        blank_phone = phone_position is not None and not normalize(values[phone_position], options)
        if matched is None and not blank_phone and key in lookup:
            # Only compare names within the same identifier's database rows.
            matched = next((number for number in lookup[key] if not _signature_difference_count(
                signature, database_signatures[number], name_positions)), None)
        if blank_phone:
            matched = None
        if not key:
            status = "BLANK KEY"
        elif key not in lookup:
            status = "NOT FOUND"
        elif matched is not None:
            status = "MATCHED"
        else:
            status = "DIFFERENT"
            # Show the candidate with the fewest differing mapped values.
            matched = min(lookup[key], key=lambda number: _signature_difference_count(
                signature, database_signatures[number], name_positions))
            expected = database_rows[matched]
            for file_pos, db_pos, item in positions:
                actual, wanted = values[file_pos], expected[db_pos]
                name_field = item.file_column in customer_name_fields
                actual_text = compare_value(actual, item.comparison, options, customer_name=name_field,
                                            ignore_name_title=item.file_column in last_name_fields)
                wanted_text = compare_value(wanted, item.comparison, options, customer_name=name_field,
                                            ignore_name_title=item.file_column in last_name_fields)
                name_match = (item.comparison == "Text" and name_field and item.file_column != source.column)
                if not (_name_text_matches(actual_text, wanted_text) if name_match else actual_text == wanted_text):
                    label = item.file_column if item.file_column == item.database_column else f"{item.file_column} → {item.database_column}"
                    result.differences.append((row_number, matched, label, actual, wanted))
                elif blank_phone and item.file_column == required_phone:
                    label = item.file_column if item.file_column == item.database_column else f"{item.file_column} → {item.database_column}"
                    result.differences.append((row_number, matched, label, actual, "(phone number required)"))
        result.records.append((row_number, status, str(matched) if matched is not None else "", values))


def check_files(database: Source, files: list[Source], options: Options,
                progress: Callable[[int, int, FileResult], None] | None = None,
                cancelled: Callable[[], bool] = lambda: False) -> list[FileResult]:
    if not database.has_headers:
        raise ValueError("The database must have a header row containing CHCODE or the supported DataGrid fields.")
    table = read_table(database.path, database.sheet)
    if database.column not in table.headers:
        raise ValueError(f"Database column '{database.column}' no longer exists.")
    comparison_headers = [name for name in table.headers if not (options.require_flow_stage and name.casefold() == "flow_stage")]
    customer_name_columns = {name for name in comparison_headers if is_customer_name_column(name)}
    comparison_positions = [table.headers.index(name) for name in comparison_headers]
    position = table.headers.index(database.column)
    lookup: dict[str, list[int]] = {}
    database_rows = dict(table.rows)
    signatures = {}
    blank_rows = []
    for row_number, values in table.rows:
        if cancelled():
            raise Cancelled()
        key = compare_value(values[position], "Text", options, customer_name=database.column in customer_name_columns)
        if key:
            lookup.setdefault(key, []).append(row_number)
            if not options.mapped_fields:
                signatures[row_number] = tuple(compare_value(values[index], "Text", options,
                                                              customer_name=table.headers[index] in customer_name_columns)
                                               for index in comparison_positions)
        else:
            blank_rows.append(row_number)
    if not lookup:
        raise ValueError("The database match column has no nonblank values.")
    if blank_rows:
        key_name = "CHCODEs" if database.column.casefold() == "chcode" else f"values in '{database.column}'"
        raise ValueError(f"The database has blank {key_name} at row(s) {', '.join(map(str, blank_rows[:20]))}. Fill these before comparing rows.")
    results = []
    viber_lookup = None
    pl_viber_lookup = None
    for index, source in enumerate(files, 1):
        if cancelled():
            raise Cancelled()
        result = FileResult(source, [], [], database_headers=table.headers,
                            comparison_mode="Uploaded rows" if options.mapped_fields else "Full database copy",
                            flow_stage_required=options.require_flow_stage and source.has_headers,
                            database_match_column=database.column)
        try:
            if source.detection_error:
                raise ValueError(source.detection_error)
            if not source.has_headers and not options.mapped_fields:
                raise ValueError("Headerless files support uploaded-row checks only. Choose Validate uploaded rows.")
            target = read_table(source.path, source.sheet, has_headers=source.has_headers)
            result.headers = target.headers
            if source.column not in target.headers:
                raise ValueError(f"Match column '{source.column}' no longer exists.")
            comparison_database, comparison_lookup = database, lookup
            if (options.mapped_fields and source.has_headers and is_salad_endo_database(table.headers)
                    and not _has_chcode(table.headers) and not _has_chcode(target.headers)
                    and not is_salad_endo_pair(target.headers, table.headers, source.column, database.column)):
                raise ValueError("SBF endorsement email checks require accountnumber, chname, and email, using accountnumber to match caseAccountNo.")
            if (options.mapped_fields and is_datagrid_pl_viber_blast(target.headers)
                    and source.column.casefold() == "debtorid"):
                if not is_datagrid_pl_viber_pair(target.headers, table.headers, source.column, database.column):
                    raise ValueError("PL Viber checks require a DataGrid database with Name, Old I.C., Account No., Email, Last Called No., and Debtor ID columns, using Old I.C. or Debtor ID as its database key.")
                debtor_column = find_header(table.headers, "Debtor ID")
                comparison_database = replace(database, column=debtor_column)
                result.database_match_column = debtor_column
                if pl_viber_lookup is None:
                    candidate_lookup = {}
                    debtor_position = table.headers.index(debtor_column)
                    for number, values in table.rows:
                        if cancelled():
                            raise Cancelled()
                        debtor = normalize(values[debtor_position], options)
                        if not debtor:
                            raise ValueError(f"The database Debtor ID column has a blank value at row {number}.")
                        candidate_lookup.setdefault(debtor, []).append(number)
                    pl_viber_lookup = candidate_lookup
                comparison_lookup = pl_viber_lookup
            if (options.mapped_fields and is_datagrid_viber_blast(target.headers)
                    and source.column.casefold() == "phone"):
                if not is_datagrid_viber_pair(target.headers, table.headers, source.column, database.column):
                    raise ValueError("SALAD Viber checks require a DataGrid database with Name, Old I.C., Account No., Email, and Last Called No. columns, using Old I.C. or Last Called No. as its database key.")
                database_phone = find_header(table.headers, "Last Called No.")
                comparison_database = replace(database, column=database_phone)
                result.database_match_column = database_phone
                if viber_lookup is None:
                    viber_lookup = {}
                    phone_position = table.headers.index(database_phone)
                    for number, values in table.rows:
                        if cancelled():
                            raise Cancelled()
                        phone = compare_value(values[phone_position], "Phone", options)
                        if phone:
                            viber_lookup.setdefault(phone, []).append(number)
                if not viber_lookup:
                    raise ValueError("The database Last Called No. column has no nonblank phone numbers.")
                comparison_lookup = viber_lookup
            if (options.mapped_fields and options.allow_missing_flow_stage and source.has_headers
                    and not find_header(target.headers, "FLOW_STAGE") and phone_column(target.headers)
                    and not phone_column(table.headers)):
                raise ValueError("The file has no FLOW_STAGE and contains phone numbers, but the database has no supported phone column.")
            if (options.mapped_fields and result.flow_stage_required
                    and not find_header(target.headers, "FLOW_STAGE")
                    and (options.allow_missing_flow_stage or is_filename_identified_pair(
                        target.headers, table.headers, source.column, database.column))):
                result.flow_stage_required = False
            if result.flow_stage_required:
                _identify_flow_stage(target, result)
            if options.mapped_fields:
                _compare_mapped(table, target, comparison_database, source, result, options, comparison_lookup, cancelled)
                results.append(result)
                if progress:
                    progress(index, len(files), result)
                continue
            target_position = target.headers.index(source.column)
            target_columns = {name.casefold(): i for i, name in enumerate(target.headers) if name != result.flow_stage_column}
            database_columns = {name.casefold(): i for i, name in enumerate(table.headers)}
            result.missing_columns = [name for name in comparison_headers if name.casefold() not in target_columns]
            result.extra_columns = [name for name in target.headers if name.casefold() not in database_columns and name != result.flow_stage_column]
            schema_matches = not result.missing_columns and not result.extra_columns
            # Reserve identical rows before pairing modified rows with the same CHCODE.
            # This keeps duplicate CHCODEs independent of source row order.
            available = defaultdict(deque)
            for number, signature in signatures.items():
                if cancelled():
                    raise Cancelled()
                available[signature].append(number)
            consumed = set()
            exact = {}
            if schema_matches:
                for number, values in target.rows:
                    if cancelled():
                        raise Cancelled()
                    signature = tuple(compare_value(values[target_columns[name.casefold()]], "Text", options,
                                                      customer_name=name in customer_name_columns)
                                      for name in comparison_headers)
                    candidates = available.get(signature)
                    if candidates:
                        matched = candidates.popleft()
                        exact[number] = matched
                        consumed.add(matched)
                # Reserve compatible damaged-name rows after exact rows, before
                # pairing changed rows. Duplicate accounts still consume one row.
                name_positions = {index for index, name in enumerate(comparison_headers)
                                  if name in customer_name_columns and name != database.column}
                for number, values in target.rows:
                    if cancelled():
                        raise Cancelled()
                    if number in exact:
                        continue
                    key = compare_value(values[target_position], "Text", options,
                                        customer_name=database.column in customer_name_columns)
                    signature = tuple(compare_value(values[target_columns[name.casefold()]], "Text", options,
                                                      customer_name=name in customer_name_columns)
                                      for name in comparison_headers)
                    matched = next((candidate for candidate in lookup.get(key, [])
                                    if candidate not in consumed and not _signature_difference_count(
                                        signature, signatures[candidate], name_positions)), None)
                    if matched is not None:
                        exact[number] = matched
                        consumed.add(matched)
            remaining = {key: deque(number for number in numbers if number not in consumed)
                         for key, numbers in lookup.items()}
            for row_number, values in target.rows:
                if cancelled():
                    raise Cancelled()
                key = compare_value(values[target_position], "Text", options,
                                    customer_name=database.column in customer_name_columns)
                matched = exact.get(row_number)
                if matched is not None:
                    status = "MATCHED"
                elif not key:
                    status = "BLANK KEY"
                elif key not in lookup:
                    status = "NOT FOUND"
                elif remaining[key]:
                    matched = remaining[key].popleft()
                    consumed.add(matched)
                    status = "DIFFERENT"
                    expected = database_rows[matched]
                    for name in comparison_headers:
                        name_key = name.casefold()
                        expected_value = expected[database_columns[name_key]]
                        if name_key not in target_columns:
                            result.differences.append((row_number, matched, name, "(missing column)", expected_value))
                        else:
                            actual = values[target_columns[name_key]]
                            name_field = name in customer_name_columns
                            actual_text = compare_value(actual, "Text", options, customer_name=name_field)
                            wanted_text = compare_value(expected_value, "Text", options, customer_name=name_field)
                            name_match = name_field and name != database.column
                            if not (_name_text_matches(actual_text, wanted_text) if name_match else actual_text == wanted_text):
                                result.differences.append((row_number, matched, name, actual, expected_value))
                    for name in result.extra_columns:
                        result.differences.append((row_number, matched, name, values[target_columns[name.casefold()]], "(no database column)"))
                else:
                    status = "EXTRA ROW"
                result.records.append((row_number, status, str(matched) if matched is not None else "", values))
            result.missing_rows = [(number, values) for number, values in table.rows if number not in consumed]
        except Cancelled:
            raise
        except Exception as error:
            result.error = str(error)
        results.append(result)
        if progress:
            progress(index, len(files), result)
    return results


def export_report(path: Path, database: Source, options: Options, results: list[FileResult]) -> None:
    from openpyxl.cell import WriteOnlyCell

    book = Workbook(write_only=True)

    def append(sheet, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)

    summary = book.create_sheet("Summary")
    append(summary, ["Database", str(database.path)])
    append(summary, ["Database sheet", database.sheet])
    append(summary, ["Database match column", database.column])
    append(summary, ["Ignore letter case", str(options.ignore_case)])
    append(summary, ["Trim outer spaces", str(options.trim_spaces)])
    append(summary, ["Created", datetime.now().isoformat(timespec="seconds")])
    append(summary, [])
    append(summary, ["File", "Full path", "Sheet", "Match column", "Rows", "Matched rows", "Not found", "Blank key", "Error", "Report sheet",
                     "Comparison", "Different rows", "Extra rows", "Missing database rows", "Missing columns", "Extra columns", "Differences sheet", "Missing rows sheet",
                     "Scope", "Fields checked", "Unchecked file columns", "Mappings sheet",
                     "FLOW_STAGE", "File identifier column", "Database match column"])
    for index, result in enumerate(results, 1):
        name = f"File {index:03d}"
        append(summary, [result.source.path.name, str(result.source.path), result.source.sheet,
                         result.source.column, len(result.records), result.count("MATCHED"),
                         result.count("NOT FOUND"), result.count("BLANK KEY"), result.error, name,
                         result.verdict, result.count("DIFFERENT"), result.count("EXTRA ROW"), len(result.missing_rows),
                         ", ".join(result.missing_columns), ", ".join(result.extra_columns), f"Differences {index:03d}",
                         f"Missing {index:03d}" if result.comparison_mode != "Uploaded rows" else "",
                         result.comparison_mode,
                         "; ".join(f"{item.file_column} → {item.database_column} ({item.comparison})" for item in result.mappings),
                         ", ".join(result.ignored_columns), f"Mappings {index:03d}" if result.comparison_mode == "Uploaded rows" else "",
                         result.flow_stage, result.flow_stage_column, result.database_match_column or database.column])
        detail = book.create_sheet(name)
        detail.freeze_panes = "D2"
        append(detail, ["Source row", "Result", "Database row(s)", *result.headers])
        for number, status, matches, values in result.records:
            append(detail, [number, status, matches, *values])
        if result.error:
            append(detail, ["ERROR", result.error])
        differences = book.create_sheet(f"Differences {index:03d}")
        differences.freeze_panes = "A2"
        append(differences, ["Source row", "Database row", "Column", "File value", "Database value"])
        for difference in result.differences:
            append(differences, difference)
        if result.comparison_mode == "Uploaded rows":
            mapping_sheet = book.create_sheet(f"Mappings {index:03d}")
            mapping_sheet.freeze_panes = "A2"
            append(mapping_sheet, ["File column", "Database column", "Comparison", "Coverage", "Rule"])
            mapped = {item.file_column: item for item in result.mappings}
            for header in result.headers:
                item = mapped.get(header)
                if header == result.flow_stage_column:
                    append(mapping_sheet, [header, "", "", "FILE IDENTIFIER", "Labels the file; row data is compared against database records by CHCODE"])
                elif item is None:
                    rule = "No database mapping selected"
                    if is_salad_endo_pair(result.headers, result.database_headers,
                                          result.source.column, result.database_match_column):
                        if header.casefold() == "agentcode":
                            rule = "Staff assignments are not compared with Collector; only account number, name and email are checked"
                        elif header.casefold() in CHCODE_ALIASES[1:]:
                            rule = "Not used for this layout; accountnumber identifies the caseAccountNo record"
                    append(mapping_sheet, [header, "", "", "NOT CHECKED", rule])
                else:
                    rule = {"Text": "Uses case and space options", "Number": "Numeric equality; commas ignored",
                            "Product": "Credit Card = CARDS; Personal Loan = LOAN",
                            "Phone": "Matches the phone number for the same CHCODE; Philippine mobile formats +639 / 639 / 09 / 9 are equivalent; other values use text comparison"}[item.comparison]
                    if is_salad_endo_pair(result.headers, result.database_headers,
                                          result.source.column, result.database_match_column):
                        rule = ("Finds database rows by accountnumber = caseAccountNo; all mapped fields must match the same row"
                                if header == result.source.column else
                                "Compares the value on the same caseAccountNo record; uses case and space options")
                    if item.comparison == "Phone" and item.file_column == result.source.column:
                        rule = "Finds database rows by phone; all mapped fields must match the same row. Philippine mobile formats +639 / 639 / 09 / 9 are equivalent; other values use text comparison"
                    if is_datagrid_pl_viber_pair(result.headers, result.database_headers,
                                                 result.source.column, result.database_match_column):
                        if item.file_column == result.source.column:
                            rule = "Finds database rows by Debtor ID; all mapped fields must match the same row"
                        elif item.comparison == "Phone":
                            rule = "Matches the phone number for the same Debtor ID; a nonblank phone is required. Philippine mobile formats +639 / 639 / 09 / 9 are equivalent; other values use text comparison"
                        else:
                            rule = "Compares the value on the same database row as Debtor ID; uses case and space options"
                    if item.comparison == "Text" and (is_customer_name_column(item.file_column)
                                                      or is_customer_name_column(item.database_column)):
                        rule += "; question marks are ignored in customer names"
                        rule += "; replacement character U+FFFD matches only a corresponding Ñ/ñ; other letters must match"
                    if (item.comparison == "Text" and item.file_column != result.source.column
                            and re.sub(r"[^a-z0-9]", "", item.database_column.casefold()) == "lastname"
                            and is_hsbc_mobile_pair(result.headers, result.database_headers,
                                                    result.source.column, result.database_match_column)):
                        rule += "; leading Mr./Ms., Mr. and Ms. titles are ignored; surname and suffixes remain checked on the same CHCODE"
                    append(mapping_sheet, [header, item.database_column, item.comparison, "CHECKED", rule])
        else:
            missing = book.create_sheet(f"Missing {index:03d}")
            missing.freeze_panes = "B2"
            append(missing, ["Database row", *result.database_headers])
            for number, values in result.missing_rows:
                append(missing, [number, *values])
    summary.freeze_panes = "A9"
    # Save beside the destination first, then replace atomically.
    import os
    import tempfile
    destination = Path(path)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".xlsx", delete=False) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        os.replace(temporary, destination)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
