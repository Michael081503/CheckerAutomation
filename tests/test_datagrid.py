import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from checker import (
    ColumnMapping, Options, Source, check_files, export_report,
    is_datagrid_database, is_datagrid_email_blast, is_datagrid_pair,
    match_column, read_table, suggest_mappings,
)


class DataGridTests(unittest.TestCase):
    database_headers = ["Name", "Old I.C.", "Account No.", "Email"]
    campaign_headers = ["email", "cc", "bcc", "indexone_uid", "primary_key", "chname", "accountnumber"]

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.options = Options(mapped_fields=True, require_flow_stage=True)

    def csv(self, name, rows, delimiter=","):
        path = self.folder / name
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            csv.writer(stream, delimiter=delimiter).writerows(rows)
        return path

    def excel(self, name, rows):
        path = self.folder / name
        book = Workbook()
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        return path

    def source(self, path):
        table = read_table(path)
        return Source(path, table.sheet, match_column(table.headers))

    def database(self, rows=None):
        return self.source(self.excel("DataGrid.xlsx", [self.database_headers, *(
            rows if rows is not None else [
                ["Alice Smith", "00123", "0005678", "alice@example.test"],
                ["Bob Jones", "00456", "0098765", "bob@example.test"],
            ]
        )]))

    def campaign(self, rows=None, headers=None, name="email-blast.csv"):
        return self.source(self.csv(name, [headers or self.campaign_headers, *(
            rows if rows is not None else [
                ["alice@example.test", "", "", "00123", "00123", "Alice Smith", "0005678"],
            ]
        )]))

    def test_supported_schemas_choose_keys_and_preserve_chcode_precedence(self):
        self.assertEqual(match_column(self.database_headers), "Old I.C.")
        self.assertEqual(match_column(self.campaign_headers), "indexone_uid")
        self.assertEqual(match_column([*self.campaign_headers, "ChCode"]), "ChCode")
        self.assertEqual(match_column([*self.database_headers, "CHCODE"]), "CHCODE")
        self.assertEqual(match_column(["primary_key", "Other"]), "")
        self.assertFalse(is_datagrid_database(["Old I.C."]))
        self.assertFalse(is_datagrid_email_blast(["indexone_uid", "email"]))
        self.assertTrue(is_datagrid_pair(self.campaign_headers, self.database_headers))
        self.assertFalse(is_datagrid_pair(self.campaign_headers, self.database_headers, "email", "Old I.C."))
        self.assertFalse(is_datagrid_pair([*self.campaign_headers, "CHCODE"], self.database_headers,
                                          "indexone_uid", "Old I.C."))
        self.assertFalse(is_datagrid_pair(self.campaign_headers, [*self.database_headers, "CHCODE"],
                                          "indexone_uid", "Old I.C."))

    def test_supported_headers_are_found_below_preambles_in_csv_and_excel(self):
        for name, headers, values in [
            ("database", self.database_headers, ["Alice Smith", "00123", "0005678", "alice@example.test"]),
            ("campaign", self.campaign_headers, ["alice@example.test", "", "", "00123", "00123", "Alice Smith", "0005678"]),
        ]:
            rows = [["Export report"], [], headers, values]
            for path in [self.csv(f"{name}.csv", rows, delimiter=";"), self.excel(f"{name}.xlsx", rows)]:
                with self.subTest(file=path.name):
                    table = read_table(path)
                    self.assertEqual(table.header_row, 3)
                    self.assertEqual(table.headers, headers)
                    expected = tuple(None if value == "" and path.suffix == ".xlsx" else value for value in values)
                    self.assertEqual(table.rows, [(4, expected)])
                    preview = read_table(path, headers_only=True)
                    self.assertEqual(preview.headers, headers)
                    self.assertEqual(preview.rows, [])

    def test_actual_blast_schema_validates_five_fields_and_discloses_recipients(self):
        campaign = self.campaign()
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertEqual(result.error, "")
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertFalse(result.flow_stage_required)
        self.assertEqual(result.file_identifier, campaign.path.name)
        self.assertEqual(result.file_identifier_column, "Filename")
        self.assertEqual(result.ignored_columns, ["cc", "bcc"])
        self.assertEqual(result.mappings, [
            ColumnMapping("email", "Email"),
            ColumnMapping("indexone_uid", "Old I.C."),
            ColumnMapping("primary_key", "Old I.C."),
            ColumnMapping("chname", "Name"),
            ColumnMapping("accountnumber", "Account No."),
        ])
        self.assertEqual(result.missing_rows, [])

    def test_primary_key_disagreement_is_reported_despite_valid_uid(self):
        campaign = self.campaign([["alice@example.test", "", "", "00123", "00456", "Alice Smith", "0005678"]])
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertEqual(result.verdict, "ROW ISSUES")
        self.assertEqual(result.count("DIFFERENT"), 1)
        self.assertEqual(result.differences, [(2, 2, "primary_key → Old I.C.", "00456", "00123")])

    def test_account_numbers_remain_text_and_detect_lost_leading_zeroes(self):
        campaign = self.campaign([["alice@example.test", "", "", "00123", "00123", "Alice Smith", "5678"]])
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertEqual(result.count("DIFFERENT"), 1)
        self.assertEqual(result.differences[0][2:], ("accountnumber → Account No.", "5678", "0005678"))

    def test_primary_key_can_identify_blast_when_uid_column_is_absent(self):
        campaign = self.campaign([["alice@example.test", "00123", "Alice Smith", "0005678"]],
                                 headers=["email", "primary_key", "chname", "accountnumber"])
        self.assertEqual(campaign.column, "primary_key")
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual(result.count("MATCHED"), 1)

    def test_duplicate_database_ids_select_row_with_matching_account(self):
        database = self.database([
            ["Alice Smith", "00123", "11111", "alice@example.test"],
            ["Alice Smith", "00123", "0005678", "alice@example.test"],
        ])
        result = check_files(database, [self.campaign()], self.options)[0]
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertEqual(result.records[0][2], "3")

    def test_missing_and_blank_ids_remain_row_issues(self):
        campaign = self.campaign([
            ["alice@example.test", "", "", "99999", "99999", "Alice Smith", "0005678"],
            ["alice@example.test", "", "", "", "", "Alice Smith", "0005678"],
        ])
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertEqual([row[1] for row in result.records], ["NOT FOUND", "BLANK KEY"])
        self.assertEqual(result.verdict, "ROW ISSUES")

    def test_arbitrary_uid_schema_cannot_bypass_flow_stage(self):
        path = self.csv("unknown.csv", [["indexone_uid", "email"], ["00123", "alice@example.test"]])
        campaign = Source(path, "CSV", "indexone_uid")
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertIn("FLOW_STAGE header", result.error)

    def test_recognized_blast_needs_matching_database_to_bypass_stage(self):
        path = self.csv("chcode-database.csv", [["CHCODE", "Name", "Account No.", "Email"],
                                               ["00123", "Alice Smith", "0005678", "alice@example.test"]])
        result = check_files(self.source(path), [self.campaign()], self.options)[0]
        self.assertIn("FLOW_STAGE header", result.error)

    def test_chcode_file_keeps_stage_requirement_with_datagrid_database(self):
        headers = [*self.campaign_headers, "CHCODE"]
        campaign = self.campaign([["alice@example.test", "", "", "00123", "00123", "Alice Smith", "0005678", "00123"]],
                                 headers=headers)
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertIn("FLOW_STAGE header", result.error)

    def test_present_flow_stage_is_validated_instead_of_bypassed(self):
        headers = [*self.campaign_headers, "FLOW_STAGE"]
        values = ["alice@example.test", "", "", "00123", "00123", "Alice Smith", "0005678"]
        for stage, expected_error in [("", "FLOW_STAGE is blank"), ("MAIL", "")]:
            with self.subTest(stage=stage):
                campaign = self.campaign([[*values, stage]], headers=headers)
                result = check_files(self.database(), [campaign], self.options)[0]
                self.assertEqual(result.flow_stage_required, True)
                if expected_error:
                    self.assertIn(expected_error, result.error)
                else:
                    self.assertEqual(result.verdict, "EXACT MATCH")
                    self.assertEqual(result.file_identifier, "MAIL")
                    self.assertEqual(result.file_identifier_column, "FLOW_STAGE")

    def test_full_copy_mode_does_not_bypass_flow_stage(self):
        result = check_files(self.database(), [self.campaign()], Options(require_flow_stage=True))[0]
        self.assertIn("FLOW_STAGE header", result.error)

    def test_custom_mapping_must_include_selected_uid_key(self):
        campaign = self.campaign()
        campaign = Source(campaign.path, campaign.sheet, campaign.column, (ColumnMapping("email", "Email"),))
        result = check_files(self.database(), [campaign], self.options)[0]
        self.assertIn("indexone_uid-to-Old I.C. mapping is required", result.error)

    def test_datagrid_aliases_take_precedence_over_unrelated_same_named_database_fields(self):
        headers = [*self.database_headers, "indexone_uid", "primary_key", "accountnumber"]
        mappings = suggest_mappings(self.campaign_headers, headers, "indexone_uid", "Old I.C.")
        self.assertIn(ColumnMapping("primary_key", "Old I.C."), mappings)
        self.assertIn(ColumnMapping("accountnumber", "Account No."), mappings)

    def test_existing_chcode_aliases_do_not_switch_to_old_ic(self):
        headers = ["CHCODE", *self.database_headers, "Email address"]
        campaign_headers = ["CHCODE", *self.campaign_headers]
        mappings = suggest_mappings(campaign_headers, headers, "CHCODE", "CHCODE")
        self.assertIn(ColumnMapping("primary_key", "CHCODE"), mappings)
        self.assertIn(ColumnMapping("email", "Email"), mappings)
        self.assertNotIn(ColumnMapping("indexone_uid", "Old I.C."), mappings)

    def test_report_includes_filename_and_honest_field_coverage(self):
        database, campaign = self.database(), self.campaign()
        result = check_files(database, [campaign], self.options)[0]
        report = self.folder / "report.xlsx"
        export_report(report, database, self.options, [result])
        book = load_workbook(report)
        try:
            self.assertEqual(book["Summary"]["A9"].value, campaign.path.name)
            self.assertEqual(book["Summary"]["K9"].value, "MAPPED ROWS MATCH")
            self.assertEqual(book["Summary"]["U9"].value, "cc, bcc")
            mappings = list(book["Mappings 001"].values)[1:]
            self.assertEqual([row[0] for row in mappings if row[3] == "NOT CHECKED"], ["cc", "bcc"])
            self.assertEqual(len([row for row in mappings if row[3] == "CHECKED"]), 5)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
