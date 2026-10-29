"""Regression coverage for DataGrid exports with stale Excel dimensions."""

import io
import re
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from checker import ColumnMapping, Options, Source, check_files, match_column, read_table
from web_checker import Upload, inspect_upload, prepare_uploaded_check, run_prepared_check
from test_web_checker import csv_upload


HEADERS = ["No", "Name", "New I.C.", "Old I.C.", "Account No.", "Email",
           "Account Type", "Amount Referred", "Balance"]
ACCOUNTS = [
    [1, "Alice Smith", None, "00123", "0005678", "alice@example.test", "BUCKET 4: 91-120", "1,000.10", "900.00"],
    [2, "Bob Jones", None, "00456", "0009876", "bob@example.test", "WRITEOFF", "200.20", "150.00"],
]
TOTAL = ["2", None, None, None, None, None, None, "1,200.30", "1,050.00"]
BLAST_HEADERS = ["email", "cc", "bcc", "indexone_uid", "primary_key", "accountnumber", "chname", "campaignbucket"]


def datagrid_bytes(rows, dimension="A1:C1"):
    stream = io.BytesIO()
    book = Workbook()
    try:
        for row in rows:
            book.active.append(row)
        book.save(stream)
    finally:
        book.close()
    changed = io.BytesIO()
    with ZipFile(io.BytesIO(stream.getvalue())) as source, ZipFile(changed, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"/>',
                              f'<dimension ref="{dimension}"/>'.encode(), data)
            target.writestr(entry, data)
    return changed.getvalue()


class ExcelDimensionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)

    def database(self, rows=None, dimension="A1:C1"):
        path = self.folder / "DataGrid.xlsx"
        path.write_bytes(datagrid_bytes(rows if rows is not None else [HEADERS, *ACCOUNTS, TOTAL], dimension))
        return path

    def test_actual_cells_are_read_beyond_understated_rows_and_columns(self):
        for dimension in ("A1:C1", "A1:I1", "A1:C4", "A1", "A1:XFD1048576"):
            with self.subTest(dimension=dimension):
                path = self.database(dimension=dimension)
                original = path.read_bytes()
                preview = read_table(path, headers_only=True)
                self.assertEqual(preview.headers, HEADERS)
                self.assertEqual(match_column(preview.headers), "Old I.C.")
                self.assertEqual(preview.rows, [])
                table = read_table(path)
                self.assertEqual(table.rows, [(2, tuple(ACCOUNTS[0])), (3, tuple(ACCOUNTS[1]))])
                self.assertEqual(path.read_bytes(), original)

    def test_headerless_campaigns_keep_all_actual_rows(self):
        path = self.folder / "campaign.xlsx"
        path.write_bytes(datagrid_bytes([
            ["Alice Smith", "00123"], ["Bob Jones", "00456"],
        ], dimension="A1"))
        table = read_table(path, has_headers=False)
        self.assertEqual(table.header_row, None)
        self.assertEqual(table.rows, [(1, ("Alice Smith", "00123")), (2, ("Bob Jones", "00456"))])

    def test_four_blasts_use_primary_key_and_all_six_mapped_fields(self):
        path = self.database()
        database = Upload(path.name, path.read_bytes())
        self.assertEqual(inspect_upload(database).column, "Old I.C.")
        files = []
        for name in ("BUCKET 4 91-120", "BUCKET 7 181-360", "BUCKET 8 361+", "WRITEOFF"):
            files.append(csv_upload(f"29-10-26 CBS APD EMAIL BLAST {name}.csv", [
                BLAST_HEADERS,
                ["alice@example.test", "", "", "00123", "00123", "0005678", "Alice Smith", "BUCKET 4: 91-120"],
                ["wrong@example.test", "", "", "00456", "00456", "0009876", "Bob Jones", "WRITEOFF"],
                ["alice@example.test", "", "", "unknown", "unknown", "0005678", "Alice Smith", "BUCKET 4: 91-120"],
                ["alice@example.test", "", "", "", "", "0005678", "Alice Smith", "BUCKET 4: 91-120"],
            ]))
        prepared = prepare_uploaded_check(database, files)
        self.assertEqual(prepared.database_error, "")
        self.assertTrue(all(not item.error and item.source.column == "primary_key" for item in prepared.files))
        run = run_prepared_check(prepared)
        for result in run.results:
            self.assertEqual(result.error, "")
            self.assertEqual(result.database_match_column, "Old I.C.")
            self.assertEqual([record[:2] for record in result.records], [
                (2, "MATCHED"), (3, "DIFFERENT"), (4, "NOT FOUND"), (5, "BLANK KEY"),
            ])
            self.assertEqual(len(result.mappings), 6)
            self.assertIn(ColumnMapping("campaignbucket", "Account Type"), result.mappings)
            self.assertEqual(result.ignored_columns, ["cc", "bcc"])
            self.assertEqual(len(result.differences), 1)
            self.assertEqual(result.differences[0][3:], ("wrong@example.test", "bob@example.test"))
        book = load_workbook(io.BytesIO(run.report), read_only=True, data_only=True)
        try:
            for index in range(1, 5):
                rows = list(book[f"File {index:03d}"].values)
                self.assertEqual(len(rows), 5)
                self.assertEqual(rows[1][7], "00123")
                self.assertEqual(book[f"File {index:03d}"].cell(2, 8).data_type, "s")
        finally:
            book.close()
        self.assertEqual(path.read_bytes(), database.data)

    def test_blank_key_accounts_and_unverified_totals_still_fail(self):
        invalid = [
            ["3", None, None, None, None, None, None, "1,200.30", "1,050.00"],
            ["2", None, None, None, None, None, None, "1,200.31", "1,050.00"],
            ["2", "Unkeyed customer", None, None, None, None, None, "1,200.30", "1,050.00"],
            ["2", None, None, None, "0000001", None, None, "1,200.30", "1,050.00"],
            ["2", None, None, None, None, None, None, "invalid", "1,050.00"],
            ["2", None, None, None, None, None, None, "NaN", "1,050.00"],
            ["2", None, None, None, None, None, None, None, None],
        ]
        for footer in invalid:
            with self.subTest(footer=footer):
                path = self.database([HEADERS, *ACCOUNTS, footer])
                self.assertEqual(len(read_table(path).rows), 3)
                with self.assertRaisesRegex(ValueError, "blank values"):
                    check_files(Source(path, "Sheet", "Old I.C."), [], Options(mapped_fields=True))
                prepared = prepare_uploaded_check(Upload(path.name, path.read_bytes()), [])
                self.assertIn("blank values", prepared.database_error)

    def test_missing_database_email_is_a_difference_even_when_the_key_matches(self):
        account = list(ACCOUNTS[0])
        account[5] = None
        path = self.database([HEADERS, account])
        blast = csv_upload("CBS APD.csv", [
            BLAST_HEADERS,
            ["alice@example.test", "", "", "00123", "00123", "0005678", "Alice Smith", "BUCKET 4: 91-120"],
        ])
        prepared = prepare_uploaded_check(Upload(path.name, path.read_bytes()), [blast])
        self.assertEqual(prepared.database_error, "")
        result = run_prepared_check(prepared).results[0]
        self.assertEqual(result.error, "")
        self.assertEqual(result.count("DIFFERENT"), 1)
        self.assertEqual(result.count("MATCHED"), 0)
        self.assertEqual(result.count("NOT FOUND"), 0)
        self.assertEqual(result.differences[0][3:], ("alice@example.test", None))
        self.assertEqual(len(result.differences), 1)


if __name__ == "__main__":
    unittest.main()
