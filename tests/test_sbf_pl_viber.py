import io
import unittest

from openpyxl import load_workbook

from checker import (
    ColumnMapping, Options, is_datagrid_pl_viber_blast, is_datagrid_pl_viber_pair,
    is_filename_identified_pair, match_column,
)
from test_sbf_salad_viber import BLAST_HEADERS as SALAD_HEADERS, salad_row
from test_web_checker import csv_upload, xlsx_upload
from web_checker import UploadedSource, prepare_uploaded_check, run_prepared_check, run_uploaded_check


DATABASE_HEADERS = ["Name", "Old I.C.", "Account No.", "Email", "Last Called No.", "Debtor ID"]
BLAST_HEADERS = ["FullName", "FirstName", "LastName", "Phone", "CHCode", "OB",
                 "Monthly", "Terms", "Promo", "Level", "Debtorid"]


def pl_database(rows=None, headers=None):
    return xlsx_upload("DataGrid(1).xlsx", [headers or DATABASE_HEADERS, *(rows if rows is not None else [
        ["Alice Smith", "00123", "000001", "alice@example.test", "+639171234567", 1001],
        ["Bob Jones", "456", "000002", "bob@example.test", "09181234567", 1002],
    ])])


def pl_row(debtor="1001", name="Alice Smith", phone="09171234567", level="Level_1"):
    return [name, "Campaign Staff", "Template Salutation", phone, "STAFFCODE", "", "", "", "", level, debtor]


def pl_viber_uploads():
    return pl_database(), [csv_upload(f"27-10-26 SBF PL VIBER BLAST {level}.csv", [
        BLAST_HEADERS, pl_row(level=level), pl_row("1002", "Bob Jones", "639181234567", level),
    ]) for level in ("Level_1", "Level_3", "Level_5")]


class PlViberTests(unittest.TestCase):
    def run_blast(self, rows, database=None, headers=None):
        prepared = prepare_uploaded_check(database or pl_database(), [
            csv_upload("PL Viber.csv", [headers or BLAST_HEADERS, *rows]),
        ])
        self.assertEqual(prepared.database_error, "")
        self.assertEqual(prepared.files[0].error, "")
        return run_prepared_check(prepared)

    def test_three_levels_use_debtor_id_and_compare_name_and_phone(self):
        database, files = pl_viber_uploads()
        prepared = prepare_uploaded_check(database, files)
        self.assertEqual(prepared.database_error, "")
        self.assertEqual(prepared.database.column, "Old I.C.")
        self.assertTrue(all(item.source.column == "Debtorid" for item in prepared.files))
        self.assertTrue(all("Debtor ID" in item.note for item in prepared.files))
        run = run_prepared_check(prepared)
        for result, upload in zip(run.results, files):
            with self.subTest(file=upload.name):
                self.assertEqual(result.error, "")
                self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
                self.assertEqual(result.database_match_column, "Debtor ID")
                self.assertEqual(result.file_identifier, upload.name)
                self.assertFalse(result.flow_stage_required)
                self.assertEqual([row[:3] for row in result.records], [(2, "MATCHED", "2"), (3, "MATCHED", "3")])
                self.assertEqual(result.mappings, [ColumnMapping("FullName", "Name"),
                                                   ColumnMapping("Phone", "Last Called No.", "Phone"),
                                                   ColumnMapping("Debtorid", "Debtor ID")])
                self.assertEqual(result.ignored_columns, ["FirstName", "LastName", "CHCode", "OB",
                                                          "Monthly", "Terms", "Promo", "Level"])

    def test_phone_and_name_of_another_customer_cannot_pass_for_the_wrong_id(self):
        result = self.run_blast([
            pl_row(), pl_row("1001", "Bob Jones", "09181234567"),
            pl_row("unknown"), pl_row(""), pl_row(phone=""), pl_row(),
        ]).results[0]
        self.assertEqual([row[:3] for row in result.records], [
            (2, "MATCHED", "2"), (3, "DIFFERENT", "2"), (4, "NOT FOUND", ""),
            (5, "BLANK KEY", ""), (6, "DIFFERENT", "2"), (7, "MATCHED", "2"),
        ])
        self.assertEqual([(row[0], row[1], row[4]) for row in result.differences],
                         [(3, 2, "Alice Smith"), (3, 2, "+639171234567"), (6, 2, "+639171234567")])

    def test_duplicate_database_ids_must_match_all_fields_on_one_row(self):
        database = pl_database(rows=[
            ["Alice Smith", "00123", "000001", "alice@example.test", "09171234567", 1001],
            ["Bob Jones", "456", "000002", "bob@example.test", "09181234567", 1001],
        ])
        result = self.run_blast([pl_row("1001", "Bob Jones", "09181234567"),
                                 pl_row(phone="09181234567"), pl_row()], database).results[0]
        self.assertEqual([row[1] for row in result.records], ["MATCHED", "DIFFERENT", "MATCHED"])
        self.assertEqual([row[2] for row in result.records], ["3", "2", "2"])

    def test_missing_or_blank_database_debtor_ids_and_phone_headers_are_errors(self):
        for headers, row, message in (
            (DATABASE_HEADERS[:-1], ["Alice Smith", "00123", "000001", "a@example.test", "09171234567"], "Debtor ID"),
            (DATABASE_HEADERS, ["Alice Smith", "00123", "000001", "a@example.test", "09171234567", None], "blank"),
            (DATABASE_HEADERS[:4] + DATABASE_HEADERS[5:], ["Alice Smith", "00123", "000001", "a@example.test", 1001], "Last Called No."),
        ):
            with self.subTest(message=message):
                result = self.run_blast([pl_row()], pl_database([row], headers)).results[0]
                self.assertEqual(result.verdict, "ERROR")
                self.assertIn(message, result.error)

    def test_blank_phone_never_passes_even_when_database_phone_is_blank(self):
        database = pl_database([["Alice Smith", "00123", "000001", "a@example.test", None, 1001]])
        result = self.run_blast([pl_row(phone="")], database).results[0]
        self.assertEqual(result.records[0][1], "DIFFERENT")
        self.assertEqual(result.differences[0][4], "(phone number required)")

    def test_mixed_batch_keeps_salad_phone_and_regular_chcode_lookups(self):
        uploads = [csv_upload("pl.csv", [BLAST_HEADERS, pl_row()]),
                   csv_upload("salad.csv", [SALAD_HEADERS, salad_row()]),
                   csv_upload("regular.csv", [["CHCODE", "Name", "Phone"], ["00123", "Alice Smith", "09171234567"]]),
                   csv_upload("pl-again.csv", [BLAST_HEADERS, pl_row()])]
        prepared = prepare_uploaded_check(pl_database(), uploads)
        results = run_prepared_check(prepared).results
        self.assertEqual([result.database_match_column for result in results],
                         ["Debtor ID", "Last Called No.", "Old I.C.", "Debtor ID"])
        self.assertTrue(all(result.count("MATCHED") == 1 for result in results))

    def test_report_preserves_original_values_and_discloses_actual_lookup_and_coverage(self):
        run = self.run_blast([pl_row(name="Changed Name")])
        book = load_workbook(io.BytesIO(run.report))
        try:
            summary = dict(zip(next(book["Summary"].iter_rows(min_row=8, max_row=8, values_only=True)),
                               next(book["Summary"].iter_rows(min_row=9, max_row=9, values_only=True))))
            self.assertEqual(summary["Match column"], "Debtorid")
            self.assertEqual(summary["Database match column"], "Debtor ID")
            self.assertEqual(book["File 001"]["G2"].value, "09171234567")
            coverage = {row[0]: row[1:] for row in book["Mappings 001"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(coverage["Debtorid"][:3], ("Debtor ID", "Text", "CHECKED"))
            self.assertIn("same Debtor ID", coverage["Phone"][3])
            self.assertIn("same database row", coverage["FullName"][3])
            self.assertEqual(coverage["CHCode"][2], "NOT CHECKED")
            self.assertEqual(coverage["LastName"][2], "NOT CHECKED")
            self.assertEqual(book["Differences 001"]["E2"].value, "Alice Smith")
        finally:
            book.close()

    def test_case_preamble_column_order_leading_zeros_and_name_rules(self):
        headers = [name.upper() for name in reversed(BLAST_HEADERS)]
        values = list(reversed(pl_row("001001", "Al?ice Smith??", "+63 (917) 123-4567")))
        upload = xlsx_upload("PL.xlsx", [["Campaign export"], [], headers, values])
        database = pl_database([["Alice Smith", "00123", "000001", "a@example.test", "09171234567", "001001"]])
        prepared = prepare_uploaded_check(database, [upload])
        self.assertEqual(prepared.files[0].inspection.header_row, 3)
        result = run_prepared_check(prepared).results[0]
        self.assertEqual(result.source.column, "DEBTORID")
        self.assertEqual(result.records[0][:3], (4, "MATCHED", "2"))
        self.assertEqual(result.records[0][3][0], "001001")
        result = self.run_blast([pl_row("1001")], database).results[0]
        self.assertEqual(result.records[0][1], "NOT FOUND")

    def test_schema_detection_is_specific_and_does_not_change_ordinary_chcode_files(self):
        self.assertTrue(is_datagrid_pl_viber_blast(BLAST_HEADERS))
        self.assertTrue(is_datagrid_pl_viber_pair(BLAST_HEADERS, DATABASE_HEADERS))
        self.assertTrue(is_filename_identified_pair(BLAST_HEADERS, DATABASE_HEADERS))
        self.assertEqual(match_column(BLAST_HEADERS), "Debtorid")
        self.assertEqual(match_column(["CHCODE", "Phone", "FullName", "Debtorid"]), "CHCODE")
        for index in range(len(BLAST_HEADERS)):
            self.assertFalse(is_datagrid_pl_viber_blast(BLAST_HEADERS[:index] + BLAST_HEADERS[index + 1:]))
        self.assertFalse(is_datagrid_pl_viber_pair(BLAST_HEADERS, DATABASE_HEADERS, "CHCode", "Old I.C."))

    def test_empty_files_present_stage_and_full_copy_keep_validation(self):
        self.assertIn("no data rows", self.run_blast([]).results[0].error)
        database = UploadedSource(pl_database(), "Sheet", "Old I.C.")
        source = UploadedSource(csv_upload("staged.csv", [
            [*BLAST_HEADERS, "FLOW_STAGE"], [*pl_row(), ""],
        ]), "CSV", "Debtorid")
        result = run_uploaded_check(database, [source], Options(mapped_fields=True, require_flow_stage=True)).results[0]
        self.assertIn("FLOW_STAGE is blank", result.error)
        source = UploadedSource(csv_upload("pl.csv", [BLAST_HEADERS, pl_row()]), "CSV", "Debtorid")
        result = run_uploaded_check(database, [source], Options(require_flow_stage=True)).results[0]
        self.assertIn("FLOW_STAGE header", result.error)


if __name__ == "__main__":
    unittest.main()
