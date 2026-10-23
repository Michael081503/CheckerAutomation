import csv
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from scripts.check_files import main


class CommandLineTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.database = self.folder / "database.xlsx"
        book = Workbook()
        book.active.append(["Name", "Old I.C.", "Account No.", "Email", "Account Type"])
        book.active.append(["André", "00123", "000456", "andre@example.test", "LOAN"])
        book.save(self.database)
        book.close()
        self.headers = ["primary_key", "indexone_uid", "email", "chname", "accountnumber", "campaignbucket"]
        self.values = ["OTHER", "00123", "andre@example.test", "André", "000456", "LOAN"]
        self.target = self.csv("CBS PL.csv", [self.headers, self.values])
        self.output = self.folder / "report.xlsx"

    def csv(self, name, rows):
        path = self.folder / name
        with path.open("w", newline="", encoding="cp1252") as stream:
            csv.writer(stream).writerows(rows)
        return path

    def run_cli(self, *extra, files=None, output=None):
        arguments = ["check_files.py", "--database", str(self.database),
                     "--files", *map(str, files or [self.target]),
                     "--output", str(output or self.output), *extra]
        stdout = io.StringIO()
        with patch.object(sys, "argv", arguments), redirect_stdout(stdout):
            exit_code = main()
        results = [json.loads(line) for line in stdout.getvalue().splitlines() if line.startswith("{")]
        return exit_code, results

    def test_explicit_uid_lookup_resolves_headers_and_reports_other_identifier_difference(self):
        original_database = self.database.read_bytes()
        original_target = self.target.read_bytes()
        exit_code, results = self.run_cli("--file-key", " INDEXONE_UID ", "--database-key", " old i.c. ")
        self.assertEqual(exit_code, 0)
        self.assertEqual(results[0]["file_key"], "indexone_uid")
        self.assertEqual(results[0]["database_key"], "Old I.C.")
        self.assertEqual((results[0]["not_found"], results[0]["changed"], results[0]["checked_fields"]), (0, 1, 6))
        book = load_workbook(self.output, read_only=True, data_only=True)
        try:
            self.assertEqual(book["Summary"]["B3"].value, "Old I.C.")
            self.assertEqual(book["Summary"]["D9"].value, "indexone_uid")
            record = list(book["File 001"].values)[1]
            self.assertEqual(record[:3], (2, "DIFFERENT", "2"))
            self.assertEqual(record[6], "André")
            differences = list(book["Differences 001"].values)[1:]
            self.assertEqual(len(differences), 1)
            self.assertIn("primary_key", differences[0][2])
        finally:
            book.close()
        self.assertEqual(self.database.read_bytes(), original_database)
        self.assertEqual(self.target.read_bytes(), original_target)

    def test_default_bucket_lookup_remains_primary_key(self):
        exit_code, results = self.run_cli()
        self.assertEqual(exit_code, 0)
        self.assertEqual(results[0]["file_key"], "primary_key")
        self.assertEqual(results[0]["database_key"], "Old I.C.")
        self.assertEqual(results[0]["not_found"], 1)
        self.assertEqual(results[0]["verdict"], "ROW ISSUES")

    def test_missing_explicit_key_is_per_file_error_and_other_files_continue(self):
        missing = self.csv("missing.csv", [self.headers[:1] + self.headers[2:],
                                           ["00123", *self.values[2:]]])
        valid = self.csv("valid.csv", [self.headers, ["00123", *self.values[1:]]])
        exit_code, results = self.run_cli("--file-key", "indexone_uid", files=[missing, valid])
        self.assertEqual(exit_code, 1)
        self.assertEqual(results[0]["verdict"], "ERROR")
        self.assertIn("indexone_uid", results[0]["error"])
        self.assertEqual(results[0]["rows"], 0)
        self.assertEqual(results[1]["verdict"], "MAPPED ROWS MATCH")
        self.assertEqual(results[1]["matched"], 1)
        self.assertTrue(self.output.is_file())

    def test_output_cannot_overwrite_input_with_explicit_key(self):
        original = self.target.read_bytes()
        with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as stopped:
            self.run_cli("--file-key", "indexone_uid", output=self.target)
        self.assertEqual(stopped.exception.code, 2)
        self.assertEqual(self.target.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
