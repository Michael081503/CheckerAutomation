"""Exercise the real Tk interface without showing file dialogs."""

import csv
import sys
import tempfile
import time
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import CheckerApp, FULL_COPY, UPLOADED_ROWS
from mapping_dialog import MappingDialog, SKIP
from openpyxl import load_workbook


def wait(app):
    deadline = time.monotonic() + 15
    while app.busy:
        app.update()
        if time.monotonic() > deadline:
            raise AssertionError("The background action did not finish.")
        time.sleep(0.01)
    app.update()


def main():
    with tempfile.TemporaryDirectory() as folder:
        root = Path(folder)
        inputs = {
            "database.csv": [["CHCODE", "Name"], ["ABC", "Alice"], ["00123", "Zoe"]],
            "first.csv": [["CHCODE", "Name"], [" abc ", "Alicia"], ["MISSING", "Bob"], ["", "Blank"]],
            "second.csv": [["Name", "CHCODE"], ["Zoe", "00123"], ["Alice", "ABC"]],
            "invalid.csv": [["Other"], ["ABC"]],
            "DataGrid.csv": [["Name", "Old I.C.", "Account No.", "Email", "Account Type"],
                             ["Alice", "MSPL-001", "00123", "alice@example.test", "BUCKET 4 91-120"]],
            "email_blast.csv": [["email", "cc", "bcc", "indexone_uid", "primary_key", "chname", "accountnumber"],
                                ["alice@example.test", "", "", "MSPL-001", "MSPL-001", "Alice", "00123"]],
            "cbs_apd.csv": [["email", "cc", "bcc", "indexone_uid", "primary_key", "accountnumber", "chname", "campaignbucket"],
                            ["alice@example.test", "", "", "WRONG-UID", "MSPL-001", "00123", "Alice", "BUCKET 4 91-120"]],
        }
        for name, rows in inputs.items():
            with (root / name).open("w", newline="", encoding="utf-8") as stream:
                csv.writer(stream).writerows(rows)
        app = CheckerApp()
        app.withdraw()
        try:
            with patch("app.filedialog.askopenfilename", return_value=str(root / "database.csv")):
                app._choose_database()
            wait(app)
            with patch("app.filedialog.askopenfilenames", return_value=[str(root / name) for name in ["first.csv", "second.csv", "invalid.csv"]]):
                app._add_files()
            wait(app)
            assert len(app.files) == 3
            assert app.metrics.get().startswith("3 files")
            app._run()
            wait(app)
            assert len(app.results) == 3
            assert sum(item.count("MATCHED") for item in app.results) == 2
            assert app.results[0].count("DIFFERENT") == 1
            assert app.results[0].count("MISSING FROM FILE") == 0
            assert app.results[1].verdict == "MAPPED ROWS MATCH"
            assert sum(bool(item.error) for item in app.results) == 1
            app.file_tree.selection_set(next(iter(app.files)))
            app.update()
            app.filter_value.set("NOT FOUND")
            app._preview()
            assert len(app.preview_tree.get_children()) == 1
            app.filter_value.set("DIFFERENT")
            app._preview()
            assert "Alicia" in app.preview_tree.item(app.preview_tree.get_children()[0], "values")[4]
            app.filter_value.set("MISSING FROM FILE")
            app._preview()
            assert len(app.preview_tree.get_children()) == 0
            destination = root / "report.xlsx"
            with patch("app.filedialog.asksaveasfilename", return_value=str(destination)):
                app._export()
            wait(app)
            book = load_workbook(destination)
            assert len(book.sheetnames) == 10
            assert book["Summary"]["K10"].value == "MAPPED ROWS MATCH"
            assert book["Differences 001"]["D2"].value == "Alicia"
            assert "Mappings 001" in book.sheetnames
            book.close()
            app._edit_mappings()
            dialog = next(child for child in app.winfo_children() if isinstance(child, MappingDialog))
            for name, destination, comparison in dialog.fields:
                if name == "Name":
                    destination.set(SKIP)
            dialog._save()
            assert not app.results
            app._run()
            wait(app)
            assert app.results[0].count("MATCHED") == 1
            assert app.results[0].ignored_columns == ["Name"]
            app.mode_value.set(FULL_COPY)
            app._run()
            wait(app)
            assert app.results[0].count("DIFFERENT") == 1
            assert app.results[0].count("MISSING FROM FILE") == 1
            assert app.results[1].verdict == "SAME ROWS"
            app.ignore_case.set(False)
            app._invalidate()
            assert not app.results
            assert str(app.export_button['state']) == "disabled"
            app._clear()
            app.mode_value.set(UPLOADED_ROWS)
            with patch("app.filedialog.askopenfilename", return_value=str(root / "DataGrid.csv")):
                app._choose_database()
            wait(app)
            with patch("app.filedialog.askopenfilenames", return_value=[str(root / "email_blast.csv")]):
                app._add_files()
            wait(app)
            assert app.database['column'] == "Old I.C."
            assert next(iter(app.files.values()))['column'] == "indexone_uid"
            app._run()
            wait(app)
            assert app.results[0].verdict == "MAPPED ROWS MATCH"
            assert len(app.results[0].mappings) == 5
            assert app.results[0].ignored_columns == ["cc", "bcc"]
            app._clear()
            with patch("app.filedialog.askopenfilenames", return_value=[str(root / "cbs_apd.csv")]):
                app._add_files()
            wait(app)
            assert next(iter(app.files.values()))['column'] == "primary_key"
            app._run()
            wait(app)
            result = app.results[0]
            assert not result.error
            assert result.count("DIFFERENT") == 1
            assert result.count("NOT FOUND") == 0
            assert result.records[0][2] == "2"
            assert len(result.mappings) == 6
            assert len(result.differences) == 1
            assert result.differences[0][2].startswith("indexone_uid")
            app.file_key.set("indexone_uid")
            app._change_file_column()
            assert not app.results
            assert str(app.export_button['state']) == "disabled"
            assert next(iter(app.files.values()))['column'] == "indexone_uid"
            app._run()
            wait(app)
            assert app.results[0].source.column == "indexone_uid"
            assert app.results[0].count("NOT FOUND") == 1
            print("GUI workflow passed: uploads, mapped validation, cell differences, custom field mappings, export, full-copy mode, invalidation, MSB and CBS DataGrid email blasts, lookup column selection.")
        finally:
            app.destroy()


if __name__ == "__main__":
    main()
