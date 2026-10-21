import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from checker import Cancelled, ColumnMapping, Options, Source, check_files, compare_value, export_report, normalize, read_table, sheet_names


class CheckerTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)

    def csv(self, name, rows, delimiter=","):
        path = self.folder / name
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.writer(stream, delimiter=delimiter)
            writer.writerows(rows)
        return path

    def excel(self, name, sheets):
        path = self.folder / name
        book = Workbook()
        book.remove(book.active)
        for title, rows in sheets.items():
            sheet = book.create_sheet(title)
            for row in rows:
                sheet.append(row)
        book.save(path)
        book.close()
        return path

    def source(self, path, sheet="CSV", column="CHCODE"):
        return Source(path, sheet, column)

    def test_multiple_files_with_missing_blank_and_duplicate_codes(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A1", "Alice"], ["A1", "Again"], ["00123", "Zero"], ["B2", "Charlie"]])
        first = self.csv("one.csv", [["CHCODE", "Name"], [" a1 ", "Alice"], ["x9", "Bob"], ["", "Blank"], ["00123", "Zero"], ["123", "Other"], ["A1", "Again"]])
        second = self.csv("two.csv", [["Name", "CHCODE"], ["Charlie", "b2"]])
        progress = []
        results = check_files(self.source(db), [self.source(first), self.source(second)], Options(), lambda n, total, result: progress.append((n, total)))
        self.assertEqual([item[1] for item in results[0].records], ["MATCHED", "NOT FOUND", "BLANK KEY", "MATCHED", "NOT FOUND", "MATCHED"])
        self.assertEqual(results[0].records[0][2], "2")
        self.assertEqual(results[1].count("MATCHED"), 1)
        self.assertEqual(progress, [(1, 2), (2, 2)])

    def test_numeric_excel_values_and_text_leading_zeros(self):
        db = self.excel("database.xlsx", {"Codes": [["CHCODE"], [123], [456.0]]})
        target = self.csv("input.csv", [["CHCODE"], ["123"], ["456"], ["00123"]])
        result = check_files(self.source(db, "Codes"), [self.source(target)], Options())[0]
        self.assertEqual([row[1] for row in result.records], ["MATCHED", "MATCHED", "NOT FOUND"])

    def test_case_and_trim_can_be_disabled(self):
        db = self.csv("db.csv", [["CHCODE"], ["ABC"]])
        target = self.csv("target.csv", [["CHCODE"], ["abc"], [" ABC "], ["ABC"]])
        result = check_files(self.source(db), [self.source(target)], Options(False, False))[0]
        self.assertEqual([row[1] for row in result.records], ["NOT FOUND", "NOT FOUND", "MATCHED"])

    def test_file_errors_do_not_stop_other_files(self):
        db = self.csv("db.csv", [["CHCODE"], ["A"]])
        invalid = self.csv("invalid.csv", [["Other"], ["A"]])
        valid = self.csv("valid.csv", [["CHCODE"], ["A"]])
        results = check_files(self.source(db), [self.source(invalid), self.source(valid)], Options())
        self.assertIn("CHCODE", results[0].error)
        self.assertEqual(results[1].count("MATCHED"), 1)

    def test_missing_database_header_stops_check(self):
        db = self.csv("db.csv", [["Other"], ["A"]])
        with self.assertRaisesRegex(ValueError, "Database column"):
            check_files(self.source(db), [], Options())

    def test_empty_database_key_column_stops_check(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["", "Alice"]])
        with self.assertRaisesRegex(ValueError, "no nonblank"):
            check_files(self.source(db), [], Options())

    def test_empty_rows_skipped_and_source_row_numbers_preserved(self):
        path = self.csv("rows.csv", [["CHCODE", "Name"], ["A", "Alice"], [], ["", "Bob"]])
        table = read_table(path)
        self.assertEqual([row[0] for row in table.rows], [2, 4])

    def test_empty_headers_rejected(self):
        with self.assertRaises(ValueError):
            read_table(self.csv("empty.csv", [["", "CHCODE"], ["A", "B"]]))

    def test_duplicate_headers_preserve_all_values_with_visible_suffixes(self):
        path = self.csv("duplicates.csv", [["CHCODE", "CHCode", "blank", "blank", "CHCode [2]"], ["A", "B", "C", "D", "E"]])
        table = read_table(path)
        self.assertEqual(table.headers, ["CHCODE", "CHCode [2]", "blank", "blank [2]", "CHCode [2] [2]"])
        self.assertEqual(table.rows[0][1], ("A", "B", "C", "D", "E"))

    def test_data_under_unnamed_column_rejected(self):
        path = self.csv("bad.csv", [["CHCODE"], ["A", "extra"]])
        with self.assertRaisesRegex(ValueError, "unnamed column"):
            read_table(path)

    def test_csv_delimiter_and_encoding(self):
        path = self.folder / "windows.csv"
        path.write_bytes("CHCODE;Name\r\nA;Andr\xe9\r\n".encode("cp1252"))
        table = read_table(path)
        self.assertEqual(table.headers, ["CHCODE", "Name"])
        self.assertEqual(table.rows[0][1], ("A", "Andr\xe9"))

    def test_explicit_excel_worksheet(self):
        path = self.excel("sheets.xlsx", {"Other": [["Other"], ["B"]], "Codes": [[" CHCODE "], ["A"]]})
        self.assertEqual(sheet_names(path), ["Other", "Codes"])
        self.assertEqual(read_table(path, "Codes").headers, ["CHCODE"])

    def test_chcode_header_is_detected_below_title_and_blank_rows(self):
        rows = [["Account report"], [], ["Name", " chcode "], ["Alice", "00123"], [], ["Bob", "456"]]
        sources = [(self.csv("preamble.csv", rows), "CSV"),
                   (self.excel("preamble.xlsx", {"Accounts": rows}), "Accounts")]
        for path, sheet in sources:
            with self.subTest(format=path.suffix):
                table = read_table(path, sheet)
                self.assertEqual(table.header_row, 3)
                self.assertEqual(table.headers, ["Name", "chcode"])
                self.assertEqual(table.rows, [(4, ("Alice", "00123")), (6, ("Bob", "456"))])
                info = read_table(path, sheet, headers_only=True)
                self.assertEqual((info.header_row, info.headers, info.rows), (3, table.headers, []))

    def test_csv_preamble_does_not_break_delimiter_detection(self):
        for index, delimiter in enumerate([",", ";", "\t", "|"]):
            with self.subTest(delimiter=delimiter):
                path = self.csv(f"preamble_{index}.csv", [
                    ["Account report"], [], ["Name", "CHCODE"], ["Alice", "00123"],
                ], delimiter=delimiter)
                table = read_table(path)
                self.assertEqual(table.header_row, 3)
                self.assertEqual(table.headers, ["Name", "CHCODE"])
                self.assertEqual(table.rows, [(4, ("Alice", "00123"))])

    def test_delayed_database_and_target_headers_keep_original_row_numbers(self):
        database = self.csv("delayed_db.csv", [
            ["Database export"], [], ["Name", "CHCODE"], ["Alice", "00123"],
        ])
        target = self.csv("delayed_target.csv", [
            ["Campaign export"], [], [], [], ["Name", "CHCODE"], ["Alice", "00123"],
        ])
        result = check_files(self.source(database), [self.source(target)], Options(mapped_fields=True))[0]
        self.assertFalse(result.error)
        self.assertEqual(result.records, [(6, "MATCHED", "4", ("Alice", "00123"))])

    def test_explicit_headerless_mode_does_not_skip_a_data_value_named_chcode(self):
        path = self.csv("no_headers.csv", [["First", "00123"], ["Second", "CHCODE"]])
        table = read_table(path, has_headers=False)
        self.assertIsNone(table.header_row)
        self.assertEqual(table.rows, [(1, ("First", "00123")), (2, ("Second", "CHCODE"))])

    def test_cancellation(self):
        db = self.csv("db.csv", [["CHCODE"], ["A"]])
        with self.assertRaises(Cancelled):
            check_files(self.source(db), [], Options(), cancelled=lambda: True)

    def test_report_round_trip_preserves_text_and_original_columns(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["=1+1", "=SUM(1,2)"], ["00123", "Alice"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["=1+1", "=SUM(1,2)"], ["00123", "Alice"]])
        source = self.source(db)
        results = check_files(source, [self.source(target)], Options())
        report = self.folder / "report.xlsx"
        export_report(report, source, Options(), results)
        book = load_workbook(report, data_only=False)
        try:
            self.assertEqual(book.sheetnames, ["Summary", "File 001", "Differences 001", "Missing 001"])
            sheet = book["File 001"]
            self.assertEqual(sheet["D2"].value, "=1+1")
            self.assertEqual(sheet["D2"].data_type, "s")
            self.assertEqual(sheet["E2"].value, "=SUM(1,2)")
            self.assertEqual(sheet["E2"].data_type, "s")
            self.assertEqual(sheet["D3"].value, "00123")
            self.assertEqual(sheet["C2"].value, "2")
            self.assertEqual(book["Summary"]["F9"].value, 2)
        finally:
            book.close()

    def test_report_includes_file_errors(self):
        db = self.csv("db.csv", [["CHCODE"], ["A"]])
        target = self.csv("invalid.csv", [["Other"], ["A"]])
        source = self.source(db)
        results = check_files(source, [self.source(target)], Options())
        report = self.folder / "errors.xlsx"
        export_report(report, source, Options(), results)
        book = load_workbook(report)
        try:
            self.assertIn("CHCODE", book["Summary"]["I9"].value)
            self.assertEqual(book["File 001"]["A2"].value, "ERROR")
        finally:
            book.close()

    def test_normalization_does_not_strip_text_leading_zeros(self):
        self.assertEqual(normalize(123.0, Options()), "123")
        self.assertEqual(normalize("00123", Options()), "00123")
        self.assertEqual(normalize(None, Options()), "")

    def test_full_rows_equal_with_reordered_rows_and_columns(self):
        db = self.csv("db.csv", [["CHCODE", "Name", "Amount"], ["A", "Alice", 100], ["B", "Bob", 200]])
        target = self.csv("target.csv", [["amount", "chcode", "name"], [200, "B", "Bob"], [100, "A", "Alice"]])
        result = check_files(self.source(db), [self.source(target, column="chcode")], Options())[0]
        self.assertEqual(result.verdict, "SAME ROWS")
        self.assertEqual(result.count("MATCHED"), 2)
        self.assertEqual([row[2] for row in result.records], ["3", "2"])

    def test_same_chcode_with_different_cell_is_different(self):
        db = self.csv("db.csv", [["CHCODE", "Name", "Amount"], ["A", "Alice", 100]])
        target = self.csv("target.csv", [["CHCODE", "Name", "Amount"], ["A", "Alicia", 101]])
        result = check_files(self.source(db), [self.source(target)], Options())[0]
        self.assertEqual(result.verdict, "DIFFERENT ROWS")
        self.assertEqual(result.count("DIFFERENT"), 1)
        self.assertEqual(result.differences, [(2, 2, "Name", "Alicia", "Alice"), (2, 2, "Amount", "101", "100")])
        self.assertFalse(result.missing_rows)

    def test_subset_is_different_and_lists_missing_database_rows(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"], ["B", "Bob"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["A", "Alice"]])
        result = check_files(self.source(db), [self.source(target)], Options())[0]
        self.assertEqual(result.verdict, "DIFFERENT ROWS")
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertEqual(result.missing_rows, [(3, ("B", "Bob"))])

    def test_extra_duplicate_row_prevents_equal_verdict(self):
        db = self.csv("db.csv", [["CHCODE"], ["A"]])
        target = self.csv("target.csv", [["CHCODE"], ["A"], ["A"]])
        result = check_files(self.source(db), [self.source(target)], Options())[0]
        self.assertEqual(result.verdict, "DIFFERENT ROWS")
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertEqual(result.count("EXTRA ROW"), 1)

    def test_missing_duplicate_database_row_prevents_equal_verdict(self):
        db = self.csv("db.csv", [["CHCODE"], ["A"], ["A"]])
        target = self.csv("target.csv", [["CHCODE"], ["A"]])
        result = check_files(self.source(db), [self.source(target)], Options())[0]
        self.assertEqual(result.verdict, "DIFFERENT ROWS")
        self.assertEqual(result.missing_rows, [(3, ("A",))])

    def test_duplicate_keys_reserve_exact_rows_before_modified_rows(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"], ["A", "Bob"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["A", "Changed"], ["A", "Alice"]])
        result = check_files(self.source(db), [self.source(target)], Options())[0]
        self.assertEqual([row[1] for row in result.records], ["DIFFERENT", "MATCHED"])
        self.assertEqual([row[2] for row in result.records], ["3", "2"])
        self.assertEqual(result.differences, [(2, 3, "Name", "Changed", "Bob")])

    def test_missing_and_extra_columns_prevent_equal_verdict(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"]])
        target = self.csv("target.csv", [["CHCODE", "Other"], ["A", "Alice"]])
        result = check_files(self.source(db), [self.source(target)], Options())[0]
        self.assertEqual(result.verdict, "DIFFERENT ROWS")
        self.assertEqual(result.missing_columns, ["Name"])
        self.assertEqual(result.extra_columns, ["Other"])
        self.assertEqual(result.count("MATCHED"), 0)

    def test_header_only_file_reports_all_database_rows_missing(self):
        db = self.csv("db.csv", [["CHCODE"], ["A"]])
        target = self.csv("target.csv", [["CHCODE"]])
        result = check_files(self.source(db), [self.source(target)], Options())[0]
        self.assertEqual(result.verdict, "DIFFERENT ROWS")
        self.assertEqual(result.count("MISSING FROM FILE"), 1)

    def test_database_blank_key_rejected_to_avoid_ignoring_rows(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"], ["", "Bob"]])
        with self.assertRaisesRegex(ValueError, "blank CHCODEs"):
            check_files(self.source(db), [], Options())

    def test_report_contains_cell_differences_missing_rows_and_verdict(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"], ["B", "Bob"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["A", "=1+1"]])
        source = self.source(db)
        results = check_files(source, [self.source(target)], Options())
        report = self.folder / "differences.xlsx"
        export_report(report, source, Options(), results)
        book = load_workbook(report)
        try:
            self.assertEqual(book["Summary"]["K9"].value, "DIFFERENT ROWS")
            self.assertEqual(book["Summary"]["L9"].value, 1)
            self.assertEqual(book["Summary"]["N9"].value, 1)
            difference = book["Differences 001"]
            self.assertEqual(difference["C2"].value, "Name")
            self.assertEqual(difference["D2"].value, "=1+1")
            self.assertEqual(difference["D2"].data_type, "s")
            self.assertEqual(difference["E2"].value, "Alice")
            self.assertEqual(list(book["Missing 001"].values)[1], (3, "B", "Bob"))
        finally:
            book.close()

    def test_uploaded_rows_validate_subset_with_hsbc_aliases(self):
        db = self.excel("db.xlsx", {"Sheet1": [
            ["CHCODE", "CHCode", "Name", "Last Name", "Email address", "OB", "PROD TYPE", "blank", "blank"],
            ["A", "A", "Alice Smith", "Smith", "alice@example.test", 1234.5, "CARDS", None, None],
            ["B", "B", "Bob Jones", "Jones", "bob@example.test", 500, "LOAN", None, None],
        ]})
        target = self.csv("campaign.csv", [
            ["ChCode", "primary_key", "email", "chname", "chname_format", "OB_AT_ENDORSEMENT", "product_type", "endo", "TEMPLATE_FAMILY"],
            ["A", "A", "alice@example.test", "Alice Smith", "Smith", "1,234.50", "Credit Card", "12OCT ENDO", "CAMPAIGN"],
        ])
        result = check_files(self.source(db, "Sheet1"), [self.source(target, column="ChCode")], Options(mapped_fields=True))[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertEqual(len(result.mappings), 7)
        self.assertEqual(result.ignored_columns, ["endo", "TEMPLATE_FAMILY"])
        self.assertFalse(result.missing_rows)

    def test_uploaded_rows_detect_cell_changes_and_primary_key_disagreement(self):
        db = self.csv("db.csv", [["CHCODE", "Name", "OB"], ["A", "Alice", 100]])
        target = self.csv("target.csv", [["ChCode", "primary_key", "chname", "OB_AT_ENDORSEMENT"], ["A", "B", "Changed", 200]])
        result = check_files(self.source(db), [self.source(target, column="ChCode")], Options(mapped_fields=True))[0]
        self.assertEqual(result.verdict, "ROW ISSUES")
        self.assertEqual(result.count("DIFFERENT"), 1)
        self.assertEqual(len(result.differences), 3)
        self.assertEqual(result.differences[0][2], "primary_key → CHCODE")

    def test_uploaded_duplicate_rows_can_reuse_database_row(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"], ["B", "Bob"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["A", "Alice"], ["A", "Alice"]])
        result = check_files(self.source(db), [self.source(target)], Options(mapped_fields=True))[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual(result.count("MATCHED"), 2)
        self.assertEqual([row[2] for row in result.records], ["2", "2"])

    def test_uploaded_rows_with_duplicate_database_key_choose_matching_values(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"], ["A", "Bob"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["A", "Bob"]])
        result = check_files(self.source(db), [self.source(target)], Options(mapped_fields=True))[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual(result.records[0][2], "3")

    def test_custom_mappings_are_used_and_skipped_fields_are_disclosed(self):
        db = self.csv("db.csv", [["CHCODE", "Email address"], ["A", "alice@example.test"]])
        target = self.csv("target.csv", [["CHCODE", "Recipient", "Other"], ["A", "alice@example.test", "campaign"]])
        source = Source(target, "CSV", "CHCODE", (ColumnMapping("CHCODE", "CHCODE"), ColumnMapping("Recipient", "Email address")))
        result = check_files(self.source(db), [source], Options(mapped_fields=True))[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual(result.ignored_columns, ["Other"])

    def test_custom_mappings_cannot_skip_the_key(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["A", "Alice"]])
        source = Source(target, "CSV", "CHCODE", (ColumnMapping("Name", "Name"),))
        result = check_files(self.source(db), [source], Options(mapped_fields=True))[0]
        self.assertEqual(result.verdict, "ERROR")
        self.assertIn("CHCODE-to-CHCODE", result.error)

    def test_custom_mapping_to_removed_column_reports_error(self):
        db = self.csv("db.csv", [["CHCODE"], ["A"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["A", "Alice"]])
        source = Source(target, "CSV", "CHCODE", (ColumnMapping("CHCODE", "CHCODE"), ColumnMapping("Name", "Deleted")))
        result = check_files(self.source(db), [source], Options(mapped_fields=True))[0]
        self.assertIn("no longer exists", result.error)

    def test_empty_uploaded_file_does_not_pass_validation(self):
        db = self.csv("db.csv", [["CHCODE"], ["A"]])
        target = self.csv("target.csv", [["CHCODE"]])
        result = check_files(self.source(db), [self.source(target)], Options(mapped_fields=True))[0]
        self.assertEqual(result.verdict, "ERROR")
        self.assertIn("no data rows", result.error)

    def test_uploaded_rows_report_not_found_and_blank_keys(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"]])
        target = self.csv("target.csv", [["CHCODE", "Name"], ["X", "Extra"], ["", "Blank"]])
        result = check_files(self.source(db), [self.source(target)], Options(mapped_fields=True))[0]
        self.assertEqual([row[1] for row in result.records], ["NOT FOUND", "BLANK KEY"])
        self.assertEqual(result.verdict, "ROW ISSUES")

    def test_uploaded_report_includes_field_coverage_and_excludes_missing_database_rows(self):
        db = self.csv("db.csv", [["CHCODE", "Name"], ["A", "Alice"], ["B", "Bob"]])
        target = self.csv("target.csv", [["CHCODE", "chname", "endo"], ["A", "Alice", "12OCT ENDO"]])
        source = self.source(db)
        options = Options(mapped_fields=True)
        result = check_files(source, [self.source(target)], options)[0]
        report = self.folder / "mapped.xlsx"
        export_report(report, source, options, [result])
        book = load_workbook(report)
        try:
            self.assertEqual(book.sheetnames, ["Summary", "File 001", "Differences 001", "Mappings 001"])
            self.assertEqual(book["Summary"]["K9"].value, "MAPPED ROWS MATCH")
            self.assertEqual(book["Summary"]["S9"].value, "Uploaded rows")
            self.assertEqual(book["Summary"]["U9"].value, "endo")
            self.assertEqual(list(book["Mappings 001"].values)[3][3], "NOT CHECKED")
        finally:
            book.close()

    def test_product_and_number_comparison_rules(self):
        self.assertEqual(compare_value("Personal Loan", "Product", Options()), compare_value("LOAN", "Product", Options()))
        self.assertEqual(compare_value("1,200.00", "Number", Options()), compare_value(1200, "Number", Options()))
        self.assertEqual(compare_value("0.00", "Number", Options()), "0")
        self.assertEqual(compare_value("invalid", "Number", Options()), "invalid")
        self.assertNotEqual(compare_value("00123", "Text", Options()), compare_value(123, "Text", Options()))


if __name__ == "__main__":
    unittest.main()
