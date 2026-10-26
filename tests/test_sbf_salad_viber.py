import io
import unittest

from openpyxl import load_workbook

from checker import (
    ColumnMapping, Options, is_datagrid_viber_blast, is_datagrid_viber_pair,
    is_filename_identified_pair, match_column,
)
from test_web_checker import csv_upload, xlsx_upload
from web_checker import UploadedSource, prepare_uploaded_check, run_prepared_check, run_uploaded_check


DATABASE_HEADERS = ["Name", "Old I.C.", "Account No.", "Email", "Last Called No.",
                    "Balance", "Monthly Instalment"]
BLAST_HEADERS = ["FullName", "FirstName", "LastName", "Phone", "CHCode", "OB",
                 "Monthly", "Terms", "Promo", "Level Endo"]


def salad_database(rows=None, headers=None):
    return xlsx_upload("DataGrid salad.xlsx", [headers or DATABASE_HEADERS, *(rows if rows is not None else [
        ["Alice Smith", "00123", "000001", "alice@example.test", "+639171234567", 10000, 1000],
        ["Bob Jones", "456", "000002", "bob@example.test", "09181234567", 5000, 500],
        ["No recorded phone", "789", "000003", "none@example.test", None, 3000, 300],
    ])])


def salad_row(name="Alice Smith", phone="09171234567", level="Level_1"):
    return [name, "Alice", "Smith", phone, "campaign-label", "6,500.00", "3,900.00", "1,950.00", "", level]


def salad_viber_uploads(changed_name=False):
    files = [csv_upload(f"SBF SALAD VIBER BLAST {level}.csv", [
        BLAST_HEADERS, salad_row(level=level),
        salad_row("Changed Name" if changed_name and level == "Level_5" else "Bob Jones",
                  "639181234567", level),
    ]) for level in ("Level_1", "Level_3", "Level_5")]
    return salad_database(), files


class SaladViberTests(unittest.TestCase):
    def run_blast(self, rows, database=None, headers=None):
        prepared = prepare_uploaded_check(database or salad_database(), [
            csv_upload("viber.csv", [headers or BLAST_HEADERS, *rows]),
        ])
        self.assertEqual(prepared.database_error, "")
        self.assertEqual(prepared.files[0].error, "")
        return run_prepared_check(prepared)

    def test_all_three_levels_detect_phone_and_check_full_name_on_same_database_row(self):
        database, files = salad_viber_uploads()
        prepared = prepare_uploaded_check(database, files)
        self.assertEqual(prepared.database_error, "")
        self.assertEqual(prepared.database.column, "Old I.C.")
        self.assertTrue(all(item.source.column == "Phone" for item in prepared.files))
        self.assertTrue(all("FullName" in item.note for item in prepared.files))
        run = run_prepared_check(prepared)
        for result, upload in zip(run.results, files):
            with self.subTest(file=upload.name):
                self.assertEqual(result.error, "")
                self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
                self.assertEqual(result.database_match_column, "Last Called No.")
                self.assertEqual(result.file_identifier, upload.name)
                self.assertFalse(result.flow_stage_required)
                self.assertEqual([row[:3] for row in result.records], [(2, "MATCHED", "2"), (3, "MATCHED", "3")])
                self.assertEqual(result.mappings, [ColumnMapping("FullName", "Name"),
                                                   ColumnMapping("Phone", "Last Called No.", "Phone")])
                self.assertEqual(result.ignored_columns, ["FirstName", "LastName", "CHCode", "OB",
                                                          "Monthly", "Terms", "Promo", "Level Endo"])

    def test_phone_and_name_must_match_the_same_record(self):
        result = self.run_blast([
            salad_row(), salad_row("Alice Smith", "09181234567"),
            salad_row("Alice Smith", "09199999999"), salad_row("Alice Smith", ""), salad_row(),
        ]).results[0]
        self.assertEqual([row[:3] for row in result.records], [
            (2, "MATCHED", "2"), (3, "DIFFERENT", "3"), (4, "NOT FOUND", ""),
            (5, "BLANK KEY", ""), (6, "MATCHED", "2"),
        ])
        self.assertEqual(len(result.differences), 1)
        self.assertEqual(result.differences[0][:2], (3, 3))
        self.assertEqual(result.differences[0][3:], ("Alice Smith", "Bob Jones"))

    def test_shared_database_phone_chooses_the_matching_name_and_reuses_rows(self):
        database = salad_database(rows=[
            ["Alice Smith", "00123", "000001", "alice@example.test", "09171234567", 10000, 1000],
            ["Bob Jones", "456", "000002", "bob@example.test", "+639171234567", 5000, 500],
        ])
        result = self.run_blast([salad_row("Bob Jones"), salad_row(), salad_row("Bob Jones")], database).results[0]
        self.assertEqual([row[2] for row in result.records], ["3", "2", "3"])
        self.assertEqual(result.count("MATCHED"), 3)

    def test_mixed_batch_keeps_chcode_checks_for_other_layouts(self):
        regular = csv_upload("regular.csv", [
            ["CHCODE", "Name", "Phone"],
            ["00123", "Alice Smith", "09181234567"],
        ])
        database, files = salad_viber_uploads()
        prepared = prepare_uploaded_check(database, [files[0], regular, files[1]])
        results = run_prepared_check(prepared).results
        self.assertEqual([result.database_match_column for result in results],
                         ["Last Called No.", "Old I.C.", "Last Called No."])
        self.assertEqual([result.count("MATCHED") for result in results], [2, 0, 2])
        self.assertEqual(results[1].records[0][:3], (2, "DIFFERENT", "2"))
        self.assertEqual(results[1].differences[0][4], "+639171234567")

    def test_excel_report_shows_actual_lookup_original_values_and_checked_fields(self):
        run = self.run_blast([salad_row("Changed Name")])
        book = load_workbook(io.BytesIO(run.report))
        try:
            summary = dict(zip(next(book["Summary"].iter_rows(min_row=8, max_row=8, values_only=True)),
                               next(book["Summary"].iter_rows(min_row=9, max_row=9, values_only=True))))
            self.assertEqual(summary["Match column"], "Phone")
            self.assertEqual(summary["Database match column"], "Last Called No.")
            self.assertEqual(book["File 001"]["G2"].value, "09171234567")
            coverage = {row[0]: row[1:] for row in book["Mappings 001"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(coverage["Phone"][:3], ("Last Called No.", "Phone", "CHECKED"))
            self.assertIn("same row", coverage["Phone"][3])
            self.assertEqual(coverage["FullName"][:3], ("Name", "Text", "CHECKED"))
            self.assertEqual(coverage["CHCode"][2], "NOT CHECKED")
            self.assertEqual(coverage["Monthly"][2], "NOT CHECKED")
            self.assertEqual(book["Differences 001"]["E2"].value, "Alice Smith")
        finally:
            book.close()

    def test_schema_recognition_requires_viber_fields_and_datagrid_phone(self):
        self.assertTrue(is_datagrid_viber_blast(BLAST_HEADERS))
        self.assertTrue(is_datagrid_viber_pair(BLAST_HEADERS, DATABASE_HEADERS))
        self.assertTrue(is_filename_identified_pair(BLAST_HEADERS, DATABASE_HEADERS))
        self.assertEqual(match_column(BLAST_HEADERS), "Phone")
        self.assertEqual(match_column(["CHCODE", "Phone", "FullName"]), "CHCODE")
        for index in range(len(BLAST_HEADERS)):
            with self.subTest(missing=BLAST_HEADERS[index]):
                self.assertFalse(is_datagrid_viber_blast(BLAST_HEADERS[:index] + BLAST_HEADERS[index + 1:]))
        self.assertFalse(is_datagrid_viber_pair(BLAST_HEADERS, DATABASE_HEADERS, "CHCode", "Old I.C."))
        self.assertFalse(is_datagrid_viber_pair(BLAST_HEADERS, DATABASE_HEADERS[:-3]))

    def test_headers_after_preamble_and_letter_case_are_preserved(self):
        headers = [f" {header.upper()} " for header in BLAST_HEADERS]
        upload = xlsx_upload("blast.xlsx", [["Campaign export"], [], headers, salad_row()])
        prepared = prepare_uploaded_check(salad_database(), [upload])
        self.assertEqual(prepared.files[0].error, "")
        self.assertEqual(prepared.files[0].inspection.header_row, 3)
        result = run_prepared_check(prepared).results[0]
        self.assertEqual(result.error, "")
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual(result.source.column, "PHONE")
        self.assertEqual(result.records[0][0], 4)

    def test_missing_or_empty_database_phone_column_is_a_file_error(self):
        for database in (
            salad_database(headers=DATABASE_HEADERS[:4], rows=[["Alice Smith", "00123", "000001", "alice@example.test"]]),
            salad_database(rows=[["Alice Smith", "00123", "000001", "alice@example.test", None, 10000, 1000]]),
        ):
            with self.subTest(database=database.name):
                result = self.run_blast([salad_row()], database).results[0]
                self.assertEqual(result.verdict, "ERROR")
                self.assertIn("Last Called No.", result.error)

    def test_present_stage_and_full_copy_keep_validation(self):
        database = UploadedSource(salad_database(), "Sheet", "Old I.C.")
        staged = UploadedSource(csv_upload("staged.csv", [
            [*BLAST_HEADERS, "FLOW_STAGE"], [*salad_row(), ""],
        ]), "CSV", "Phone")
        result = run_uploaded_check(database, [staged], Options(mapped_fields=True, require_flow_stage=True)).results[0]
        self.assertIn("FLOW_STAGE is blank", result.error)
        source = UploadedSource(csv_upload("viber.csv", [BLAST_HEADERS, salad_row()]), "CSV", "Phone")
        result = run_uploaded_check(database, [source], Options(require_flow_stage=True)).results[0]
        self.assertIn("FLOW_STAGE header", result.error)

    def test_empty_viber_blast_is_an_error(self):
        result = self.run_blast([]).results[0]
        self.assertEqual(result.verdict, "ERROR")
        self.assertIn("no data rows", result.error)


if __name__ == "__main__":
    unittest.main()
