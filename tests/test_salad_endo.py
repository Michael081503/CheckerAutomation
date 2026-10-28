import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from checker import ColumnMapping, is_salad_endo_database, is_salad_endo_pair, match_column
from scripts.check_files import main as check_cli
from test_web_checker import csv_upload, xlsx_upload
from web_checker import prepare_uploaded_check, run_prepared_check


DATABASE_HEADERS = ["Collector", "caseNo", "caseAccountNo", "caseBaseAccountNo",
                    "caseCustomerName", "caseCustomerEmail", "caseCustomerMobileContactNo"]
BLAST_HEADERS = ["email", "cc", "bcc", "indexone_uid", "primary_key", "accountnumber", "chname", "agentcode"]


def endo_database(rows=None, headers=None):
    return xlsx_upload("SALAD - 10282026.xlsx", [headers or DATABASE_HEADERS, *(rows if rows is not None else [
        ["COLLECTOR-A", "case-A", "0012345", "base-A", "Ana Peña", "ana@example.test", "09171234567"],
        ["COLLECTOR-B", "case-B", "0098765", "base-B", "Bob Jones", "bob@example.test", "09181234567"],
        ["COLLECTOR-C", "case-C", "0000001", "base-C", "Carol Smith", None, "09191234567"],
    ])])


def endo_row(account="0012345", name="Ana Peña", email="ana@example.test"):
    return [email, "", "", "", "", account, name, "CAMPAIGN-AGENT"]


def regular_database():
    # Regular PL exports add columns and shift the account / customer positions.
    return xlsx_upload("REGULAR - 10282026.xlsx", [
        ["Collector", "Level_Today", "caseNo", "caseAccountNo", "caseBaseAccountNo",
         "caseCustomerName", "caseCustomerEmail", "PRODUCT_TYPE"],
        ["STAFF-A", "Level_1", "case-A", "0012345", "base-A", "Ana Peña", "ana@example.test", "PL"],
        ["STAFF-B", "Level_3", "case-B", "0098765", "base-B", "Bob Jones", "bob@example.test", "PL"],
    ])


def endo_blast(rows=None, headers=None):
    return csv_upload("SBF SALAD EMAIL BLAST NEW ENDO.csv", [headers or BLAST_HEADERS, *(rows if rows is not None else [
        endo_row(), endo_row("0098765", "Bob Jones", "bob@example.test"),
    ])])


class SaladEndorsementTests(unittest.TestCase):
    def run_blast(self, rows=None, database=None):
        prepared = prepare_uploaded_check(database or endo_database(), [endo_blast(rows)])
        self.assertEqual(prepared.database_error, "")
        self.assertEqual(prepared.files[0].error, "")
        return run_prepared_check(prepared)

    def test_account_number_is_selected_for_this_database_with_blank_legacy_keys(self):
        prepared = prepare_uploaded_check(endo_database(), [endo_blast()])
        self.assertEqual(prepared.database_info.column, "caseAccountNo")
        self.assertEqual(prepared.files[0].source.column, "accountnumber")
        self.assertIn("Staff assignments are not checked", prepared.files[0].note)
        result = run_prepared_check(prepared).results[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual([row[:3] for row in result.records], [(2, "MATCHED", "2"), (3, "MATCHED", "3")])
        self.assertEqual(result.database_match_column, "caseAccountNo")
        self.assertEqual(result.mappings, [ColumnMapping("email", "caseCustomerEmail"),
                                          ColumnMapping("accountnumber", "caseAccountNo"),
                                          ColumnMapping("chname", "caseCustomerName")])
        self.assertEqual(result.ignored_columns, ["cc", "bcc", "indexone_uid", "primary_key", "agentcode"])
        self.assertEqual(result.missing_rows, [])  # Unused endorsement records are allowed.

    def test_regular_pl_export_uses_the_same_fields_despite_extra_columns_and_different_filenames(self):
        base = endo_blast()
        upload = type(base)("28-10-26 SBF PL EMAIL BLAST NEW ENDO.csv", base.data)
        prepared = prepare_uploaded_check(regular_database(), [upload])
        self.assertEqual(prepared.database_error, "")
        self.assertEqual(prepared.files[0].error, "")
        self.assertEqual(prepared.files[0].source.column, "accountnumber")
        self.assertIn("SALAD / PL", prepared.files[0].note)
        result = run_prepared_check(prepared).results[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertEqual([row[:3] for row in result.records], [(2, "MATCHED", "2"), (3, "MATCHED", "3")])
        self.assertEqual(len(result.mappings), 3)
        self.assertIn("agentcode", result.ignored_columns)
        self.assertEqual(result.database_match_column, "caseAccountNo")

    def test_unknown_blank_and_wrong_accounts_and_real_field_differences_fail(self):
        result = self.run_blast([
            endo_row(name="Ana Pe\ufffda"),
            endo_row(email="bob@example.test"),
            endo_row(name="Bob Jones"),
            endo_row("unknown"), endo_row(""), endo_row("12345"),
        ]).results[0]
        self.assertEqual([row[1] for row in result.records],
                         ["MATCHED", "DIFFERENT", "DIFFERENT", "NOT FOUND", "BLANK KEY", "NOT FOUND"])
        self.assertEqual([row[2].split(" ")[0] for row in result.differences], ["email", "chname"])
        self.assertEqual([row[4] for row in result.differences], ["ana@example.test", "Ana Peña"])

    def test_duplicate_accounts_require_all_fields_on_one_row_and_allow_repeat_uploads(self):
        database = endo_database([
            ["COLLECTOR-A", "case-A", "0012345", "base-A", "Ana Peña", "ana@example.test", "09171234567"],
            ["COLLECTOR-B", "case-B", "0012345", "base-B", "Bob Jones", "bob@example.test", "09181234567"],
        ])
        result = self.run_blast([endo_row(), endo_row(name="Bob Jones", email="bob@example.test"),
                                 endo_row(email="bob@example.test"), endo_row()], database).results[0]
        self.assertEqual([row[1] for row in result.records], ["MATCHED", "MATCHED", "DIFFERENT", "MATCHED"])
        self.assertEqual([row[2] for row in result.records], ["2", "3", "2", "2"])

    def test_schema_is_specific_and_other_layouts_keep_their_lookup_rules(self):
        self.assertTrue(is_salad_endo_database(DATABASE_HEADERS))
        self.assertFalse(is_salad_endo_database(["caseAccountNo"]))
        self.assertTrue(is_salad_endo_pair(BLAST_HEADERS, DATABASE_HEADERS))
        self.assertEqual(match_column(BLAST_HEADERS), "indexone_uid")
        datagrid = ["Name", "Old I.C.", "Account No.", "Email"]
        self.assertEqual(match_column(BLAST_HEADERS, datagrid), "indexone_uid")
        self.assertFalse(is_salad_endo_pair(BLAST_HEADERS, datagrid))
        self.assertEqual(match_column(["CHCODE", *BLAST_HEADERS], DATABASE_HEADERS), "CHCODE")
        self.assertFalse(is_salad_endo_pair(["CHCODE", *BLAST_HEADERS], DATABASE_HEADERS))
        self.assertEqual(match_column(BLAST_HEADERS, ["CHCODE", *DATABASE_HEADERS]), "indexone_uid")
        self.assertFalse(is_salad_endo_pair(BLAST_HEADERS, DATABASE_HEADERS, "primary_key", "caseAccountNo"))

    def test_case_reordered_headers_and_preambles_skip_all_header_rows(self):
        db_headers = [header.upper() for header in reversed(DATABASE_HEADERS)]
        database = xlsx_upload("future-endo.xlsx", [["Endorsement export"], [], db_headers,
            list(reversed(["STAFF", "case-A", "0012345", "base-A", "Ana Peña", "ana@example.test", "09171234567"]))])
        upload = xlsx_upload("campaign.xlsx", [["Campaign export"], [],
            [header.upper() for header in reversed(BLAST_HEADERS)], list(reversed(endo_row()))])
        prepared = prepare_uploaded_check(database, [upload])
        self.assertEqual(prepared.database_error, "")
        self.assertEqual(prepared.files[0].error, "")
        self.assertEqual(prepared.database_info.header_row, 3)
        self.assertEqual(prepared.files[0].inspection.header_row, 3)
        result = run_prepared_check(prepared).results[0]
        self.assertEqual(result.records[0][:3], (4, "MATCHED", "4"))
        self.assertEqual(len(result.records), 1)

    def test_empty_campaign_blank_database_keys_and_present_stage_keep_validation(self):
        self.assertIn("no data rows", self.run_blast([]).results[0].error)
        database = endo_database([["STAFF", "case-A", None, "base-A", "Ana", "ana@example.test", ""]])
        prepared = prepare_uploaded_check(database, [endo_blast()])
        self.assertIn("blank", prepared.database_error)
        staged = csv_upload("staged.csv", [[*BLAST_HEADERS, "FLOW_STAGE"], [*endo_row(), ""]])
        result = run_prepared_check(prepare_uploaded_check(endo_database(), [staged])).results[0]
        self.assertIn("FLOW_STAGE is blank", result.error)

    def test_report_preserves_input_and_discloses_only_the_selected_three_fields(self):
        run = self.run_blast()
        book = load_workbook(io.BytesIO(run.report))
        try:
            summary = dict(zip(next(book["Summary"].iter_rows(min_row=8, max_row=8, values_only=True)),
                               next(book["Summary"].iter_rows(min_row=9, max_row=9, values_only=True))))
            self.assertEqual(summary["Match column"], "accountnumber")
            self.assertEqual(summary["Database match column"], "caseAccountNo")
            self.assertEqual(book["File 001"]["I2"].value, "0012345")
            coverage = {row[0]: row[1:] for row in book["Mappings 001"].iter_rows(min_row=2, values_only=True)}
            self.assertEqual(sum(row[2] == "CHECKED" for row in coverage.values()), 3)
            self.assertEqual(coverage["agentcode"][2], "NOT CHECKED")
            self.assertIn("not compared with Collector", coverage["agentcode"][3])
            self.assertEqual(coverage["primary_key"][2], "NOT CHECKED")
            self.assertIn("same caseAccountNo record", coverage["email"][3])
        finally:
            book.close()

    def test_missing_required_campaign_columns_are_errors_without_a_key_only_fallback(self):
        for missing in ("accountnumber", "chname", "email"):
            with self.subTest(missing=missing):
                position = BLAST_HEADERS.index(missing)
                headers = [header for header in BLAST_HEADERS if header != missing]
                values = endo_row()
                del values[position]
                prepared = prepare_uploaded_check(endo_database(), [csv_upload("incomplete.csv", [headers, values])])
                self.assertEqual(prepared.database_error, "")
                self.assertTrue(prepared.files[0].error)
                self.assertEqual(run_prepared_check(prepared).results[0].verdict, "ERROR")

    def test_command_line_detects_the_pair_and_preserves_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            database, upload = endo_database(), endo_blast()
            database_path, file_path = root / database.name, root / upload.name
            database_path.write_bytes(database.data)
            file_path.write_bytes(upload.data)
            output = root / "report.xlsx"
            arguments = ["check_files.py", "--database", str(database_path), "--files", str(file_path), "--output", str(output)]
            stdout = io.StringIO()
            with patch.object(sys, "argv", arguments), redirect_stdout(stdout):
                self.assertEqual(check_cli(), 0)
            result = json.loads(stdout.getvalue().splitlines()[0])
            self.assertEqual((result["file_key"], result["database_key"], result["matched"], result["checked_fields"]),
                             ("accountnumber", "caseAccountNo", 2, 3))
            self.assertEqual(database_path.read_bytes(), database.data)
            self.assertEqual(file_path.read_bytes(), upload.data)
            # An explicit override remains explicit; it cannot switch to accountnumber.
            with patch.object(sys, "argv", [*arguments, "--file-key", "primary_key"]), redirect_stdout(io.StringIO()):
                self.assertEqual(check_cli(), 1)


if __name__ == "__main__":
    unittest.main()
