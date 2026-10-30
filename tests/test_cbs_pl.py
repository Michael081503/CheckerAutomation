"""CBS PL recovery/write-off exports with an unused legacy identifier."""

import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from checker import ColumnMapping, Options, Source, check_files, read_table
from web_checker import prepare_uploaded_check, run_prepared_check
from test_web_checker import csv_upload, xlsx_upload


HEADERS = ["email", "cc", "bcc", "indexone_uid", "primary_key", "accountnumber", "chname", "campaignbucket"]
ACCOUNTS = [
    ["Alice Smith", "00123", "000001", "alice@example.test", "RECOVERY ACCOUNTS"],
    ["Bob Jones", "00456", "000002", "bob@example.test", "RECOVERY ACCOUNTS"],
    ["Carol Tan", "00789", "000003", "carol@example.test", "WRITE OFF ACCOUNTS"],
    ["Dan Lim", "00987", "000004", "dan@example.test", "WRITE OFF ACCOUNTS"],
]


def cbs_pl_account_uploads():
    database = xlsx_upload("DataGrid(1).xlsx", [
        ["Name", "Old I.C.", "Account No.", "Email", "Account Type"], *ACCOUNTS,
    ])
    files = []
    for bucket in ("RECOVERY ACCOUNTS", "WRITE OFF ACCOUNTS"):
        rows = [[email, "", "", "", key, account, name, category]
                for name, key, account, email, category in ACCOUNTS if category == bucket]
        files.append(csv_upload(f"CBS PL EMAIL BLAST {bucket}.csv", [HEADERS, *rows]))
    return database, files


class CbsPlTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.database_upload, self.files = cbs_pl_account_uploads()
        self.database = self.source(self.database_upload, "Old I.C.")
        self.options = Options(mapped_fields=True, require_flow_stage=True)

    def source(self, upload, key="primary_key", mappings=None):
        path = self.folder / upload.name
        path.write_bytes(upload.data)
        return Source(path, read_table(path).sheet, key, mappings)

    def campaign(self, rows, key="primary_key", mappings=None):
        return self.source(csv_upload("campaign.csv", [HEADERS, *rows]), key, mappings)

    def row(self, index=0, **changes):
        name, key, account, email, bucket = ACCOUNTS[index]
        row = dict(zip(HEADERS, [email, "", "", "", key, account, name, bucket]))
        row.update(changes)
        return [row[header] for header in HEADERS]

    def check(self, *sources):
        return check_files(self.database, list(sources), self.options)

    def test_both_buckets_match_five_fields_and_report_unused_identifier(self):
        prepared = prepare_uploaded_check(self.database_upload, self.files)
        self.assertEqual(prepared.database_error, "")
        self.assertTrue(all(not item.error and item.source.column == "primary_key" for item in prepared.files))
        run = run_prepared_check(prepared)
        for result in run.results:
            self.assertEqual(result.error, "")
            self.assertEqual((result.verdict, result.count("MATCHED")), ("MAPPED ROWS MATCH", 2))
            self.assertEqual(len(result.mappings), 5)
            self.assertEqual(result.database_match_column, "Old I.C.")
            self.assertEqual(result.ignored_columns, ["cc", "bcc", "indexone_uid"])
            self.assertIn("blank in every row", result.unchecked_reasons["indexone_uid"])
            self.assertEqual([row[0] for row in result.records], [2, 3])
            self.assertTrue(all(row[3][3] == "" for row in result.records))
        book = load_workbook(io.BytesIO(run.report), read_only=True, data_only=True)
        try:
            for index in (1, 2):
                coverage = {row[0]: row[1:] for row in book[f"Mappings {index:03d}"].iter_rows(min_row=2, values_only=True)}
                self.assertEqual(coverage["indexone_uid"][2], "NOT CHECKED")
                self.assertIn("blank in every row", coverage["indexone_uid"][3])
                self.assertEqual(coverage["primary_key"][:3], ("Old I.C.", "Text", "CHECKED"))
                detail = list(book[f"File {index:03d}"].values)
                self.assertEqual(detail[0][3:], tuple(HEADERS))
                self.assertTrue(detail[1][7].startswith("00"))
        finally:
            book.close()

    def test_whitespace_only_secondary_id_is_unused(self):
        result = self.check(self.campaign([self.row(indexone_uid=" \t ")]))[0]
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertIn("indexone_uid", result.unchecked_reasons)

    def test_partly_populated_secondary_id_checks_every_row(self):
        result = self.check(self.campaign([self.row(indexone_uid="00123"), self.row(1)]))[0]
        self.assertEqual([row[1] for row in result.records], ["MATCHED", "DIFFERENT"])
        self.assertEqual(result.unchecked_reasons, {})
        self.assertEqual(len(result.mappings), 6)
        self.assertEqual(result.differences[0][3:], ("", "00456"))

    def test_populated_wrong_secondary_id_is_still_different(self):
        result = self.check(self.campaign([self.row(indexone_uid="00456")]))[0]
        self.assertEqual(result.count("DIFFERENT"), 1)
        self.assertIn("indexone_uid", result.differences[0][2])

    def test_blank_and_unknown_lookup_keys_never_fall_back(self):
        source = self.campaign([self.row(primary_key=""), self.row(1, primary_key="99999")])
        result = self.check(source)[0]
        self.assertEqual([row[1] for row in result.records], ["BLANK KEY", "NOT FOUND"])

    def test_wrong_email_name_account_and_bucket_still_fail_on_same_record(self):
        for header, other_value in zip(HEADERS, self.row(1)):
            if header not in {"email", "chname", "accountnumber", "campaignbucket"}:
                continue
            if header == "campaignbucket":
                other_value = "WRITE OFF ACCOUNTS"
            with self.subTest(header=header):
                result = self.check(self.campaign([self.row(**{header: other_value})]))[0]
                self.assertEqual(result.count("DIFFERENT"), 1)
                self.assertEqual(result.records[0][2], "2")
                self.assertEqual(len(result.differences), 1)
                self.assertIn(header, result.differences[0][2])

    def test_blank_email_name_account_and_bucket_remain_checked(self):
        for header in ("email", "chname", "accountnumber", "campaignbucket"):
            with self.subTest(header=header):
                result = self.check(self.campaign([self.row(**{header: ""})]))[0]
                self.assertEqual(result.count("DIFFERENT"), 1)
                self.assertIn(header, result.differences[0][2])

    def test_explicit_secondary_id_mapping_is_honored(self):
        mappings = (ColumnMapping("primary_key", "Old I.C."), ColumnMapping("indexone_uid", "Old I.C."))
        result = self.check(self.campaign([self.row()], mappings=mappings))[0]
        self.assertEqual(result.count("DIFFERENT"), 1)
        self.assertEqual(result.unchecked_reasons, {})

    def test_rule_works_in_reverse_without_skipping_selected_uid(self):
        source = self.campaign([self.row(primary_key="", indexone_uid="00123")], key="indexone_uid")
        result = self.check(source)[0]
        self.assertEqual(result.count("MATCHED"), 1)
        self.assertEqual(result.ignored_columns, ["cc", "bcc", "primary_key"])
        self.assertIn("primary_key", result.unchecked_reasons)

    def test_non_datagrid_identifiers_keep_their_existing_comparisons(self):
        database = self.source(csv_upload("ordinary_db.csv", [["CHCODE", "Name"], ["00123", "Alice"]]), "CHCODE")
        source = self.source(csv_upload("ordinary.csv", [
            ["CHCODE", "Name", "indexone_uid", "FLOW_STAGE"], ["00123", "Alice", "", "MAIL"],
        ]), "CHCODE")
        result = check_files(database, [source], self.options)[0]
        self.assertEqual(result.count("DIFFERENT"), 1)
        self.assertEqual(result.unchecked_reasons, {})


if __name__ == "__main__":
    unittest.main()
