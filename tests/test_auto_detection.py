import csv
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from checker import ColumnMapping
from test_web_checker import csv_upload, database_upload, datagrid_upload, xlsx_upload
from web_checker import Upload, prepare_uploaded_check, run_prepared_check


def blast_upload(name, identifiers, *, campaign=False):
    headers = ["email", "cc", "bcc", "indexone_uid", "primary_key", "chname", "accountnumber"]
    if campaign:
        headers.append("campaignbucket")
    rows = [headers]
    accounts = {
        "00123": ("alice@example.test", "Alice Smith", "001234567890"),
        "456": ("bob@example.test", "Bob Jones", "009876543210"),
    }
    for uid, primary, account in identifiers:
        email, name_value, account_number = accounts[account]
        row = [email, "", "", uid, primary, name_value, account_number]
        if campaign:
            row.append("PL")
        rows.append(row)
    return csv_upload(name, rows)


def workbook_upload(name, sheets):
    stream = io.BytesIO()
    book = Workbook()
    try:
        for index, (title, rows) in enumerate(sheets):
            sheet = book.active if index == 0 else book.create_sheet()
            sheet.title = title
            for row in rows:
                sheet.append(row)
        book.save(stream)
    finally:
        book.close()
    return Upload(name, stream.getvalue())


class AutoDetectionTests(unittest.TestCase):
    def assert_prepared(self, prepared):
        self.assertEqual(prepared.database_error, "")
        self.assertIsNotNone(prepared.database)
        self.assertTrue(prepared.files)
        for item in prepared.files:
            self.assertEqual(item.error, "", item.source.upload.name)
            self.assertIsNotNone(item.inspection)

    def test_database_hits_choose_the_key_independently_for_each_file(self):
        primary = blast_upload("primary.csv", [("other-A", "00123", "00123"), ("other-B", "456", "456")])
        uid = blast_upload("uid.csv", [("00123", "other-A", "00123"), ("456", "other-B", "456")])
        prepared = prepare_uploaded_check(datagrid_upload(), [primary, uid])
        self.assert_prepared(prepared)
        self.assertEqual(prepared.database_info.header_row, 3)
        self.assertEqual([item.source.column for item in prepared.files], ["primary_key", "indexone_uid"])
        self.assertTrue(all(item.inspection.has_headers for item in prepared.files))
        run = run_prepared_check(prepared)
        for result in run.results:
            self.assertEqual(result.count("DIFFERENT"), 2)
            self.assertEqual(result.count("NOT FOUND"), 0)

    def test_identical_matching_identifier_columns_prefer_primary_and_compare_both(self):
        upload = blast_upload("SBF.csv", [("00123", "00123", "00123"), ("456", "456", "456")])
        prepared = prepare_uploaded_check(datagrid_upload(), [upload])
        self.assert_prepared(prepared)
        self.assertEqual(prepared.files[0].source.column, "primary_key")
        result = run_prepared_check(prepared).results[0]
        self.assertEqual(result.verdict, "MAPPED ROWS MATCH")
        self.assertIn(ColumnMapping("primary_key", "Old I.C."), result.mappings)
        self.assertIn(ColumnMapping("indexone_uid", "Old I.C."), result.mappings)

    def test_unequal_tied_or_unmatched_identifiers_keep_the_schema_preference(self):
        for campaign, expected in ((False, "indexone_uid"), (True, "primary_key")):
            for identifiers in (
                [("00123", "456", "00123"), ("456", "00123", "456")],
                [("unknown-uid", "unknown-primary", "00123")],
            ):
                with self.subTest(campaign=campaign, identifiers=identifiers):
                    upload = blast_upload("tied.csv", identifiers, campaign=campaign)
                    prepared = prepare_uploaded_check(datagrid_upload(), [upload])
                    self.assert_prepared(prepared)
                    self.assertEqual(prepared.files[0].source.column, expected)

    def test_explicit_chcode_header_takes_precedence_over_other_identifiers(self):
        upload = csv_upload("explicit.csv", [
            ["CHCODE", "primary_key", "indexone_uid", "FLOW_STAGE"],
            ["missing-account", "00123", "00123", "MAIL"],
        ])
        prepared = prepare_uploaded_check(database_upload(), [upload])
        self.assert_prepared(prepared)
        self.assertEqual(prepared.files[0].source.column, "CHCODE")
        result = run_prepared_check(prepared).results[0]
        self.assertEqual(result.count("NOT FOUND"), 1)
        self.assertEqual(result.records[0][3][0], "missing-account")

    def test_selected_identifier_never_falls_back_per_row_and_secondary_is_checked(self):
        upload = blast_upload("mixed.csv", [
            ("other-A", "00123", "00123"),
            ("other-B", "456", "456"),
            ("00123", "missing-primary", "00123"),
        ])
        prepared = prepare_uploaded_check(datagrid_upload(), [upload])
        self.assert_prepared(prepared)
        self.assertEqual(prepared.files[0].source.column, "primary_key")
        result = run_prepared_check(prepared).results[0]
        self.assertEqual([record[1] for record in result.records], ["DIFFERENT", "DIFFERENT", "NOT FOUND"])
        self.assertIn(ColumnMapping("indexone_uid", "Old I.C."), result.mappings)
        self.assertIn(ColumnMapping("primary_key", "Old I.C."), result.mappings)
        self.assertEqual([difference[3] for difference in result.differences], ["other-A", "other-B"])

    def test_generic_lookup_headers_use_filename_and_validate_rows_without_stage(self):
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            for key in ("CHCODE", "primary_key", "indexone_uid"):
                for key_position in range(3):
                    with self.subTest(format=suffix, key=key, position=key_position):
                        headers = ["Name", "Delivery label"]
                        headers.insert(key_position, f" {key.upper()} ")
                        rows = [["Campaign export"], [], headers]
                        for identifier, name in (("00123", "Alice Smith"), ("456", "Changed Name"),
                                                 ("00999", "Unknown Account"), ("", "Blank Account")):
                            values = [name, "EMAIL"]
                            values.insert(key_position, identifier)
                            rows.append(values)
                        upload = make_upload(f"generic.{suffix}", rows)
                        prepared = prepare_uploaded_check(database_upload(), [upload])
                        self.assert_prepared(prepared)
                        item = prepared.files[0]
                        self.assertTrue(item.source.has_headers)
                        self.assertEqual(item.source.column, key.upper())
                        self.assertEqual(item.inspection.header_row, 3)
                        result = run_prepared_check(prepared).results[0]
                        self.assertEqual(result.error, "")
                        self.assertEqual(result.file_identifier, upload.name)
                        self.assertFalse(result.flow_stage_required)
                        self.assertEqual([record[:2] for record in result.records], [
                            (4, "MATCHED"), (5, "DIFFERENT"), (6, "NOT FOUND"), (7, "BLANK KEY"),
                        ])
                        self.assertIn(ColumnMapping(key.upper(), "CHCODE"), result.mappings)
                        self.assertIn(ColumnMapping("Name", "Name"), result.mappings)
                        self.assertEqual([difference[3] for difference in result.differences], ["Changed Name"])

    def test_generic_identifier_columns_are_both_checked_with_one_lookup_for_the_file(self):
        for chosen in ("primary_key", "indexone_uid"):
            with self.subTest(chosen=chosen):
                rows = [["primary_key", "Name", "indexone_uid"]]
                for key, secondary, name in (("00123", "other-A", "Alice Smith"),
                                             ("456", "other-B", "Bob Jones"),
                                             ("missing-key", "00123", "Alice Smith")):
                    primary, uid = (key, secondary) if chosen == "primary_key" else (secondary, key)
                    rows.append([primary, name, uid])
                prepared = prepare_uploaded_check(database_upload(), [csv_upload("identifiers.csv", rows)])
                self.assert_prepared(prepared)
                self.assertEqual(prepared.files[0].source.column, chosen)
                result = run_prepared_check(prepared).results[0]
                self.assertEqual(result.error, "")
                self.assertEqual([record[1] for record in result.records], ["DIFFERENT", "DIFFERENT", "NOT FOUND"])
                self.assertIn(ColumnMapping("primary_key", "CHCODE"), result.mappings)
                self.assertIn(ColumnMapping("indexone_uid", "CHCODE"), result.mappings)
                self.assertEqual([difference[3] for difference in result.differences], ["other-A", "other-B"])

    def test_generic_files_still_validate_present_flow_stage_values(self):
        for key in ("CHCODE", "primary_key", "indexone_uid"):
            for stages, expected_error in ((("", "MAIL"), "FLOW_STAGE is blank"),
                                           (("MAIL", "NEW_ENDO"), "multiple FLOW_STAGE identifiers"),
                                           (("mail", "MAIL"), "")):
                with self.subTest(key=key, stages=stages):
                    upload = csv_upload("staged.csv", [
                        ["Name", "FLOW_STAGE", key],
                        ["Alice Smith", stages[0], "00123"],
                        ["Bob Jones", stages[1], "456"],
                    ])
                    prepared = prepare_uploaded_check(database_upload(), [upload])
                    self.assert_prepared(prepared)
                    result = run_prepared_check(prepared).results[0]
                    if expected_error:
                        self.assertIn(expected_error, result.error)
                        self.assertEqual(result.records, [])
                    else:
                        self.assertEqual(result.error, "")
                        self.assertEqual(result.verdict, "EXACT MATCH")
                        self.assertEqual(result.file_identifier, "MAIL")
                    self.assertTrue(result.flow_stage_required)

    def test_headerless_files_include_row_one_and_preserve_unknown_and_blank_keys_in_report(self):
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            with self.subTest(format=suffix):
                upload = make_upload(f"raw.{suffix}", [
                    ["", "00123", "First original row"],
                    ["Unknown account", "00999", "Keep unknown row"],
                    ["Blank account", "", "Keep blank row"],
                ])
                prepared = prepare_uploaded_check(database_upload(), [upload])
                self.assert_prepared(prepared)
                self.assertFalse(prepared.files[0].source.has_headers)
                self.assertIsNone(prepared.files[0].inspection.header_row)
                run = run_prepared_check(prepared)
                result = run.results[0]
                self.assertEqual([record[:3] for record in result.records], [
                    (1, "MATCHED", "2"), (2, "NOT FOUND", ""), (3, "BLANK KEY", ""),
                ])
                self.assertEqual(result.records[1][3], ("Unknown account", "00999", "Keep unknown row"))
                book = load_workbook(io.BytesIO(run.report))
                try:
                    self.assertEqual(book["File 001"]["A2"].value, 1)
                    self.assertEqual(book["File 001"]["E2"].value, "00123")
                    self.assertEqual(book["File 001"]["E2"].data_type, "s")
                    self.assertEqual(book["File 001"]["F2"].value, "First original row")
                    self.assertEqual(book["File 001"]["E3"].value, "00999")
                    self.assertEqual(book["File 001"]["B4"].value, "BLANK KEY")
                finally:
                    book.close()

    def test_literal_chcode_in_headerless_data_does_not_discard_original_rows(self):
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            for keys in (("CHCODE", "00999"), ("00999", "CHCODE"), ("00123", "CHCODE")):
                with self.subTest(format=suffix, keys=keys):
                    rows = [[f"Account {index}", key, f"Original {index}"] for index, key in enumerate(keys, 1)]
                    prepared = prepare_uploaded_check(database_upload(), [make_upload(f"literal.{suffix}", rows)])
                    self.assert_prepared(prepared)
                    self.assertFalse(prepared.files[0].source.has_headers)
                    result = run_prepared_check(prepared).results[0]
                    self.assertEqual([record[0] for record in result.records], [1, 2])
                    self.assertEqual([record[3][1] for record in result.records], list(keys))
                    self.assertEqual(result.count("NOT FOUND"), 1 if keys[0] == "00123" else 2)

    def test_supported_headers_after_preambles_and_csv_delimiters_are_detected(self):
        rows = [["Campaign export"], [], ["Name", "CHCODE", "FLOW_STAGE"], ["Alice Smith", "00123", "MAIL"]]
        uploads = [xlsx_upload("preamble.xlsx", rows)]
        for delimiter in (",", ";", "\t", "|"):
            stream = io.StringIO(newline="")
            csv.writer(stream, delimiter=delimiter).writerows(rows)
            uploads.append(Upload(f"preamble_{ord(delimiter)}.csv", stream.getvalue().encode("utf-8-sig")))
        prepared = prepare_uploaded_check(database_upload(), uploads)
        self.assert_prepared(prepared)
        self.assertTrue(all(item.inspection.has_headers for item in prepared.files))
        self.assertEqual([item.inspection.header_row for item in prepared.files], [3] * len(uploads))
        for result in run_prepared_check(prepared).results:
            self.assertEqual([record[:3] for record in result.records], [(4, "MATCHED", "2")])

    def test_headered_files_retain_unknown_and_blank_selected_keys(self):
        upload = blast_upload("all_rows.csv", [
            ("00123", "00123", "00123"),
            ("00999", "00999", "00123"),
            ("", "", "00123"),
        ])
        prepared = prepare_uploaded_check(datagrid_upload(), [upload])
        self.assert_prepared(prepared)
        self.assertEqual(prepared.files[0].source.column, "primary_key")
        result = run_prepared_check(prepared).results[0]
        self.assertEqual([record[1] for record in result.records], ["MATCHED", "NOT FOUND", "BLANK KEY"])
        self.assertEqual([record[0] for record in result.records], [2, 3, 4])

    def test_malformed_or_unsupported_named_headers_cannot_bypass_validation_as_raw_data(self):
        cases = [
            [["Name", "Email", "Account ID"], ["Alice Smith", "alice@example.test", "00123"]],
            [["CHCODE", "", "FLOW_STAGE"], ["00123", "Alice Smith", "MAIL"]],
        ]
        cases += [[["Campaign export"], [], *rows] for rows in cases]
        for make_upload, suffix in ((csv_upload, "csv"), (xlsx_upload, "xlsx")):
            for rows in cases:
                with self.subTest(format=suffix, headers=rows[-2], preamble=len(rows) > 2):
                    upload = make_upload(f"invalid.{suffix}", rows)
                    prepared = prepare_uploaded_check(database_upload(), [upload])
                    self.assertTrue(prepared.files[0].error)
                    self.assertTrue(prepared.files[0].source.has_headers)
                    result = run_prepared_check(prepared).results[0]
                    self.assertEqual(result.verdict, "ERROR")
                    self.assertEqual(result.records, [])

    def test_only_eligible_worksheet_is_selected_while_empty_and_notes_are_ignored(self):
        database = workbook_upload("database.xlsx", [
            ("Notes", [["Database notes"], ["Generated export"]]),
            ("Empty", []),
            ("Accounts", [["Name", "CHCODE"], ["Alice Smith", "00123"]]),
        ])
        upload = workbook_upload("campaign.xlsx", [
            ("Empty", []),
            ("Notes", [["Campaign notes"], ["Use current database"]]),
            ("Delivery", [["Name", "CHCODE", "FLOW_STAGE"], ["Alice Smith", "00123", "MAIL"]]),
        ])
        prepared = prepare_uploaded_check(database, [upload])
        self.assert_prepared(prepared)
        self.assertEqual(prepared.database.sheet, "Accounts")
        self.assertEqual(prepared.files[0].source.sheet, "Delivery")
        self.assertEqual(run_prepared_check(prepared).results[0].count("MATCHED"), 1)

    def test_multiple_eligible_worksheets_are_reported_instead_of_guessing(self):
        database_rows = [["CHCODE", "Name"], ["00123", "Alice Smith"]]
        ambiguous_database = workbook_upload("ambiguous.xlsx", [("One", database_rows), ("Two", database_rows)])
        prepared = prepare_uploaded_check(ambiguous_database, [])
        self.assertIsNone(prepared.database)
        self.assertIn("Multiple database worksheets", prepared.database_error)
        with self.assertRaisesRegex(ValueError, "Multiple database worksheets"):
            run_prepared_check(prepared)
        target_rows = [["CHCODE", "FLOW_STAGE"], ["00123", "MAIL"]]
        ambiguous_target = workbook_upload("ambiguous.xlsx", [("One", target_rows), ("Two", target_rows)])
        prepared = prepare_uploaded_check(database_upload(), [ambiguous_target])
        self.assertIn("Multiple worksheets", prepared.files[0].error)
        self.assertEqual(run_prepared_check(prepared).results[0].verdict, "ERROR")

    def test_duplicate_filenames_keep_independent_uploaded_data(self):
        database = csv_upload("same.csv", [["CHCODE", "Name"], ["00123", "Alice Smith"]])
        first = csv_upload("same.csv", [["CHCODE", "FLOW_STAGE"], ["00123", "MAIL"]])
        second = csv_upload("same.csv", [["CHCODE", "FLOW_STAGE"], ["missing", "MAIL"]])
        prepared = prepare_uploaded_check(database, [first, second])
        self.assert_prepared(prepared)
        run = run_prepared_check(prepared)
        self.assertEqual(run.results[0].count("MATCHED"), 1)
        self.assertEqual(run.results[1].count("NOT FOUND"), 1)
        self.assertEqual(run.results[0].records[0][3][0], "00123")
        self.assertEqual(run.results[1].records[0][3][0], "missing")

    def test_invalid_upload_does_not_block_valid_file_and_remains_in_download(self):
        invalid = Upload("broken.xlsx", b"not an Excel workbook")
        valid = blast_upload("valid.csv", [("00123", "00123", "00123")])
        prepared = prepare_uploaded_check(datagrid_upload(), [invalid, valid])
        self.assertTrue(prepared.files[0].error)
        self.assertEqual(prepared.files[1].error, "")
        run = run_prepared_check(prepared)
        self.assertEqual(run.results[0].verdict, "ERROR")
        self.assertEqual(run.results[1].verdict, "MAPPED ROWS MATCH")
        book = load_workbook(io.BytesIO(run.report))
        try:
            self.assertEqual(book["Summary"]["B9"].value, "broken.xlsx")
            self.assertEqual(book["Summary"]["B10"].value, "valid.csv")
            self.assertIn("ERROR", next(book["Summary"].iter_rows(min_row=9, max_row=9, values_only=True)))
        finally:
            book.close()

    def test_detection_temporary_files_are_cleaned_after_success_and_errors(self):
        original = tempfile.TemporaryDirectory
        directories = []

        def tracked(*args, **kwargs):
            directory = original(*args, **kwargs)
            directories.append(Path(directory.name))
            return directory

        good = csv_upload("good.csv", [["CHCODE", "FLOW_STAGE"], ["00123", "MAIL"]])
        with patch("web_checker.tempfile.TemporaryDirectory", side_effect=tracked):
            successful = prepare_uploaded_check(database_upload(), [good])
            self.assert_prepared(successful)
            run_prepared_check(successful)
            failed_file = prepare_uploaded_check(database_upload(), [Upload("broken.xlsx", b"invalid")])
            self.assertTrue(failed_file.files[0].error)
            failed_database = prepare_uploaded_check(Upload("broken.xlsx", b"invalid"), [good])
            self.assertTrue(failed_database.database_error)
        self.assertEqual(len(directories), 4)
        self.assertTrue(all(not directory.exists() for directory in directories))


if __name__ == "__main__":
    unittest.main()
